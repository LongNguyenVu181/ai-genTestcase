# -*- coding: utf-8 -*-
from __future__ import annotations

import io
import re
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


DETAIL_HEADERS = [
    "ID", "Name", "PreConditions", "Importance", "Step", "Data test",
    "Expected Result", "Actual Result", "Lần 1", "Lần 2", "Lần 3",
    "Kết quả hiện tại", "Ghi chú", "Mã lỗi", "QC viết testcase",
    "Sprint viết testcase", "QC thực hiện test", "Sprint thực hiện test",
    "Người review", "Ngày review", "Nội dung review",
    "Case cần Auto (Yes/No)", "Case Đã Auto (Yes/No)",
    "Smoke test (Yes/No)", "Regression test (Yes/No)",
    "TCs Out of date (Yes/No)", "Ngày TCs Out of date",
]

EXCEL_FIELDS = [
    "id", "name", "preCondition", "importance", "steps", "testData",
    "expectedResult", "actualResult", "run1", "run2", "run3",
    "currentResult", "note", "errorCode", "qcWriter", "sprintWriter",
    "qcExecutor", "sprintExecutor", "reviewer", "reviewDate", "reviewContent",
    "needAuto", "automated", "smoke", "regression", "outdated", "outdatedDate",
]


def _sheet_title(title: str, used: set[str]) -> str:
    name = re.sub(r'[\\/*?:\[\]]', "_", str(title or "Test Cases")).strip() or "Test Cases"
    name = name[:31]
    base = name
    i = 2
    while name.casefold() in used:
        suffix = f"_{i}"
        name = (base[:31-len(suffix)] + suffix)[:31]
        i += 1
    used.add(name.casefold())
    return name


def _cell_value(tc: dict, field: str):
    value = tc.get(field, "")
    if field == "steps":
        return "\n".join(f"{i+1}. {step}" for i, step in enumerate(value or []))
    return "" if value is None else value


def build_excel(testcases: list[dict], *, workbook_title: str = "Test Cases") -> bytes:
    if not testcases:
        raise ValueError("Không có testcase để xuất Excel.")

    wb = Workbook()
    wb.remove(wb.active)

    fill_green = PatternFill(start_color="81C784", fill_type="solid")
    fill_light_green = PatternFill(start_color="C8E6C9", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", fill_type="solid")
    fill_header_bg = PatternFill(start_color="4DD0E1", fill_type="solid")
    fill_pink = PatternFill(start_color="F8BBD0", fill_type="solid")
    fill_blue_exec = PatternFill(start_color="90CAF9", fill_type="solid")

    thin = Side(style="thin", color="B0BEC5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    font_title = Font(name="Segoe UI", size=11, bold=True)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_body = Font(name="Segoe UI", size=10)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    group_headers = {
        "A11": "Test Case", "E11": "Steps", "I11": "Chrome",
        "L11": "Kết quả hiện tại", "M11": "Ghi chú", "O11": "QC viết testcase",
        "P11": "Sprint viết testcase", "R11": "Sprint thực hiện test",
        "V11": "Case cần Auto (Yes/No)", "W11": "Case Đã Auto (Yes/No)",
        "X11": "Smoke test (Yes/No)", "Y11": "Regression test (Yes/No)",
        "Z11": "TCs Out of date (Yes/No)", "AA11": "Ngày TCs Out of date",
    }
    widths = {
        'A': 14, 'B': 52, 'C': 38, 'D': 12, 'E': 52, 'F': 30, 'G': 58, 'H': 30,
        'I': 10, 'J': 10, 'K': 10, 'L': 16, 'M': 28, 'N': 16, 'O': 18, 'P': 18,
        'Q': 18, 'R': 18, 'S': 18, 'T': 16, 'U': 34, 'V': 18, 'W': 18, 'X': 16,
        'Y': 18, 'Z': 18, 'AA': 20,
    }

    # Preserve incoming deterministic order and split sheets by screen.
    by_screen: OrderedDict[str, list[dict]] = OrderedDict()
    for tc in testcases:
        by_screen.setdefault(str(tc.get("screen") or workbook_title), []).append(tc)

    used = set()
    summary_rows = []
    total = 0

    for screen_name, rows in by_screen.items():
        ws = wb.create_sheet(_sheet_title(screen_name, used))
        ws.views.sheetView[0].showGridLines = True
        ws["D1"] = "KỊCH BẢN KIỂM THỬ *"
        ws["D1"].font = font_title
        ws["D1"].alignment = center
        ws["C2"] = "Tên màn hình/Tên chức năng"
        ws["C2"].font = font_bold
        ws["D2"] = screen_name
        ws["C4"] = "Mã Testcase"
        ws["C4"].font = font_bold
        ws["D4"] = "TC_"
        ws["B13"] = "Màn hình test"
        ws["C13"] = "Link test:"

        for r in (1, 2, 4):
            for col in ("C", "D"):
                ws[f"{col}{r}"].border = border

        for col_idx in range(1, 28):
            c = ws.cell(row=11, column=col_idx)
            c.fill = fill_header_bg
            c.border = border
            c.alignment = center
            c.font = font_bold
        for pos, label in group_headers.items():
            ws[pos] = label

        for col_idx, label in enumerate(DETAIL_HEADERS, start=1):
            c = ws.cell(row=12, column=col_idx, value=label)
            c.font = font_bold
            c.alignment = center
            c.border = border
            c.fill = fill_pink if col_idx <= 8 else fill_blue_exec

        for col, width in widths.items():
            ws.column_dimensions[col].width = width

        row_idx = 15
        last_group = None
        last_feature = None
        screen_count = 0

        for tc in rows:
            group = str(tc.get("featureGroup") or "").strip()
            feature = str(tc.get("featureName") or "").strip()

            if group and group != last_group:
                for col in range(1, 28):
                    c = ws.cell(row=row_idx, column=col)
                    c.fill = fill_green
                    c.border = border
                ws.cell(row=row_idx, column=2, value=group).font = font_bold
                row_idx += 1
                last_group = group
                last_feature = None

            if feature and feature != last_feature:
                for col in range(1, 28):
                    c = ws.cell(row=row_idx, column=col)
                    c.fill = fill_light_green
                    c.border = border
                ws.cell(row=row_idx, column=2, value=f"  {feature}").font = font_bold
                row_idx += 1
                last_feature = feature

            for col_idx, field in enumerate(EXCEL_FIELDS, start=1):
                ws.cell(row=row_idx, column=col_idx, value=_cell_value(tc, field))

            for col in range(1, 28):
                c = ws.cell(row=row_idx, column=col)
                c.font = font_body
                c.fill = fill_white
                c.border = border
                c.alignment = center if col in (1,4,9,10,11,12,14,22,23,24,25,26,27) else left
            row_idx += 1
            screen_count += 1
            total += 1

        ws.freeze_panes = "A13"
        ws.auto_filter.ref = f"A12:AA{max(12, row_idx - 1)}"
        summary_rows.append((screen_name, screen_count, ws.title))

    summary = wb.create_sheet("Summary", 0)
    summary["A1"] = "ĐỐI SOÁT TEST CASE"
    summary["A1"].font = Font(name="Segoe UI", size=12, bold=True)
    summary["A3"] = "Màn hình / Chức năng"
    summary["B3"] = "Số Test Case"
    summary["C3"] = "Worksheet"
    for cell in summary[3]:
        cell.font = font_bold
        cell.fill = fill_header_bg
        cell.border = border
        cell.alignment = center

    r = 4
    for screen_name, count, sheet_name in summary_rows:
        summary.cell(r, 1, screen_name)
        summary.cell(r, 2, count)
        summary.cell(r, 3, sheet_name)
        for col in range(1, 4):
            summary.cell(r, col).border = border
            summary.cell(r, col).alignment = left if col != 2 else center
        r += 1
    summary.cell(r, 1, "TOTAL").font = font_bold
    summary.cell(r, 2, total).font = font_bold
    for col in range(1, 4):
        summary.cell(r, col).border = border
    summary.column_dimensions["A"].width = 55
    summary.column_dimensions["B"].width = 18
    summary.column_dimensions["C"].width = 35

    stream = io.BytesIO()
    wb.save(stream)
    stream.seek(0)
    return stream.getvalue()
