# -*- coding: utf-8 -*-
from __future__ import annotations

import asyncio
import io
import json
import os
import re
import tempfile
import time
import uuid
from pathlib import Path
from urllib.parse import quote
from typing import Any

from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.middleware.gzip import GZipMiddleware
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field

import db
import excel_export
import pipeline_core as core
import testcase_mapper as mapper

APP_ROOT = Path(__file__).resolve().parent.parent
FRONTEND_DIST = APP_ROOT / "frontend" / "dist"
DEFAULT_BASE_URL = "https://ws-2vuxxf5tta2cjplh.ap-southeast-1.maas.aliyuncs.com/compatible-mode/v1"
DEFAULT_MODEL = "qwen-max"

app = FastAPI(title="TestPilot AI API", version="1.11.5")
app.add_middleware(GZipMiddleware, minimum_size=1000)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173",
        "http://127.0.0.1:5173",
        "http://localhost:8000",
        "http://127.0.0.1:8000",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


@app.middleware("http")
async def use_transient_browser_session(request, call_next):
    # API calls include the browser-generated ID in a header. Download links cannot
    # carry custom headers, so they pass the same ID as a query parameter.
    session_id = request.headers.get("X-TestPilot-Session") or request.query_params.get("session_id")
    token = db.activate_session(session_id)
    try:
        return await call_next(request)
    finally:
        db.reset_session(token)


@app.middleware("http")
async def add_frontend_cache_headers(request, call_next):
    response = await call_next(request)
    path = request.url.path
    if path.startswith("/assets/"):
        response.headers["Cache-Control"] = "public, max-age=31536000, immutable"
    elif not path.startswith("/api/") and response.headers.get("content-type", "").startswith("text/html"):
        response.headers["Cache-Control"] = "no-cache"
    return response

WEB_AGENT1_CACHE: dict[str, Any] = {}
API_AGENT1_CACHE: dict[str, Any] = {}


class ConfigTestRequest(BaseModel):
    api_key: str = Field(min_length=1)
    base_url: str = DEFAULT_BASE_URL
    model: str = DEFAULT_MODEL


class ProjectPayload(BaseModel):
    id: str
    name: str
    description: str = ""
    type: str = "both"
    webFolders: list[dict] = Field(default_factory=list)
    apiFolders: list[dict] = Field(default_factory=list)
    customFolders: list[dict] = Field(default_factory=list)
    enabledScopes: list[str] = Field(default_factory=list)
    updated: str = "vừa xong"
    createdAt: float | None = None
    updatedAt: float | None = None


class TestcasePayload(BaseModel):
    id: str
    name: str
    preCondition: str = ""
    importance: str = ""
    steps: list[str] = Field(default_factory=list)
    testData: str = ""
    expectedResult: str
    actualResult: str = ""
    run1: str = ""
    run2: str = ""
    run3: str = ""
    currentResult: str = ""
    note: str = ""
    errorCode: str = ""
    qcWriter: str = ""
    sprintWriter: str = ""
    qcExecutor: str = ""
    sprintExecutor: str = ""
    reviewer: str = ""
    reviewDate: str = ""
    reviewContent: str = ""
    needAuto: str = ""
    automated: str = ""
    smoke: str = ""
    regression: str = ""
    outdated: str = ""
    outdatedDate: str = ""
    type: str = "Chức năng"
    sourceRuleId: str = ""
    sourceRequirement: str = ""
    featureGroup: str = "FUNCTION"
    featureName: str = "Testcase thủ công"
    category: str = "ACTION"
    screen: str = ""


class CreateTestcaseRequest(BaseModel):
    project_id: str
    folder_id: str
    scope: str = "web"
    run_id: str | None = None
    testcase: TestcasePayload


class UpdateTestcaseRequest(BaseModel):
    testcase: TestcasePayload


class MemoryUpload:
    def __init__(self, name: str, data: bytes):
        self.name = name
        self._data = data

    def getvalue(self):
        return self._data


def _extract_docx(data: bytes) -> str:
    """Extract DOCX in document order and preserve heading/table structure.

    The previous implementation emitted all paragraphs first and all tables afterward. In BA
    documents that destroys ownership (screen heading -> field table -> logic table), which is
    exactly the structure the Web discovery stage needs.
    """
    try:
        from docx import Document
        from docx.text.paragraph import Paragraph
        from docx.table import Table
    except ImportError as exc:
        raise HTTPException(500, "Thiếu python-docx. Chạy pip install -r requirements.txt") from exc

    doc = Document(io.BytesIO(data))
    parts: list[str] = []

    def clean(text: str) -> str:
        return str(text or "").replace("\r\n", "\n").replace("\r", "\n").strip()

    def paragraph_text(paragraph: Paragraph) -> str:
        # Preserve strikethrough because BA revisions often keep old requirements in the file.
        # The discovery prompt treats [STRIKETHROUGH] content as inactive.
        if not paragraph.runs:
            return clean(paragraph.text)
        pieces: list[str] = []
        for run in paragraph.runs:
            text = str(run.text or "")
            if not text:
                continue
            if bool(run.font.strike) or bool(run.font.double_strike):
                pieces.append(f"[STRIKETHROUGH]{text}[/STRIKETHROUGH]")
            else:
                pieces.append(text)
        return clean("".join(pieces))

    def cell_text(cell) -> str:
        values = [paragraph_text(p) for p in cell.paragraphs]
        return clean("\n".join(x for x in values if x))

    def emit_table(table: Table) -> None:
        for row_idx, row in enumerate(table.rows, start=1):
            cells = [cell_text(c).replace("|", "\\|") for c in row.cells]
            if not any(cells):
                continue
            while cells and cells[-1] == "":
                cells.pop()
            # Markdown-style row keeps each BA table row atomic for semantic chunking.
            parts.append(f"| TROW {row_idx} | " + " | ".join(cells) + " |")

    # python-docx >=1.1 exposes iter_inner_content and preserves Paragraph/Table order.
    iterator = getattr(doc, "iter_inner_content", None)
    if callable(iterator):
        blocks = iterator()
    else:
        # Compatibility fallback using document body XML order.
        from docx.oxml.text.paragraph import CT_P
        from docx.oxml.table import CT_Tbl
        blocks = []
        for child in doc.element.body.iterchildren():
            if isinstance(child, CT_P):
                blocks.append(Paragraph(child, doc))
            elif isinstance(child, CT_Tbl):
                blocks.append(Table(child, doc))

    for block in blocks:
        if isinstance(block, Paragraph):
            text = paragraph_text(block)
            if not text:
                continue
            style_name = clean(getattr(getattr(block, "style", None), "name", ""))
            m = re.match(r"(?i)^heading\s*(\d+)$", style_name)
            if m:
                level = max(1, min(6, int(m.group(1))))
                parts.append("#" * level + " " + text)
            else:
                parts.append(text)
        elif isinstance(block, Table):
            emit_table(block)
    return "\n".join(parts)


def _extract_doc(data: bytes) -> str:
    with tempfile.NamedTemporaryFile(suffix=".doc", delete=False) as tmp:
        tmp.write(data)
        tmp_path = tmp.name
    try:
        import subprocess
        result = subprocess.run(
            ["antiword", tmp_path],
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
            timeout=60,
        )
        if result.returncode != 0:
            raise HTTPException(400, "Không đọc được file .doc. Hãy chuyển sang DOCX hoặc PDF.")
        return result.stdout
    except FileNotFoundError as exc:
        raise HTTPException(400, "Máy chưa có antiword để đọc .doc. Hãy dùng DOCX hoặc PDF.") from exc
    finally:
        try:
            os.unlink(tmp_path)
        except OSError:
            pass


def _extract_xlsx(data: bytes) -> str:
    """Extract BA-style Excel with sheet/row boundaries and table semantics."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise HTTPException(500, "Thiếu openpyxl. Chạy pip install -r requirements.txt") from exc
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except Exception as exc:
        raise HTTPException(400, "Không đọc được file XLSX.") from exc

    heading_re = re.compile(
        r"(?i)^(?:\d+(?:\.\d+)*[.)]?\s+)?(?:màn hình|screen|mô tả màn hình|logic|luồng|quy tắc nghiệp vụ|mockup|đặc tả)\b"
    )
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"=== SHEET: {ws.title} ===")
        for row_number, row in enumerate(ws.iter_rows(), start=1):
            raw_values: list[str] = []
            for cell in row:
                value = cell.value
                text = "" if value is None else str(value).replace("\r\n", "\n").replace("\r", "\n").strip()
                raw_values.append(text)
            if not any(raw_values):
                continue
            while raw_values and raw_values[-1] == "":
                raw_values.pop()
            non_empty = [x for x in raw_values if x]
            if len(non_empty) == 1 and heading_re.match(non_empty[0]):
                parts.append(f"# {non_empty[0]} [ROW {row_number}]")
                continue
            cells = [x.replace("|", "\\|") for x in raw_values]
            parts.append(f"| ROW {row_number} | " + " | ".join(cells) + " |")
    return "\n".join(parts)


def extract_upload_text(filename: str, data: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".docx"):
        return _extract_docx(data)
    if lower.endswith(".doc"):
        return _extract_doc(data)
    if lower.endswith(".xlsx"):
        return _extract_xlsx(data)
    return core.extract_text_from_file(MemoryUpload(filename, data))


async def _read_upload(upload: UploadFile) -> tuple[str, bytes]:
    data = await upload.read()
    if not data:
        raise HTTPException(400, f"File {upload.filename or 'không tên'} rỗng")
    return upload.filename or "document.txt", data


def _get_run_or_404(run_id: str) -> dict:
    run = db.get_run(run_id)
    if not run:
        raise HTTPException(404, "Không tìm thấy phiên phân tích.")
    return run


def _source_excel_name(run: dict | None, fallback: str = "TestCases") -> str:
    names = (run or {}).get("source_names") or []
    base = Path(str(names[0])).stem if names else str(fallback or "TestCases")
    base = re.sub(r'[\\/:*?"<>|]+', '_', base).strip(' ._') or "TestCases"
    return f"{base}.xlsx"


def _attachment_header(filename: str) -> dict[str, str]:
    ascii_name = re.sub(r'[^A-Za-z0-9._-]+', '_', filename) or "TestCases.xlsx"
    return {
        "Content-Disposition": f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
    }


def _ensure_manual_run(project_id: str, folder_id: str, scope: str) -> str:
    # Manual records still need a DB run FK. Create a stable synthetic run for the folder.
    run_id = f"manual::{scope}::{project_id}::{folder_id}"
    if not db.get_run(run_id):
        db.save_run(
            run_id=run_id,
            kind=scope,
            project_id=project_id,
            folder_id=folder_id,
            source_names=[],
            matrix={},
            agent1_summary={"manual": True},
        )
    return run_id


@app.get("/api/health")
def health():
    return {
        "ok": True,
        "version": "1.11.5",
        "web_pipeline": core.WEB_TEST_DESIGN_VERSION,
        "ai_stages": 2,
        "web_architecture": "DISCOVERY->CANONICAL_INVENTORY->QA_RULES",
        "persistence": "session-scoped in-memory sqlite",
        "analysis_mode": "async_job_polling",
        "database": str(db.DB_PATH),
        "models": ["qwen-max", "qwen3.8-max", "qwen-plus"],
    }


@app.post("/api/session/end")
def end_browser_session(session_id: str = ""):
    db.discard_session(session_id)
    return {"ok": True}


@app.get("/api/projects")
def list_projects():
    return {"projects": db.list_projects()}


@app.get("/api/projects/{project_id}/metadata")
def get_project_metadata(project_id: str):
    project = db.get_project(project_id)
    if not project:
        raise HTTPException(404, "Không tìm thấy dự án")
    return {"project": project}


@app.put("/api/projects/{project_id}/metadata")
def save_project_metadata(project_id: str, payload: ProjectPayload):
    data = payload.model_dump()
    data["id"] = project_id
    data["updatedAt"] = data.get("updatedAt") or time.time()
    project = db.upsert_project(data)
    return {"ok": True, "project": project}


@app.post("/api/projects/sync")
def sync_projects(payload: dict):
    projects = payload.get("projects") if isinstance(payload, dict) else []
    if not isinstance(projects, list):
        raise HTTPException(400, "projects phải là array")
    saved = []
    for project in projects:
        if isinstance(project, dict) and (project.get("id") or "").strip():
            saved.append(db.upsert_project(project))
    return {"ok": True, "projects": db.list_projects(), "saved": len(saved)}


@app.post("/api/config/test")
async def test_config(payload: ConfigTestRequest):
    base_url = payload.base_url.strip().rstrip("/")
    model = payload.model.strip()

    def _call(max_tokens: int):
        return core.call_qwen_max_agent_detailed(
            content="ping",
            api_key=payload.api_key.strip(),
            base_url=base_url,
            model=model,
            prompt_template="Đây là health check. Hãy trả lời đúng một từ: OK. {content}",
            max_tokens=max_tokens,
            agent_name="CONFIG_TEST",
        )

    result = await asyncio.to_thread(_call, 1024)
    if (not result.ok) and result.finish_reason == "length":
        result = await asyncio.to_thread(_call, 4096)
    if not result.ok:
        raise HTTPException(
            400,
            detail={
                "message": result.error or "Không kết nối được model",
                "model": model,
                "base_url": base_url,
                "finish_reason": result.finish_reason,
                "elapsed": result.elapsed,
            },
        )
    return {
        "ok": True,
        "model": model,
        "base_url": base_url,
        "result": result.text.strip(),
        "elapsed": result.elapsed,
        "finish_reason": result.finish_reason,
    }


ANALYSIS_TASKS: set[asyncio.Task] = set()


def _launch_analysis_job(coro) -> None:
    task = asyncio.create_task(coro)
    ANALYSIS_TASKS.add(task)
    task.add_done_callback(ANALYSIS_TASKS.discard)


def _job_failure_payload(exc: Exception) -> dict:
    if isinstance(exc, HTTPException):
        detail = exc.detail
        if isinstance(detail, dict):
            return detail
        return {"message": str(detail), "status_code": exc.status_code}
    return {"message": str(exc) or exc.__class__.__name__, "type": exc.__class__.__name__}


async def _process_web_analysis_job(
    *,
    job_id: str,
    uploads: list[tuple[str, bytes]],
    api_key: str,
    base_url: str,
    model: str,
    project_id: str,
    folder_id: str,
) -> None:
    try:
        db.update_analysis_job(
            job_id,
            status="processing",
            stage="extracting",
            message="Đang đọc và trích xuất nội dung tài liệu...",
            progress=5,
        )
        sources: list[str] = []
        names: list[str] = []
        total = max(1, len(uploads))
        for index, (name, data) in enumerate(uploads, start=1):
            text = await asyncio.to_thread(extract_upload_text, name, data)
            if text.strip():
                sources.append(f"=== SOURCE DOCUMENT: {name} ===\n{text}")
                names.append(name)
            db.update_analysis_job(
                job_id,
                message=f"Đã đọc {index}/{total} tài liệu.",
                progress=min(15, 5 + int(index / total * 10)),
            )

        if not sources:
            raise HTTPException(400, "Không trích xuất được nội dung từ tài liệu")

        raw_text = "\n\n".join(sources)
        base_filename = Path(names[0]).stem if len(names) == 1 else "multi_document_web"
        db.update_analysis_job(
            job_id,
            stage="analyzing",
            message="Bước 1/2: AI đang xác định màn hình, field/control và logic BA...",
            progress=20,
        )

        def _progress(current, total, message):
            try:
                ratio = float(current) / max(1.0, float(total)) if total else 0.0
            except Exception:
                ratio = 0.0
            db.update_analysis_job(
                job_id,
                status="processing",
                stage="analyzing",
                message=message,
                progress=max(20, min(80, 20 + int(60 * ratio))),
            )

        def _call():
            return core.run_agent1_document_pipeline(
                raw_text=raw_text,
                base_filename=base_filename,
                api_key=api_key,
                base_url=base_url,
                model=model,
                prompt_template=core.PROMPT_AGENT1_EXTRACT_RULE_MATRIX,
                cache=WEB_AGENT1_CACHE,
                progress_callback=_progress,
            )

        ok, matrix, summary = await asyncio.to_thread(_call)
        if not ok or matrix is None:
            stage = str((summary or {}).get("stage") or "unknown")
            stage_messages = {
                "web_discovery": "Không hoàn tất bước đọc BA và xác định màn hình/field/logic.",
                "web_inventory": "Không lập được inventory màn hình hợp lệ từ tài liệu BA.",
                "web_qa_rules": "Đã đọc được cấu trúc BA nhưng không hoàn tất bước áp dụng QA Rule cho một màn hình.",
                "final_schema": "Rule Matrix Web cuối cùng không đạt schema sau khi chuẩn hóa.",
            }
            raise HTTPException(422, detail={
                "message": stage_messages.get(stage, "AI phân tích Web thất bại."),
                "stage": stage,
                "summary": summary,
            })

        db.update_analysis_job(
            job_id,
            stage="mapping",
            message="Đang chuyển Rule Matrix thành testcase...",
            progress=82,
        )
        testcases = mapper.map_web_matrix_to_testcases(matrix)
        expected_rules = (summary.get("merge") or {}).get("final_rules")
        if expected_rules is not None and len(testcases) != int(expected_rules):
            raise HTTPException(500, detail={
                "message": "Mapping testcase không bảo toàn số Rule Matrix",
                "final_rules": expected_rules,
                "mapped_testcases": len(testcases),
            })

        db.update_analysis_job(
            job_id,
            stage="saving",
            message="Đang lưu testcase vào workspace...",
            progress=92,
        )
        run_id = uuid.uuid4().hex
        db.save_run(
            run_id=run_id,
            kind="web",
            project_id=project_id,
            folder_id=folder_id,
            source_names=names,
            matrix=matrix,
            agent1_summary=summary,
        )
        saved = db.replace_run_testcases(
            run_id=run_id,
            project_id=project_id,
            folder_id=folder_id,
            scope="web",
            testcases=testcases,
        )
        invalid = sum(1 for tc in saved if not tc.get("validation", {}).get("valid", False))
        db.update_analysis_job(
            job_id,
            status="completed",
            stage="completed",
            message=f"Hoàn tất {len(saved)} testcase" + (f", {invalid} case cần kiểm tra." if invalid else "."),
            progress=100,
            run_id=run_id,
            error_payload={},
        )
    except Exception as exc:
        payload = _job_failure_payload(exc)
        db.update_analysis_job(
            job_id,
            status="failed",
            stage="failed",
            message=payload.get("message") or "Phân tích tài liệu thất bại.",
            error_payload=payload,
        )


async def _process_api_analysis_job(
    *,
    job_id: str,
    design_upload: tuple[str, bytes],
    ba_upload: tuple[str, bytes],
    api_key: str,
    base_url: str,
    model: str,
    project_id: str,
    folder_id: str,
) -> None:
    try:
        design_name, design_bytes = design_upload
        ba_name, ba_bytes = ba_upload
        db.update_analysis_job(
            job_id,
            status="processing",
            stage="extracting",
            message="Đang đọc API Design và tài liệu BA...",
            progress=5,
        )
        design_text, ba_text = await asyncio.gather(
            asyncio.to_thread(extract_upload_text, design_name, design_bytes),
            asyncio.to_thread(extract_upload_text, ba_name, ba_bytes),
        )
        if not design_text.strip() or not ba_text.strip():
            raise HTTPException(400, "Không trích xuất được đầy đủ nội dung tài liệu API")

        base_filename = f"{Path(design_name).stem}_{Path(ba_name).stem}"
        db.update_analysis_job(
            job_id,
            stage="analyzing",
            message="Đang nhận diện API mục tiêu từ API Design...",
            progress=20,
        )

        def _progress(current, total, message):
            try:
                progress = int((float(current) / max(1.0, float(total))) * 100) if total else int(current)
            except Exception:
                progress = int(current or 20)
            db.update_analysis_job(
                job_id,
                status="processing",
                stage="analyzing",
                message=message,
                progress=max(10, min(90, progress)),
            )

        def _call():
            return core.run_api_agent1_document_pipeline(
                design_text=design_text,
                ba_text=ba_text,
                base_filename=base_filename,
                api_key=api_key,
                base_url=base_url,
                model=model,
                prompt_template=core.PROMPT_API_AGENT1_RULE_MATRIX,
                cache=API_AGENT1_CACHE,
                progress_callback=_progress,
            )

        ok, matrix, summary = await asyncio.to_thread(_call)
        if not ok or matrix is None:
            diagnostics = (summary or {}).get("diagnostics") or []
            last_diag = diagnostics[-1] if diagnostics else {}
            reason = str(last_diag.get("reason") or (summary or {}).get("stage") or "UNKNOWN")
            reason_messages = {
                "JSON_PARSE_FAIL": "AI trả JSON không hợp lệ sau các bước phục hồi tự động.",
                "SCHEMA_FAIL": "AI trả dữ liệu không đúng cấu trúc Rule Matrix API.",
                "MAX_TOKENS": "Kết quả AI vượt giới hạn output sau khi đã tự chia nhỏ.",
                "API_FAILED": "Kết nối tới model AI thất bại.",
                "ba_deep": "Không hoàn tất được bước phân tích phần BA liên quan tới API mục tiêu.",
                "final_schema": "Rule Matrix còn dữ liệu không hợp lệ sau các bước chuẩn hóa an toàn.",
            }
            raise HTTPException(422, detail={
                "message": reason_messages.get(reason, "AI phân tích API thất bại."),
                "reason": reason,
                "summary": summary,
            })

        db.update_analysis_job(
            job_id,
            stage="mapping",
            message="Đang chuyển Rule Matrix API thành testcase...",
            progress=82,
        )
        testcases = mapper.map_api_matrix_to_testcases(matrix)
        expected_rules = (summary.get("merge") or {}).get("final_rules")
        if expected_rules is not None and len(testcases) != int(expected_rules):
            raise HTTPException(500, detail={
                "message": "Mapping testcase API không bảo toàn số Rule Matrix",
                "final_rules": expected_rules,
                "mapped_testcases": len(testcases),
            })

        db.update_analysis_job(
            job_id,
            stage="saving",
            message="Đang lưu testcase API vào workspace...",
            progress=92,
        )
        run_id = uuid.uuid4().hex
        db.save_run(
            run_id=run_id,
            kind="api",
            project_id=project_id,
            folder_id=folder_id,
            source_names=[design_name, ba_name],
            matrix=matrix,
            agent1_summary=summary,
        )
        saved = db.replace_run_testcases(
            run_id=run_id,
            project_id=project_id,
            folder_id=folder_id,
            scope="api",
            testcases=testcases,
        )
        invalid = sum(1 for tc in saved if not tc.get("validation", {}).get("valid", False))
        db.update_analysis_job(
            job_id,
            status="completed",
            stage="completed",
            message=f"Hoàn tất {len(saved)} testcase" + (f", {invalid} case cần kiểm tra." if invalid else "."),
            progress=100,
            run_id=run_id,
            error_payload={},
        )
    except Exception as exc:
        payload = _job_failure_payload(exc)
        db.update_analysis_job(
            job_id,
            status="failed",
            stage="failed",
            message=payload.get("message") or "Phân tích tài liệu API thất bại.",
            error_payload=payload,
        )


@app.post("/api/web/analyze", status_code=202)
async def web_analyze(
    files: list[UploadFile] = File(...),
    api_key: str = Form(...),
    base_url: str = Form(DEFAULT_BASE_URL),
    model: str = Form(DEFAULT_MODEL),
    project_id: str = Form("default-project"),
    folder_id: str = Form("default-folder"),
):
    if not files:
        raise HTTPException(400, "Chưa có tài liệu")
    uploads = [await _read_upload(upload) for upload in files]
    job_id = uuid.uuid4().hex
    db.create_analysis_job(
        job_id=job_id,
        kind="web",
        project_id=project_id,
        folder_id=folder_id,
        source_names=[name for name, _ in uploads],
    )
    _launch_analysis_job(_process_web_analysis_job(
        job_id=job_id,
        uploads=uploads,
        api_key=api_key,
        base_url=base_url,
        model=model,
        project_id=project_id,
        folder_id=folder_id,
    ))
    return {
        "ok": True,
        "job_id": job_id,
        "status": "queued",
        "message": "Đã tiếp nhận tài liệu. AI sẽ xử lý ở background.",
    }


@app.post("/api/api/analyze", status_code=202)
async def api_analyze(
    design_file: UploadFile = File(...),
    ba_file: UploadFile = File(...),
    api_key: str = Form(...),
    base_url: str = Form(DEFAULT_BASE_URL),
    model: str = Form(DEFAULT_MODEL),
    project_id: str = Form("default-project"),
    folder_id: str = Form("default-folder"),
):
    design_upload = await _read_upload(design_file)
    ba_upload = await _read_upload(ba_file)
    job_id = uuid.uuid4().hex
    db.create_analysis_job(
        job_id=job_id,
        kind="api",
        project_id=project_id,
        folder_id=folder_id,
        source_names=[design_upload[0], ba_upload[0]],
    )
    _launch_analysis_job(_process_api_analysis_job(
        job_id=job_id,
        design_upload=design_upload,
        ba_upload=ba_upload,
        api_key=api_key,
        base_url=base_url,
        model=model,
        project_id=project_id,
        folder_id=folder_id,
    ))
    return {
        "ok": True,
        "job_id": job_id,
        "status": "queued",
        "message": "Đã tiếp nhận tài liệu API. AI sẽ xử lý ở background.",
    }


@app.get("/api/jobs/{job_id}")
def get_analysis_job(job_id: str):
    job = db.get_analysis_job(job_id)
    if not job:
        raise HTTPException(404, "Không tìm thấy tác vụ phân tích.")
    return job




def _ordered_testcases(testcases: list[dict], scope: str | None = None) -> list[dict]:
    return mapper.order_and_renumber_testcases(testcases, scope)


@app.get("/api/runs/{run_id}")
def get_run(run_id: str):
    run = _get_run_or_404(run_id)
    return {
        "ok": True,
        "run_id": run_id,
        "kind": run.get("kind"),
        "project_id": run.get("project_id"),
        "folder_id": run.get("folder_id"),
        # Matrix is intentionally not returned to normal UI.
        "testcases": _ordered_testcases(db.list_run_testcases(run_id), run.get("kind")),
        "agent1_summary": run.get("agent1_summary"),
        "source_names": run.get("source_names", []),
    }


@app.get("/api/runs/{run_id}/diagnostics")
def get_diagnostics(run_id: str):
    run = _get_run_or_404(run_id)
    return {"analysis": run.get("agent1_summary")}


@app.get("/api/projects/{project_id}/folders/{folder_id}/testcases")
def get_folder_testcases(project_id: str, folder_id: str, scope: str = "web"):
    return {
        "ok": True,
        "project_id": project_id,
        "folder_id": folder_id,
        "scope": scope,
        "testcases": _ordered_testcases(db.list_folder_testcases(project_id, folder_id, scope), scope),
    }


@app.get("/api/projects/{project_id}/testcase-tree")
def get_project_testcase_tree(project_id: str):
    rows = db.list_project_tree_summary(project_id)
    scopes = {
        "web": {"key": "web", "label": "Web App", "screens": []},
        "api": {"key": "api", "label": "API", "screens": []},
    }
    screen_maps: dict[str, dict[str, dict]] = {"web": {}, "api": {}}
    total_testcases = 0

    for row in rows:
        scope = row.get("scope") if row.get("scope") in scopes else "web"
        screen_name = str(row.get("screen") or "Chưa xác định màn hình").strip()
        count = int(row.get("count") or 0)
        total_testcases += count
        smap = screen_maps[scope]
        if screen_name not in smap:
            node = {
                "screen": screen_name,
                "count": 0,
                "folder_ids": [],
                "folderCounts": {},
            }
            smap[screen_name] = node
            scopes[scope]["screens"].append(node)
        node = smap[screen_name]
        node["count"] += count
        folder_id = row.get("folder_id")
        if folder_id:
            if folder_id not in node["folder_ids"]:
                node["folder_ids"].append(folder_id)
            node["folderCounts"][folder_id] = node["folderCounts"].get(folder_id, 0) + count

    return {
        "ok": True,
        "project_id": project_id,
        "total_testcases": total_testcases,
        "total_screens": sum(len(item["screens"]) for item in scopes.values()),
        "scopes": [scopes["web"], scopes["api"]],
    }


@app.get("/api/projects/{project_id}/screen-testcases")
def get_screen_testcases(project_id: str, scope: str = "web", screen: str = ""):
    target = str(screen or "").strip()
    testcases = db.list_project_testcases(project_id, scope)
    if target:
        testcases = [tc for tc in testcases if str(tc.get("screen") or "").strip() == target]
    testcases = _ordered_testcases(testcases, scope)
    return {
        "ok": True,
        "project_id": project_id,
        "scope": scope,
        "screen": target,
        "testcases": testcases,
    }


@app.get("/api/projects/{project_id}/screen-excel")
def download_screen_excel(project_id: str, scope: str = "web", screen: str = ""):
    target = str(screen or "").strip()
    testcases = [
        tc for tc in db.list_project_testcases(project_id, scope)
        if str(tc.get("screen") or "").strip() == target
    ]
    if not testcases:
        raise HTTPException(409, "Màn hình chưa có testcase")
    testcases = _ordered_testcases(testcases, scope)
    run = db.get_run(testcases[0].get("runId")) if testcases and testcases[0].get("runId") else None
    run = run or db.get_latest_run_for_project_scope(project_id, scope)
    filename = _source_excel_name(run, target or "TestCases")
    data = excel_export.build_excel(testcases, workbook_title=Path(filename).stem)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_attachment_header(filename),
    )


@app.get("/api/projects/{project_id}/screen-xmind")
def download_screen_xmind(project_id: str, scope: str = "web", screen: str = ""):
    target = str(screen or "").strip()
    testcases = [
        tc for tc in db.list_project_testcases(project_id, scope)
        if str(tc.get("screen") or "").strip() == target
    ]
    if not testcases:
        raise HTTPException(409, "Màn hình chưa có testcase")
    testcases = _ordered_testcases(testcases, scope)
    tree = mapper.build_tree_text(testcases)
    with tempfile.NamedTemporaryFile(suffix=".xmind", delete=False) as tmp:
        path = tmp.name
    ok, err = core.create_xmind_from_text(
        tree,
        path,
        root_title=target or ("Bộ Test Case API" if scope == "api" else "Bộ Test Case Web App"),
    )
    if not ok:
        try:
            os.unlink(path)
        except OSError:
            pass
        raise HTTPException(500, err or "Không tạo được XMind")
    safe = re.sub(r'[\\/:*?"<>|]+', '_', target).strip(' ._') or "TestCases"
    return FileResponse(path, filename=f"{safe}.xmind", media_type="application/octet-stream")


@app.post("/api/testcases")
def create_testcase(payload: CreateTestcaseRequest):
    tc = mapper.validate_and_attach(payload.testcase.model_dump())
    run_id = payload.run_id or _ensure_manual_run(payload.project_id, payload.folder_id, payload.scope)
    item = db.create_testcase(
        project_id=payload.project_id,
        folder_id=payload.folder_id,
        scope=payload.scope,
        run_id=run_id,
        payload=tc,
    )
    return {"ok": True, "testcase": item}


@app.put("/api/testcases/{record_id}")
def update_testcase(record_id: int, payload: UpdateTestcaseRequest):
    tc = mapper.validate_and_attach(payload.testcase.model_dump())
    item = db.update_testcase(record_id, tc)
    if not item:
        raise HTTPException(404, "Không tìm thấy testcase")
    return {"ok": True, "testcase": item}


@app.delete("/api/testcases/{record_id}")
def delete_testcase(record_id: int):
    if not db.delete_testcase(record_id):
        raise HTTPException(404, "Không tìm thấy testcase")
    return {"ok": True}


@app.get("/api/projects/{project_id}/folders/{folder_id}/excel")
def download_folder_excel(project_id: str, folder_id: str, scope: str = "web"):
    testcases = db.list_folder_testcases(project_id, folder_id, scope)
    if not testcases:
        raise HTTPException(409, "Thư mục chưa có testcase")
    testcases = _ordered_testcases(testcases, scope)
    run = db.get_latest_run_for_folder(project_id, folder_id, scope)
    filename = _source_excel_name(run, folder_id)
    data = excel_export.build_excel(testcases, workbook_title=Path(filename).stem)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_attachment_header(filename),
    )


@app.get("/api/projects/{project_id}/folders/{folder_id}/xmind")
def download_folder_xmind(project_id: str, folder_id: str, scope: str = "web"):
    testcases = db.list_folder_testcases(project_id, folder_id, scope)
    if not testcases:
        raise HTTPException(409, "Thư mục chưa có testcase")
    testcases = _ordered_testcases(testcases, scope)
    tree = mapper.build_tree_text(testcases)
    with tempfile.NamedTemporaryFile(suffix=".xmind", delete=False) as tmp:
        path = tmp.name
    ok, err = core.create_xmind_from_text(
        tree,
        path,
        root_title="Bộ Test Case API" if scope == "api" else "Bộ Test Case UI/UX & Web App",
    )
    if not ok:
        try:
            os.unlink(path)
        except OSError:
            pass
        raise HTTPException(500, err or "Không tạo được XMind")
    return FileResponse(path, filename="TestCases.xmind", media_type="application/octet-stream")


@app.get("/api/runs/{run_id}/excel")
def download_run_excel(run_id: str):
    run = _get_run_or_404(run_id)
    testcases = db.list_run_testcases(run_id)
    if not testcases:
        raise HTTPException(409, "Run chưa có testcase")
    testcases = _ordered_testcases(testcases, run.get("kind"))
    filename = _source_excel_name(run, "TestCases")
    data = excel_export.build_excel(testcases, workbook_title=Path(filename).stem)
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers=_attachment_header(filename),
    )


@app.get("/api/runs/{run_id}/xmind")
def download_run_xmind(run_id: str):
    run = _get_run_or_404(run_id)
    testcases = db.list_run_testcases(run_id)
    if not testcases:
        raise HTTPException(409, "Run chưa có testcase")
    testcases = _ordered_testcases(testcases, run.get("kind"))
    tree = mapper.build_tree_text(testcases)
    with tempfile.NamedTemporaryFile(suffix=".xmind", delete=False) as tmp:
        path = tmp.name
    ok, err = core.create_xmind_from_text(
        tree,
        path,
        root_title="Bộ Test Case API" if run.get("kind") == "api" else "Bộ Test Case UI/UX & Web App",
    )
    if not ok:
        try:
            os.unlink(path)
        except OSError:
            pass
        raise HTTPException(500, err or "Không tạo được XMind")
    return FileResponse(path, filename="TestCases.xmind", media_type="application/octet-stream")


@app.get("/api/runs/{run_id}/rule-matrix")
def download_rule_matrix_debug(run_id: str):
    """Debug-only endpoint. The normal UI intentionally does not expose the internal matrix."""
    run = _get_run_or_404(run_id)
    payload = json.dumps(run.get("matrix", {}), ensure_ascii=False, indent=2).encode("utf-8")
    return Response(
        content=payload,
        media_type="application/json; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="internal_rule_matrix.json"'},
    )


if FRONTEND_DIST.exists():
    assets_dir = FRONTEND_DIST / "assets"
    if assets_dir.exists():
        app.mount("/assets", StaticFiles(directory=str(assets_dir)), name="assets")

    @app.get("/{full_path:path}", include_in_schema=False)
    def spa(full_path: str):
        candidate = FRONTEND_DIST / full_path
        if full_path and candidate.exists() and candidate.is_file():
            return FileResponse(candidate)
        return FileResponse(FRONTEND_DIST / "index.html")
