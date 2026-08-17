# -*- coding: utf-8 -*-
import json
import io
import os
import re
import sys
import tempfile
import time
import traceback
import zipfile 
from pypdf import PdfReader
from datetime import datetime
import streamlit as st
from openai import OpenAI
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==========================================
# 1. CẤU HÌNH TRANG STREAMLIT & SIDEBAR
# ==========================================
st.set_page_config(
    page_title="Chuyển đổi Tài liệu sang XMind Test Case (Multi-Agent)",
    page_icon="🧠",
    layout="wide"
)

st.title("🧠 Công cụ Tự động Tạo Test Case XMind từ Tài liệu (Kiến trúc Multi-Agent)")
st.caption("Giải pháp Mã Nguồn Mở dành riêng cho QA Lead: Phân tích tài liệu/Figma, Review JSON Scope Plan và viết bộ Test Case chi tiết ra sơ đồ tư duy (.xmind)")

with st.sidebar:
    st.header("⚙️ Cấu hình Qwen API (DashScope)")
    api_key = st.text_input("Nhập API Key:", type="password", key="sidebar_api_key")
    
    base_url = st.text_input(
        "Base URL Endpoint:", 
        value="https://token-plan.ap-southeast-1.maas.aliyuncs.com/compatible-mode/v1",
        key="sidebar_base_url"
    )
    
    model_name = st.selectbox(
        "Chọn Model cho Agent 2:", 
        ["qwen-max", "qwen3.8-max", "qwen-plus"], 
        index=0,
        key="sidebar_model_name"
    )

    st.divider()
    st.markdown("### 🔄 Trạng thái Tiến trình")
    if "step_ui" not in st.session_state:
        st.session_state.step_ui = 1
    if "step_api" not in st.session_state:
        st.session_state.step_api = 1

    if st.button("🔴 Reset/Chạy lại từ đầu", key="sidebar_reset_btn"):
        st.session_state.clear()
        st.rerun()

# ==========================================
# 2. KHAI BÁO CÁC HÀM XỬ LÝ (HELPER FUNCTIONS)
# ==========================================
def log_info(msg: str):
    timestamp = datetime.now().strftime("%H:%M:%S")
    log_msg = f"[LOG {timestamp}] ℹ️ {msg}"
    print(log_msg)
    sys.stdout.flush()

def log_error(msg: str, exc: Exception = None):
    timestamp = datetime.now().strftime("%H:%M:%S")
    err_msg = f"[ERROR {timestamp}] ❌ {msg}"
    if exc:
        err_msg += f"\nDetails: {str(exc)}\n{traceback.format_exc()}"
    print(err_msg)
    sys.stdout.flush()
    return err_msg

def extract_text_from_file(uploaded_file) -> str:
    if uploaded_file is None:
        return ""
    file_bytes = uploaded_file.read()
    filename = uploaded_file.name.lower()
    
    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        full_pdf_text = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                full_pdf_text.append(f"--- TRANG {idx+1} ---\n{page_text}")
        return "\n\n".join(full_pdf_text)
    else:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1")

def _log_response_summary(agent_name: str, elapsed: float, result: str, chunk_count: int, finish_reason: str | None, reasoning_chunk_count: int = 0):
    """Log chẩn đoán đầy đủ cho request Qwen mà không dump toàn bộ output ra console."""
    log_info(
        f"[{agent_name}] ✅ Request kết thúc | "
        f"Time={elapsed:.2f}s | Chunks={chunk_count} | "
        f"ReasoningChunks={reasoning_chunk_count} | "
        f"OutputChars={len(result):,} | "
        f"FinishReason={finish_reason or 'UNKNOWN'}"
    )


def call_qwen_agent(prompt: str, api_key: str, base_url: str, model: str, max_tokens: int = 16384) -> tuple[bool, str]:
    """Generic Qwen caller cho API/flow cũ. Có log finish_reason + kích thước output."""
    client = OpenAI(
        api_key=api_key,
        base_url=base_url,
        timeout=3600.0
    )
    start_time = time.time()
    log_info(
        f"[QWEN_AGENT] 🚀 Request | Model={model} | MaxTokens={max_tokens} | "
        f"PromptChars={len(prompt):,}"
    )

    try:
        response = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.1,
            max_tokens=max_tokens,
            stream=True
        )

        full_response = []
        chunk_count = 0
        reasoning_chunk_count = 0
        finish_reason = None
        last_log_time = time.time()

        for chunk in response:
            if not chunk.choices:
                continue

            choice = chunk.choices[0]
            delta = choice.delta

            if getattr(delta, "reasoning_content", None):
                reasoning_chunk_count += 1

            if delta and delta.content:
                full_response.append(delta.content)
                chunk_count += 1

            if choice.finish_reason:
                finish_reason = choice.finish_reason

            now = time.time()
            if now - last_log_time >= 10:
                elapsed = int(now - start_time)
                log_info(
                    f"[QWEN_AGENT] ⏳ Streaming | Elapsed={elapsed}s | "
                    f"ContentChunks={chunk_count} | ReasoningChunks={reasoning_chunk_count}"
                )
                last_log_time = now

        elapsed_total = time.time() - start_time
        result = "".join(full_response)
        _log_response_summary("QWEN_AGENT", elapsed_total, result, chunk_count, finish_reason, reasoning_chunk_count)
        return True, result

    except Exception as e:
        elapsed_total = time.time() - start_time
        err_details = log_error(
            f"[QWEN_AGENT] ❌ Request failed after {elapsed_total:.2f}s | "
            f"Model={model} | MaxTokens={max_tokens}",
            e
        )
        return False, err_details


def call_qwen_max_agent(content: str, api_key: str, base_url: str, model: str, prompt_template: str, max_tokens: int = 16384) -> tuple[bool, str]:
    """Qwen caller chính cho Web/App UI. Không parse JSON ở tầng transport."""
    client = OpenAI(
        api_key=api_key,
        base_url=base_url,
        timeout=3600.0
    )

    # Chỉ thay placeholder {content}; không dùng str.format() vì prompt có rất nhiều { } JSON.
    full_prompt = prompt_template.replace("{content}", content)
    start_time = time.time()
    log_info(
        f"[QWEN_MAX_AGENT] 🚀 Request | Model={model} | MaxTokens={max_tokens} | "
        f"PromptChars={len(full_prompt):,} | ContentChars={len(content):,}"
    )

    try:
        response = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": full_prompt}],
            temperature=0.1,
            max_tokens=max_tokens,
            stream=True
        )

        full_response = []
        chunk_count = 0
        reasoning_chunk_count = 0
        finish_reason = None
        last_log_time = time.time()

        for chunk in response:
            if not chunk.choices:
                continue

            choice = chunk.choices[0]
            delta = choice.delta

            # Qwen3 reasoning/thinking có thể stream riêng trước content.
            if getattr(delta, "reasoning_content", None):
                reasoning_chunk_count += 1

            if delta and delta.content:
                full_response.append(delta.content)
                chunk_count += 1

            if choice.finish_reason:
                finish_reason = choice.finish_reason

            now = time.time()
            if now - last_log_time >= 10:
                elapsed = int(now - start_time)
                log_info(
                    f"[QWEN_MAX_AGENT] ⏳ Streaming | Elapsed={elapsed}s | "
                    f"ContentChunks={chunk_count} | ReasoningChunks={reasoning_chunk_count}"
                )
                last_log_time = now

        elapsed_total = time.time() - start_time
        result = "".join(full_response)
        _log_response_summary(
            "QWEN_MAX_AGENT",
            elapsed_total,
            result,
            chunk_count,
            finish_reason,
            reasoning_chunk_count
        )
        if finish_reason == "length":
            log_error(
                f"[QWEN_MAX_AGENT] ❌ OUTPUT BỊ CẮT DO MAX TOKENS | "
                f"OutputChars={len(result):,} | "
                f"MaxTokens={max_tokens}"
            )
            return False, result
        if not result.strip():
            log_error(
                f"[QWEN_MAX_AGENT] ⚠️ Request kết thúc nhưng Content rỗng | "
                f"FinishReason={finish_reason or 'UNKNOWN'} | "
                f"ReasoningChunks={reasoning_chunk_count}"
            )

        return True, result

    except Exception as e:
        elapsed_total = time.time() - start_time
        err_details = log_error(
            f"[QWEN_MAX_AGENT] ❌ Request failed/timeout after {elapsed_total:.2f}s | "
            f"Model={model} | MaxTokens={max_tokens}",
            e
        )
        return False, err_details


def extract_json_from_model_response(raw_text: str) -> tuple[bool, dict | list | None, str]:
    """Parse JSON robustly từ output model.

    Hỗ trợ:
    - JSON thuần
    - ```json ... ```
    - model có text thừa trước/sau JSON
    - JSON object/array bằng JSONDecoder.raw_decode

    Trả về (ok, parsed, diagnostic_message).
    """
    if raw_text is None:
        return False, None, "Raw response = None"

    text = raw_text.strip()
    if not text:
        return False, None, "Raw response rỗng"

    # 1) Ưu tiên lấy nội dung trong code fence.
    fenced = re.search(r"```(?:json)?\s*(.*?)\s*```", text, flags=re.IGNORECASE | re.DOTALL)
    candidate = fenced.group(1).strip() if fenced else text

    # 2) Parse trực tiếp trước.
    try:
        return True, json.loads(candidate), "JSON parse OK (direct)"
    except json.JSONDecodeError as direct_err:
        direct_error = direct_err

    # 3) Tìm JSON object/array đầu tiên và dùng raw_decode để không bị text thừa phía sau.
    decoder = json.JSONDecoder()

# Chỉ cho phép parse ROOT JSON đầu tiên.
# Tuyệt đối không scan các object con vì có thể biến JSON bị truncate
# thành một object con "hợp lệ" giả.
    root_obj = candidate.find("{")
    root_arr = candidate.find("[")

    root_positions = [p for p in (root_obj, root_arr) if p >= 0]

    if root_positions:
        root_start = min(root_positions)

    try:
        parsed, end = decoder.raw_decode(candidate[root_start:])

        trailing = candidate[root_start + end:].strip()

        diagnostic = "JSON parse OK (root raw_decode)"

        if trailing:
            diagnostic += f" | Ignored trailing chars={len(trailing)}"

        return True, parsed, diagnostic

    except json.JSONDecodeError:
        pass

    # Diagnostic rõ ràng: vị trí lỗi + preview quanh vị trí đó.
    pos = getattr(direct_error, "pos", None)
    if pos is not None:
        line = getattr(direct_error, "lineno", None)
        col = getattr(direct_error, "colno", None)
        start = max(0, pos - 180)
        end = min(len(candidate), pos + 180)
        preview = candidate[start:end].replace("\n", "\\n")
        diagnostic = (
            f"JSON parse FAILED | {direct_error.msg} | "
            f"pos={pos} | line={line} | col={col} | "
            f"candidate_chars={len(candidate):,} | "
            f"context={preview!r}"
        )
    else:
        diagnostic = f"JSON parse FAILED | candidate_chars={len(candidate):,}"

    return False, None, diagnostic


def validate_rule_matrix_schema(data) -> tuple[bool, str]:
    """Kiểm tra schema tối thiểu mà Agent 2/Review UI đang expect."""
    if not isinstance(data, dict):
        return False, f"Root phải là object, nhận {type(data).__name__}"

    screens = data.get("screens")
    if not isinstance(screens, list):
        return False, "Thiếu field 'screens' hoặc 'screens' không phải array"

    required_rule_fields = {
        "rule_id", "target", "category", "rule_type", "rule_name",
        "test_objective", "source_requirement", "applied_qa_rule", "generation_reason"
    }

    total_rules = 0
    for idx, screen in enumerate(screens, start=1):
        if not isinstance(screen, dict):
            return False, f"screens[{idx-1}] phải là object"
        if not isinstance(screen.get("test_rules"), list):
            return False, f"screens[{idx-1}] thiếu 'test_rules' hoặc không phải array"
        for ridx, rule in enumerate(screen["test_rules"], start=1):
            total_rules += 1
            if not isinstance(rule, dict):
                return False, f"screens[{idx-1}].test_rules[{ridx-1}] phải là object"
            missing = sorted(required_rule_fields - set(rule.keys()))
            if missing:
                return False, f"Rule {idx}.{ridx} thiếu field: {', '.join(missing)}"

    return True, f"Schema OK | Screens={len(screens)} | TestRules={total_rules}"

def python_smart_split_md(raw_text: str, base_filename: str, max_chunk_size: int = 5000) -> tuple[list[dict], str]:
    sections = []
    heading_pattern = re.compile(r'^(#{1,3}\s+.+)$', re.MULTILINE)
    splits = heading_pattern.split(raw_text)

    if len(splits) > 1:
        current_title = f"{base_filename}_Part_1"
        current_content = ""

        for element in splits:
            if heading_pattern.match(element):
                if current_content.strip():
                    sections.append({"title": current_title, "content": current_content.strip()})
                current_title = re.sub(r'^[#\s]+', '', element).strip()
                current_content = element + "\n"
            else:
                current_content += element

        if current_content.strip():
            sections.append({"title": current_title, "content": current_content.strip()})
    else:
        sections.append({"title": base_filename, "content": raw_text})

    final_chunks = []
    for sec in sections:
        content = sec["content"]
        title = sec["title"]
        
        if len(content) > max_chunk_size:
            for i in range(0, len(content), max_chunk_size):
                part_text = content[i:i+max_chunk_size]
                part_title = f"{title}_Part_{i//max_chunk_size + 1}"
                final_chunks.append({"title": part_title, "content": part_text})
        else:
            final_chunks.append({"title": title, "content": content})

    temp_dir = tempfile.mkdtemp()
    extracted_files = []
    zip_filename = f"{base_filename}_MD_Parsed.zip"
    zip_path = os.path.join(tempfile.gettempdir(), zip_filename)

    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zip_out:
        for idx, item in enumerate(final_chunks):
            safe_title = re.sub(r'[\\/*?:"<>|]', '_', item["title"])
            sub_file_name = f"{idx+1:02d}_{safe_title}.md"
            file_path = os.path.join(temp_dir, sub_file_name)

            with open(file_path, "w", encoding="utf-8") as f:
                f.write(item["content"])

            zip_out.write(file_path, arcname=sub_file_name)

            extracted_files.append({
                "screen_name": item["title"],
                "file_name": sub_file_name,
                "file_path": file_path,
                "content": item["content"],
                "char_count": len(item["content"])
            })

    return extracted_files, zip_path

def create_xmind_from_text(tree_data: str, output_path: str, root_title: str = "Kế hoạch & Kịch bản Kiểm thử"):
    try:
        log_info(f"Tiến hành parse dữ liệu text sang cấu trúc file .xmind...")
        root_topic = {
            "id": "root_node",
            "title": root_title,
            "children": {"attached": []}
        }
        
        clean_data_str = re.sub(r'^```[a-zA-Z]*\n', '', tree_data, flags=re.MULTILINE)
        clean_data_str = re.sub(r'\n```$', '', clean_data_str, flags=re.MULTILINE).strip()
        
        node_id_counter = [1]
        lines = clean_data_str.split("\n")
        stack = [(0, root_topic)]

        for line in lines:
            if not line.strip():
                continue
                
            expanded_line = line.replace("\t", "    ")
            indent_spaces = len(expanded_line) - len(expanded_line.lstrip(" "))
            
            content = expanded_line.strip()
            
            if content.startswith("- "):
                content = content[2:].strip()
            elif content.startswith("* "):
                content = content[2:].strip()
                
            content = content.replace("\\n", "\n")

            if not content:
                continue

            level = (indent_spaces // 2) + 1

            new_node = {
                "id": f"node_{node_id_counter[0]}",
                "title": content,
                "children": {"attached": []}
            }
            node_id_counter[0] += 1

            while stack and stack[-1][0] >= level:
                stack.pop()

            parent_node = stack[-1][1]
            parent_node["children"]["attached"].append(new_node)
            stack.append((level, new_node))

        content_json = [{
            "id": "sheet_1",
            "title": "Sơ đồ Kiểm thử QA",
            "rootTopic": root_topic
        }]
        
        manifest_json = {"file-entries": {"content.json": {}, "metadata.json": {}}}
        metadata_json = {}

        with zipfile.ZipFile(output_path, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            zip_file.writestr('content.json', json.dumps(content_json, ensure_ascii=False, indent=2))
            zip_file.writestr('manifest.json', json.dumps(manifest_json, ensure_ascii=False, indent=2))
            zip_file.writestr('metadata.json', json.dumps(metadata_json, ensure_ascii=False, indent=2))
            
        log_info(f"✅ Đã đóng gói thành công file XMind tại: {output_path}")
        return True, None
    except Exception as e:
        err_msg = log_error("Lỗi khi đóng gói file XMind", e)
        return False, err_msg

# ==============================================================================
# HÀM BÓC TÁCH CÂY XMIND VÀ XUẤT EXCEL CHUẨN THEO NHÁNH NHÓM (DỌC THEO TÊN MÀN HÌNH)
# ==============================================================================
# ==============================================================================
# HÀM BÓC TÁCH CÂY XMIND VÀ XUẤT EXCEL CHUẨN (CẬP NHẬT DOCSTRING & QUÉT ĐỘ SÂU)
# ==============================================================================
def convert_tree_to_ui_excel(tree_text: str) -> bytes:
    """
    Chuyển đổi cây XMind/Text thụt lề sang file Excel chuẩn template:
    - Root (Cấp 0): Tên Màn Hình -> Đổ vào ô D2
    - Level 1: Nhóm kiểm thử (1. UI, 2. Validate, 3. Chức năng, 4. Ngoại lệ...) -> Dòng Header xanh lá đậm
    - Level 2: Nhóm con / Trường cần validate (1.1, 1.2...) -> Dòng Header xanh lá nhạt
    - Level 3: Test Case (Tách TC ID -> Cột A, Tên TC -> Cột B)
    - Level 4: Pre-condition -> Cột C
    - Level 5: Steps -> Cột E
    - Level 6: Expected Result -> Cột G
    - Cột D (Importance) và Cột F (Data test) để NULL.
    Chuyển đổi cây XMind/Text thụt lề sang file Excel chuẩn template:
    - Tuân thủ tuyệt đối phân nhóm từ Prompt (UI vs Validation vs Action Controls...).
    - Giữ nguyên các Test Case Validation thuộc đúng phân vùng Validation, không đưa nhầm vào UI.
    """
    clean_text = re.sub(r'^```[a-zA-Z]*\n', '', tree_text, flags=re.MULTILINE)
    clean_text = re.sub(r'\n```$', '', clean_text, flags=re.MULTILINE).strip()
    
    lines = clean_text.split('\n')
    
    root_node = {"title": "Root", "children": []}
    stack = [(0, root_node)]
    
    for line in lines:
        if not line.strip():
            continue
        expanded = line.replace('\t', '    ')
        indent = len(expanded) - len(expanded.lstrip(' '))
        level = (indent // 2) + 1
        content = expanded.strip().lstrip('-* ').replace('\\n', '\n')
        
        node = {"title": content, "children": []}
        
        while stack and stack[-1][0] >= level:
            stack.pop()
            
        stack[-1][1]["children"].append(node)
        stack.append((level, node))
        
    screen_node = root_node["children"][0] if root_node["children"] else {"title": "Tên Màn Hình", "children": []}
    screen_name = screen_node["title"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Test Execution"
    ws.views.sheetView[0].showGridLines = True
    
    fill_green = PatternFill(start_color="81C784", fill_type="solid")       # Level 1: Nhóm lớn (1. UI, 2. Validate...)
    fill_light_green = PatternFill(start_color="C8E6C9", fill_type="solid") # Level 2+: Nhóm con / Trường / Popup
    fill_white = PatternFill(start_color="FFFFFF", fill_type="solid")       # Dòng Test Case
    fill_header_bg = PatternFill(start_color="4DD0E1", fill_type="solid")
    fill_pink = PatternFill(start_color="F8BBD0", fill_type="solid")
    fill_blue_exec = PatternFill(start_color="90CAF9", fill_type="solid")
    
    thin_border = Border(
        left=Side(style='thin', color='B0BEC5'),
        right=Side(style='thin', color='B0BEC5'),
        top=Side(style='thin', color='B0BEC5'),
        bottom=Side(style='thin', color='B0BEC5')
    )
    
    font_title = Font(name="Segoe UI", size=11, bold=True)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_body = Font(name="Segoe UI", size=10)
    
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. Header Metadata Block
    ws['D1'] = "KỊCH BẢN KIỂM THỬ *"
    ws['D1'].font = font_title
    ws['D1'].alignment = align_center
    
    ws['C2'] = "Tên màn hình/Tên chức năng"
    ws['C2'].font = font_bold
    ws['D2'] = screen_name
    ws['D2'].font = font_body
    
    for r in range(1, 5):
        for col in ["C", "D"]:
            ws[f"{col}{r}"].border = thin_border

    # 2. Table Column Headers
    headers = [
        ("A11", "External ID", fill_header_bg),
        ("B11", "Name", fill_pink),
        ("C11", "PreConditions", fill_pink),
        ("D11", "Importance", fill_pink),
        ("E11", "Step", fill_pink),
        ("F11", "Data test", fill_pink),
        ("G11", "Expected Result", fill_pink),
        ("H11", "Actual Result", fill_pink),
        ("I11", "Lần 1", fill_blue_exec),
        ("J11", "Lần 2", fill_blue_exec)
    ]
    for pos, text, fill_color in headers:
        ws[pos] = text
        ws[pos].font = font_bold
        ws[pos].alignment = align_center
        ws[pos].border = thin_border
        ws[pos].fill = fill_color

    current_row = 14

    def clean_prefix(text, prefixes):
        t = str(text).strip()
        for p in prefixes:
            if t.lower().startswith(p.lower()):
                return t[len(p):].strip()
        return t

    def extract_details(node):
        """Hàm đệ quy rút gọn Pre-condition, Step, Expected Result từ bất kỳ cấp con nào"""
        pre, step, exp = "", "", ""
        for child in node.get("children", []):
            title = child["title"]
            title_lower = title.lower()
            
            if any(k in title_lower for k in ["pre-condition", "precondition", "điều kiện"]):
                pre = clean_prefix(title, ["Pre-condition:", "Precondition:", "Tiền điều kiện:"])
                sub_pre, sub_step, sub_exp = extract_details(child)
                step = step or sub_step
                exp = exp or sub_exp
            elif any(k in title_lower for k in ["bước", "step", "steps"]):
                step = clean_prefix(title, ["Các bước thực hiện:", "Steps:", "Step:", "Các bước:"])
                sub_pre, sub_step, sub_exp = extract_details(child)
                exp = exp or sub_exp
            elif any(k in title_lower for k in ["kết quả", "expected", "response"]):
                exp = clean_prefix(title, ["Kết quả mong đợi:", "Expected Result:", "Expected:", "Response (Kết quả mong đợi):"])
            else:
                sub_pre, sub_step, sub_exp = extract_details(child)
                pre = pre or sub_pre
                step = step or sub_step
                exp = exp or sub_exp
                
        return pre, step, exp

    def is_test_case_node(node):
        """Kiểm tra xem node có phải là Test Case hay không"""
        t = node["title"].lower()
        return ("tc_" in t) or (" - " in node["title"] and not node.get("children"))

    def has_test_cases_in_subtree(node):
        """Kiểm tra xem toàn bộ cây con bên dưới có chứa Test Case nào không"""
        if is_test_case_node(node):
            return True
        for child in node.get("children", []):
            if has_test_cases_in_subtree(child):
                return True
        return False

    # 3. Duyệt Cây tổng thể
    for lv1 in screen_node.get("children", []):
        if not has_test_cases_in_subtree(lv1):
            continue

        # Header Nhóm cấp 1 (Xanh lá đậm)
        ws.cell(row=current_row, column=2, value=lv1["title"]).font = font_bold
        for col in range(1, 11):
            c = ws.cell(row=current_row, column=col)
            c.fill = fill_green
            c.border = thin_border
        current_row += 1

        def traverse_and_write(node, indent_level=1):
            nonlocal current_row
            
            if is_test_case_node(node):
                # Xuất dòng Test Case (Cột A: ID, Cột B: Name, Cột C: PreConditions, Cột E: Steps, Cột G: Expected Result)
                tc_full_name = node["title"]
                tc_id, tc_name = "", tc_full_name
                if " - " in tc_full_name:
                    parts = tc_full_name.split(" - ", 1)
                    tc_id, tc_name = parts[0].strip(), parts[1].strip()
                elif ":" in tc_full_name:
                    parts = tc_full_name.split(":", 1)
                    tc_id, tc_name = parts[0].strip(), parts[1].strip()

                pre_condition, steps, expected_result = extract_details(node)

                ws.cell(row=current_row, column=1, value=tc_id).alignment = align_center # Cột A
                ws.cell(row=current_row, column=2, value=tc_name)                       # Cột B
                ws.cell(row=current_row, column=3, value=pre_condition)                 # Cột C
                ws.cell(row=current_row, column=4, value=None)                          # Cột D (NULL)
                ws.cell(row=current_row, column=5, value=steps)                         # Cột E
                ws.cell(row=current_row, column=6, value=None)                          # Cột F (NULL)
                ws.cell(row=current_row, column=7, value=expected_result)               # Cột G

                for col in range(1, 11):
                    c = ws.cell(row=current_row, column=col)
                    c.font = font_body
                    c.fill = fill_white
                    c.border = thin_border
                    if col not in [1, 4]:
                        c.alignment = align_left

                current_row += 1
            else:
                # Nếu là nút Nhóm con / Trường / Popup / Sub-group (Level 2+) -> Dòng Xanh lá nhạt
                if has_test_cases_in_subtree(node):
                    prefix_space = "  " * indent_level
                    ws.cell(row=current_row, column=2, value=f"{prefix_space}{node['title']}").font = font_bold
                    for col in range(1, 11):
                        c = ws.cell(row=current_row, column=col)
                        c.fill = fill_light_green
                        c.border = thin_border
                    current_row += 1

                    for child in node.get("children", []):
                        traverse_and_write(child, indent_level + 1)

        for child in lv1.get("children", []):
            traverse_and_write(child, indent_level=1)

    col_widths = {'A': 15, 'B': 40, 'C': 30, 'D': 12, 'E': 40, 'F': 15, 'G': 40, 'H': 20, 'I': 10, 'J': 10}
    for col, w in col_widths.items():
        ws.column_dimensions[col].width = w

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()

# ==========================================
# 3. KHAI BÁO PROMPTS WEB/APP
# ==========================================
# ==============================================================================
# 1. PROMPT DÀNH CHO BÓC TÁCH MỘT BƯỚC (DÙNG TRỰC TIẾP PROMPT_WEB_UI CHUẨN)
# ==============================================================================
# QA TEST DESIGN PIPELINE — PROMPT PACK v1
# Mục tiêu:
# Agent 1 = QA Brain: đọc Spec + áp dụng QA Test Design Rules -> Test Design / Rule Matrix
# Agent 2 = Renderer: chỉ biến Rule Matrix thành Test Case/XMind, không suy luận thêm
# Web = Renderer cuối nếu cần, không tự nghĩ Business Rule


# ============================================================
# 1. PROMPT AGENT 1 — EXTRACT + APPLY QA RULES
# ============================================================

# NEW 2 — QA TEST DESIGN PIPELINE v2

# ============================================================
# 1. AGENT 1 — QA BRAIN
# ============================================================

PROMPT_AGENT1_EXTRACT_RULE_MATRIX = """
Bạn là SENIOR QA TEST DESIGN ENGINE / QA BRAIN cho hệ thống Banking / Enterprise.

==================================================
NHIỆM VỤ DUY NHẤT
==================================================

Đọc TOÀN BỘ tài liệu nguồn và tạo TEST DESIGN / RULE MATRIX.

Bạn là Agent DUY NHẤT được phép:
- đọc hiểu requirement;
- xác định phạm vi kiểm thử;
- áp dụng QA Test Design Rules;
- tạo EXPLICIT / DERIVED TEST RULE.

Agent 2 và Web Generator phía sau CHỈ được render Rule Matrix thành Test Case.
Agent 2 / Web KHÔNG được suy luận thêm, bù thiếu hoặc tự tạo Rule.

MỤC TIÊU CỐT LÕI:

SPEC
→ Requirement / Constraint có ý nghĩa kiểm thử
→ Target
→ QA Test Design Rule phù hợp
→ Atomic Test Rule
→ Traceability

KHÔNG biến tài liệu thành bản tóm tắt.
KHÔNG biến mỗi Field thành hàng loạt testcase máy móc.
KHÔNG được nén nhiều requirement độc lập thành một Rule chỉ để làm output ngắn.

==================================================
1. SOURCE OF TRUTH
==================================================

TÀI LIỆU NGUỒN là SOURCE OF TRUTH DUY NHẤT.

Chỉ được sử dụng thông tin có trong tài liệu để xác định:
- màn hình / module / page;
- field / control;
- button / action;
- popup;
- grid;
- API / request / response;
- business flow;
- business rule;
- permission;
- validation;
- trạng thái;
- message;
- dữ liệu;
- điều kiện;
- dependency;
- exception;
- constraint.

Không được bịa:
- Business Rule;
- API;
- endpoint;
- HTTP status;
- error message;
- DB behavior;
- permission;
- timeout;
- concurrency;
- validation;
- edge case;
- dữ liệu;
- trạng thái.

Nếu Spec không có căn cứ → KHÔNG tạo Rule.

==================================================
2. QUY TRÌNH BẮT BUỘC — KHÔNG ĐƯỢC DỪNG SỚM
==================================================

Trước khi tạo JSON, PHẢI thực hiện phân tích nội bộ theo 4 pass:

PASS 1 — DOCUMENT INVENTORY
Quét TOÀN BỘ tài liệu từ đầu đến cuối và xác định:
- tất cả Screen / Page / Module;
- tất cả section / subsection có requirement;
- tất cả luồng chính / luồng thay thế / luồng ngoại lệ;
- tất cả bảng mô tả field / control;
- tất cả API / integration;
- tất cả business rule;
- tất cả permission / role;
- tất cả validation / constraint;
- tất cả popup / message;
- tất cả grid / action / trạng thái.

QUAN TRỌNG:
Không được chỉ tập trung vào màn hình / section xuất hiện đầu tiên.
Không được kết thúc phân tích chỉ vì đã có "đủ" một số lượng Rule.
Phải tiếp tục đọc đến cuối tài liệu.

PASS 2 — REQUIREMENT EXTRACTION
Với từng section / requirement có ý nghĩa kiểm thử:
- xác định requirement;
- xác định target;
- xác định điều kiện;
- xác định dữ liệu / giá trị quan trọng;
- xác định hành vi mong đợi;
- xác định dependency / state / permission nếu có.

PASS 3 — QA TEST DESIGN
Với từng requirement:
- nếu Rule được mô tả trực tiếp → EXPLICIT;
- nếu áp dụng một QA Rule được phép để suy ra cách kiểm thử → DERIVED;
- nếu chỉ là kiến thức QA chung không có căn cứ → INFERRED và KHÔNG OUTPUT.

PASS 4 — COVERAGE & QUALITY GATE
Sau khi đã tạo Rule Matrix nội bộ:
- kiểm tra lại từng Screen / Module;
- kiểm tra từng section có requirement quan trọng;
- kiểm tra các business flow;
- kiểm tra các API / integration;
- kiểm tra các validation / constraint;
- kiểm tra permission;
- kiểm tra popup / message;
- kiểm tra grid / action;
- kiểm tra các màn hình xuất hiện ở phần sau tài liệu.

Nếu một requirement có ý nghĩa kiểm thử nhưng chưa có Rule:
→ tạo Rule nếu có QA Rule phù hợp.
Nếu không có QA Rule phù hợp:
→ không bịa Rule.

CHỈ sau khi hoàn tất 4 pass mới được trả JSON.

==================================================
3. SCREEN / MODULE COMPLETENESS
==================================================

Mỗi Screen / Page / Module có requirement kiểm thử phải được phản ánh trong:
screens[].

Nếu tài liệu có nhiều màn hình:
- tạo riêng từng screen;
- không dồn Rule của màn hình sau vào màn hình đầu;
- không bỏ màn hình chỉ vì đã có nhiều Rule ở màn hình trước.

Tên screen_name phải bám theo tên trong tài liệu.

Một Screen có thể có nhiều loại Rule:
UI / VALIDATION / ACTION / GRID / POPUP / BUSINESS_FLOW / EXCEPTION.

==================================================
4. ATOMIC TEST RULE — CRITICAL
==================================================

Mỗi test_rule phải là một mục tiêu kiểm thử có thể đứng độc lập.

ĐỊNH NGHĨA:
1 Rule = 1 Test Objective độc lập.

KHÔNG GỘP nếu các phần trong cùng một câu có thể:
- có điều kiện khác nhau;
- có expected result khác nhau;
- có target khác nhau;
- có dữ liệu / parameter khác nhau;
- có thể fail độc lập;
- hoặc Agent 2 sẽ phải tạo hai Test Case khác nhau để kiểm tra đầy đủ.

Ví dụ KHÔNG ĐƯỢC GỘP:

"Kiểm tra dropdown Chi nhánh: gọi API Active, mặc định Tất cả, chọn Tất cả chọn toàn bộ, bỏ một giá trị thì bỏ Tất cả."

→ phải tách thành các Rule độc lập nếu mỗi hành vi có objective riêng.

Tuy nhiên KHÔNG được over-split một Rule thành nhiều Rule chỉ vì có nhiều giá trị test.

Ví dụ:
"Field CIF tối đa 10 ký tự + Boundary Value"
→ có thể là 1 DERIVED Rule nếu cùng một objective kiểm tra boundary của Max Length.

NGUYÊN TẮC:
Tách theo TEST OBJECTIVE, không tách theo từng từ / từng giá trị.

==================================================
5. REQUIREMENT PRESERVATION — KHÔNG ĐƯỢC SUMMARY LOSS
==================================================

source_requirement KHÔNG được là câu tóm tắt quá chung chung nếu Spec có chi tiết quan trọng.

Phải giữ lại các thông tin ảnh hưởng trực tiếp đến Test Case, ví dụ:
- giá trị mặc định;
- Min / Max;
- Boundary;
- điều kiện;
- trạng thái;
- role / permission;
- API;
- parameter;
- request field;
- response;
- message;
- thời gian;
- thứ tự;
- dependency;
- hành vi khi chọn / bỏ chọn;
- hành vi khi lỗi;
- điều kiện chuyển trạng thái.

Ví dụ KHÔNG ĐƯỢC:
"Kiểm tra tham số tìm kiếm gửi API."

Nếu Spec mô tả cụ thể nhiều parameter:
→ source_requirement phải nêu các parameter / điều kiện quan trọng đó.

Ví dụ KHÔNG ĐƯỢC:
"Kiểm tra dropdown loại phát hành."

Nếu Spec có:
codeType = BOND_TYPE
status = Null
→ phải giữ thông tin này trong source_requirement / test_objective khi cần.

Không được làm mất thông tin chỉ vì muốn JSON ngắn.

==================================================
6. EXPLICIT vs DERIVED
==================================================

EXPLICIT:
Rule / hành vi được mô tả trực tiếp trong Spec.

Ví dụ:
Spec:
"Nếu From Date > To Date thì tự động cập nhật lại ngày."

→ EXPLICIT.

Không được biến requirement trực tiếp thành DERIVED chỉ vì bạn dùng một QA Rule để kiểm tra nó.

DERIVED:
Spec cung cấp requirement / constraint,
sau đó QA Rule được phép dùng để thiết kế cách kiểm thử.

Ví dụ:
Spec:
"CIF tối đa 10 ký tự."

QA Rule:
Boundary Value / Max Length.

→ DERIVED.

DERIVED bắt buộc có:
- source_requirement;
- applied_qa_rule;
- generation_reason.

EXPLICIT bắt buộc có source_requirement.
applied_qa_rule có thể là:
"EXPLICIT FROM SPEC"

==================================================
7. INFERRED — TUYỆT ĐỐI KHÔNG OUTPUT
==================================================

Nếu Rule chỉ đến từ kiến thức QA chung mà không có requirement / constraint tương ứng trong Spec:
→ INFERRED
→ KHÔNG đưa vào JSON.

Ví dụ Spec không nói:
- concurrency;
- timeout;
- 500;
- 403;
- session expired;
- duplicate;
- retry;

→ không tự tạo các Rule này.

==================================================
8. QA RULE LIBRARY
==================================================

Chỉ được sử dụng QA Rule trong danh sách dưới đây.

1. TEXTBOX
- Default Value
- Placeholder
- Required
- Enable/Disable
- Readonly
- Character Set
- Unicode
- Special Character
- Whitespace
- Trim Space
- Min Length
- Max Length
- Boundary Value
- Valid Value
- Invalid Value
- Copy/Paste

2. NUMERIC TEXTBOX
- Required
- Enable/Disable
- Readonly
- Numeric Format
- Min
- Max
- Boundary Value
- Zero
- Positive
- Negative
- Decimal
- Thousand Separator
- Rounding
- Copy/Paste

3. DATEPICKER
- Default
- Placeholder
- Required
- Readonly
- Calendar Selection
- Keyboard Input
- Invalid Format
- Invalid Date
- Clear
- Min Date
- Max Date
- Boundary Date
- Past/Future
- Leap Year
- Start Date <= End Date

4. DROPDOWN / COMBOBOX
- Default
- Placeholder
- Required
- Enable/Disable
- Readonly
- Data List
- Display Order
- Duplicate
- Empty Data
- Single/Multiple Selection
- Clear Selection
- Search
- Dependency

5. SEARCH / FILTER
- Default
- Placeholder
- Required
- Enable/Disable
- Readonly
- Trim Space
- Max Length
- Copy/Paste
- Exact Search
- Partial Search
- Special Character
- No Result
- Multiple Conditions
- Reset Filter

6. ACTION CONTROL
- Label
- Enable/Disable
- Action Trigger
- Navigation
- Double Click / Multi Click

Không tạo Action Control chỉ vì có:
- focus textbox;
- open dropdown;
- open datepicker;
- click checkbox.

7. DATA GRID
UI:
- Column
- Order
- Width
- Alignment
- Format
- Null/Empty
- Wrap
- Tooltip
- Loading

Mapping:
- API → UI
- Missing/Extra/Duplicate
- Refresh

Pagination:
- Page size
- Next/Previous
- Page number
- STT

Sorting:
- Ascending
- Descending
- Text
- Number
- Date

Row Action:
- View
- Edit
- Delete
- Select
- Expand/Collapse

8. POPUP
- Trigger condition
- Title
- Content
- Button
- Close
- Cancel
- Confirm
- Overlay

Business result sau Confirm → BUSINESS_FLOW.

9. BUSINESS FLOW
- Happy Path
- Negative Flow
- Business Rule Validation
- Cross-field Dependency
- State Transition
- Permission
- API Request/Response
- Status Code
- DB Update
- Message
- Navigation

10. EXCEPTION
- Timeout
- Network Error
- System Error
- Duplicate
- Concurrent Editing
- Empty Data

Mọi QA Rule trên CHỈ được áp dụng khi Spec có căn cứ.

==================================================
9. VALIDATION RULE — KHÔNG AUTO-GENERATE
==================================================

Field tồn tại KHÔNG đồng nghĩa phải tạo:
Required / Unicode / Special Character / Trim / Copy-Paste / Min / Max / Boundary / Invalid.

Chỉ tạo khi:
- Spec mô tả trực tiếp; hoặc
- Spec có constraint đủ để áp dụng một QA Rule.

Kiểu dữ liệu chỉ giúp chọn QA Rule phù hợp.
Kiểu dữ liệu KHÔNG phải căn cứ duy nhất để tạo toàn bộ validation.

==================================================
10. BUSINESS FLOW / API / UI
==================================================

Nếu Spec có API request:
phải giữ các thông tin quan trọng như:
- API name / endpoint nếu có;
- Method nếu có;
- request parameter;
- field;
- giá trị mặc định;
- điều kiện;
- response;
- status / message nếu Spec có.

Nếu Spec có permission:
tạo Rule riêng cho từng behavior/role khi behavior khác nhau.

Nếu Spec có state transition:
tạo Rule cho từng transition được mô tả.

Nếu Spec có popup:
tách:
- trigger;
- content / message;
- action;
- business result sau confirm/cancel
nếu đây là các objective độc lập.

==================================================
11. TARGET
==================================================

target phải là đối tượng thực sự được kiểm thử:
- screen;
- field;
- button;
- dropdown;
- grid;
- popup;
- API;
- business flow;
- trạng thái;
- v.v.

Không dùng target quá chung như:
"Màn hình"
nếu có thể xác định chính xác target.

==================================================
12. CATEGORY
==================================================

Chỉ sử dụng:
- UI
- VALIDATION
- ACTION
- GRID
- POPUP
- BUSINESS_FLOW
- EXCEPTION

==================================================
13. TRACEABILITY
==================================================

Mỗi Rule bắt buộc có:

- rule_id
- target
- category
- rule_type
- rule_name
- test_objective
- source_requirement
- applied_qa_rule
- generation_reason

rule_id phải unique trong toàn bộ output.

Không reset rule_id theo từng screen.

source_requirement phải đủ cụ thể để QA Lead có thể lần ngược về requirement gốc.

generation_reason phải giải thích ngắn gọn:
- vì sao Rule này được tạo;
- nếu DERIVED thì QA Rule nào đã được áp dụng.

==================================================
14. COVERAGE GATE — CRITICAL
==================================================

Trước khi output, tự kiểm tra nội bộ:

A. SCREEN COVERAGE
- Có bỏ Screen / Page / Module nào không?

B. SECTION COVERAGE
- Có bỏ section / subsection có requirement không?

C. REQUIREMENT COVERAGE
- Mỗi requirement có ý nghĩa kiểm thử đã được map chưa?

D. FLOW COVERAGE
- Luồng chính?
- Luồng thay thế?
- Luồng ngoại lệ?
- State transition?
- Permission?

E. CONTROL COVERAGE
- Field?
- Dropdown?
- Date?
- Button?
- Popup?
- Grid?

F. INTEGRATION COVERAGE
- API?
- Request?
- Response?
- Mapping?
- Parameter?

G. DETAIL PRESERVATION
- Có làm mất giá trị / điều kiện / message / parameter / status / default quan trọng không?

H. ATOMICITY
- Có Rule nào đang chứa nhiều objective độc lập không?

Nếu phát hiện thiếu:
→ quay lại bước phân tích và bổ sung trước khi output.

Không được tự ép coverage bằng cách bịa Rule.

==================================================
15. CHỐNG "RULE COMPRESSION" — CRITICAL
==================================================

Không được dùng các câu tổng quát như:
- "Kiểm tra các field tìm kiếm."
- "Kiểm tra dropdown."
- "Kiểm tra các tham số API."
- "Kiểm tra các chức năng của màn hình."
- "Kiểm tra validation theo Spec."

nếu bên trong có nhiều requirement độc lập.

Rule phải chỉ ra:
- target cụ thể;
- hành vi cụ thể;
- điều kiện cụ thể;
- expected behavior cụ thể.

Mục tiêu là:
REQUIREMENT → TEST RULE

không phải:
REQUIREMENT → SUMMARY → TEST RULE.

==================================================
16. OUTPUT JSON — BẮT BUỘC
==================================================

CHỈ trả về JSON hợp lệ.
KHÔNG markdown.
KHÔNG ```json.
KHÔNG giải thích trước hoặc sau JSON.
KHÔNG thêm field ngoài schema.

Schema:

{
  "test_design_version": "3.0",
  "screens": [
    {
      "screen_name": "",
      "test_rules": [
        {
          "rule_id": "",
          "target": "",
          "category": "UI | VALIDATION | ACTION | GRID | POPUP | BUSINESS_FLOW | EXCEPTION",
          "rule_type": "EXPLICIT | DERIVED",
          "rule_name": "",
          "test_objective": "",
          "source_requirement": "",
          "applied_qa_rule": "",
          "generation_reason": ""
        }
      ]
    }
  ]
}

Không tạo screen giả.
Không tạo test_rule không có source_requirement.
Không tạo DERIVED không có applied_qa_rule.
Không tạo duplicate rule.

==================================================
17. FINAL SELF-CHECK
==================================================

Trước khi trả JSON, tự kiểm tra:

1. Đã đọc đến CUỐI tài liệu chưa?
2. Đã nhận diện TẤT CẢ screen/page/module chưa?
3. Có bỏ section quan trọng ở phần sau tài liệu không?
4. Có Rule nào chỉ vì "best practice QA" không?
5. Có tự bịa Business Rule không?
6. Có tự bịa API / Message / Status / DB không?
7. EXPLICIT có thực sự được mô tả trực tiếp trong Spec không?
8. DERIVED có requirement + QA Rule làm căn cứ không?
9. Mỗi Rule có source_requirement cụ thể không?
10. Có Rule nào đang gộp nhiều objective độc lập không?
11. Có Rule nào quá chung chung không?
12. Có làm mất parameter / value / condition / message / status / default quan trọng không?
13. rule_id có unique toàn bộ output không?
14. JSON có đúng schema không?
15. Output có chứa DUY NHẤT JSON không?

Nếu bất kỳ câu nào chưa đạt:
→ sửa nội bộ trước khi trả kết quả.

=== TÀI LIỆU NGUỒN ===
{content}
"""



# ============================================================
# 2. AGENT 2 — RULE MATRIX -> TEST CASE
# ============================================================

PROMPT_AGENT2_GEN_XMIND_FROM_RULE_MATRIX = """
Bạn là TEST CASE RENDERER.

Bạn KHÔNG phải QA Analyst.

Bạn KHÔNG phân tích lại tài liệu.
Bạn KHÔNG suy luận Business Rule.
Bạn KHÔNG bổ sung Test Rule.
Bạn KHÔNG mở rộng scope bằng kiến thức QA riêng.

NHIỆM VỤ DUY NHẤT:
Đọc TEST DESIGN / RULE MATRIX JSON và chuyển từng test_rule
thành đúng 1 Test Case dạng cây XMind.

==================================================
=== SOURCE OF TRUTH
==================================================

TEST DESIGN / RULE MATRIX là SOURCE OF TRUTH DUY NHẤT.

Không được:
- thêm Test Rule
- bỏ Test Rule
- gộp Test Rule
- thay đổi ý nghĩa Test Rule
- thêm Business Rule
- thêm Validation
- thêm Edge Case
- thêm API
- thêm Status Code
- thêm DB behavior
- thêm Permission
- thêm Exception

Nếu JSON có 100 test_rule:
phải tạo đủ 100 testcase.

==================================================
=== MAPPING
==================================================

Mỗi test_rule → đúng 1 Test Case.

EXPLICIT → generate.
DERIVED → generate.
INFERRED → không generate.

==================================================
=== CÁCH VIẾT
==================================================

Dựa vào:
- target
- category
- rule_name
- test_objective
- source_requirement
- applied_qa_rule

Nếu JSON không cung cấp:
- API
- Message
- Status
- Giá trị
- Tên trạng thái
- Điều kiện cụ thể

thì KHÔNG được tự bịa.

==================================================
=== CẤU TRÚC XMIND
==================================================

- Tên Màn hình
  - 1. Kiểm tra UI
    - Target
      - TC_001 - Tên Test Case
        - Pre-condition: ...
          - Các bước thực hiện: 1. ... -> 2. ... -> 3. ...
            - Kết quả mong đợi: 1. ... -> 2. ... -> 3. ...

Cấp 1: Tên màn hình.
Cấp 2: Nhóm kiểm thử.
Cấp 3: Target.
Cấp 4: Test Case.
Cấp 5: Pre-condition.
Cấp 6: Các bước thực hiện.
Cấp 7: Kết quả mong đợi.

Pre-condition là child trực tiếp của Test Case.
Các bước thực hiện là child trực tiếp của Pre-condition.
Kết quả mong đợi là child trực tiếp của Các bước thực hiện.

Steps và Expected Result phải nằm trên một dòng.
Không Enter để tạo node mới.

==================================================
=== CATEGORY
==================================================

UI → Kiểm tra UI
VALIDATION → Kiểm tra Validation
ACTION → Kiểm tra Action Controls
GRID → Kiểm tra Data Grid
POPUP → Kiểm tra Popup
BUSINESS_FLOW → Kiểm tra Luồng nghiệp vụ
EXCEPTION → Kiểm tra Ngoại lệ

==================================================
=== THỨ TỰ
==================================================

1. Kiểm tra UI
2. Kiểm tra Validation
3. Kiểm tra Action Controls
4. Kiểm tra Data Grid
5. Kiểm tra Popup
6. Kiểm tra Luồng nghiệp vụ
7. Kiểm tra Ngoại lệ

Chỉ tạo category có dữ liệu.

==================================================
=== OUTPUT
==================================================

CHỈ trả plain text dạng cây bằng dấu "-".

Không JSON.
Không markdown code block.
Không giải thích.
Không thống kê.
Không thêm testcase.
Không bỏ testcase.
Tất cả bằng tiếng Việt.

=== TEST DESIGN / RULE MATRIX JSON ===
{content}
"""


# ============================================================
# 3. WEB UI — FINAL GENERATOR
# ============================================================

PROMPT_WEB_UI = """
Bạn là Test Case Generator.

NHIỆM VỤ DUY NHẤT:
Đọc TEST DESIGN / RULE MATRIX và chuyển từng Test Rule
thành Test Case chi tiết bằng tiếng Việt.

==================================================
=== SOURCE OF TRUTH — CRITICAL
==================================================

TEST DESIGN / RULE MATRIX là SOURCE OF TRUTH DUY NHẤT.

KHÔNG:
- tự phân tích lại Business Rule
- tự bổ sung Test Rule
- tự tạo Edge Case
- tự tạo API
- tự tạo Status Code
- tự tạo Database behavior
- tự tạo Permission
- tự tạo Error Message
- tự tạo Timeout
- tự tạo Concurrent Editing
- tự tạo Validation
- tự mở rộng scope

Nếu Rule Matrix không có thông tin:
KHÔNG được bịa.

==================================================
=== 1 RULE = 1 TEST CASE
==================================================

Mỗi test_rule → đúng 1 testcase.

Không gộp.
Không bỏ.
Không tạo thêm testcase.

==================================================
=== RULE TYPE
==================================================

EXPLICIT → Generate.
DERIVED → Generate.
INFERRED → Không generate.

==================================================
=== TEST CASE WRITING
==================================================

Sử dụng:
- target
- category
- rule_name
- test_objective
- source_requirement
- applied_qa_rule

Nếu Rule Matrix cung cấp:
- Field → dùng đúng Field.
- Button → dùng đúng Button.
- Min/Max → dùng đúng giá trị.
- API → dùng đúng API.
- Message → dùng đúng Message.
- Status → dùng đúng Status.
- Condition → dùng đúng Condition.

Nếu không có:
không tự bịa.

==================================================
=== CẤU TRÚC XMIND
==================================================

- Cấp 1: Tên Màn hình / Module
  - Cấp 2: Nhóm kiểm thử
    - Cấp 3: Target
      - TC_(STT) - Tên Test Case
        - Pre-condition: ...
          - Các bước thực hiện: 1. ... -> 2. ... -> 3. ...
            - Kết quả mong đợi: 1. ... -> 2. ... -> 3. ...

Pre-condition là child trực tiếp của Test Case.
Các bước thực hiện là child trực tiếp của Pre-condition.
Kết quả mong đợi là child trực tiếp của Các bước thực hiện.

Steps và Expected Result:
KHÔNG xuống dòng để tạo node mới.

==================================================
=== CATEGORY
==================================================

UI → Kiểm tra UI
VALIDATION → Kiểm tra Validation
ACTION → Kiểm tra Action Controls
GRID → Kiểm tra Data Grid
POPUP → Kiểm tra Popup
BUSINESS_FLOW → Kiểm tra Luồng nghiệp vụ
EXCEPTION → Kiểm tra Ngoại lệ

==================================================
=== THỨ TỰ NHÁNH
==================================================

1. Kiểm tra UI
2. Kiểm tra Validation
3. Kiểm tra Action Controls
4. Kiểm tra Data Grid
5. Kiểm tra Popup
6. Kiểm tra Luồng nghiệp vụ
7. Kiểm tra Ngoại lệ

Chỉ tạo nhánh có Test Rule.

==================================================
=== OUTPUT
==================================================

CHỈ trả plain text dạng cây thụt lề bằng dấu "-".

KHÔNG JSON.
KHÔNG markdown code block.
KHÔNG giải thích.
KHÔNG thống kê.
KHÔNG thêm testcase.
KHÔNG bỏ testcase.
Tất cả bằng tiếng Việt.

=== TEST DESIGN / RULE MATRIX ===
{content}
"""


# ==========================================
# 4. KHAI BÁO PROMPTS API
# ==========================================

PROMPT_API_SPEC = """
Bạn hãy đóng vai trò là Senior API QA Automation Lead. Hãy đọc kỹ tài liệu API Specification + Tài liệu mô tả luồng nghiệp vụ (Swagger / Postman / Markdown Spec) dưới đây và tạo danh sách Test Case API chi tiết BẰNG TIẾNG VIỆT theo đúng CẤU TRÚC PHÂN CẤP THỤT LỀ (sử dụng dấu gạch đầu dòng '-').

=== TUYỆT ĐỐI CẤM (CRITICAL RULE) ===
1. KHÔNG XUẤT DẠNG JSON, KHÔNG BỌC TRONG json ...  HOẶC markdown ... .
2. CHỈ TRẢ VỀ VĂN BẢN THUẦN (PLAIN TEXT) DẠNG CÂY THỤT LỀ DÙNG DẤU GẠCH ĐẦU DÒNG -.
3. KHÔNG TÓM TẮT: Phải bắt cặp 100% tất cả Endpoint, Parameter, Header, Payload và Status Code xuất hiện trong tài liệu.
4. Phải sử dụng tài liệu thiết kế API để sinh các case method, URL, Header, Phân quyền, Validate, Luồng chính, Luồng ngoại lệ. Không được suy diễn thêm các case không có trong tài liệu.
5. Đối với các case business rule đặc thù, nếu tài liệu có mô tả thì sinh ra các case riêng biệt, không gộp chung vào Happy Path, phải đọc kỹ tài liệu của BA để sinh ra các case nghiệp vụ đặc thù.
6. Chỉ đọc API spec để sinh các case phân quyền, validate, luồng chính, luồng ngoại lệ (nếu có). Không được suy diễn thêm các case không có trong tài liệu. Trường hợp các case trùng với tài liệu BA thì bỏ qua không sinh thêm.
7. Đối với trường hợp gọi sang các hệ thống khác nếu tài liệu mô tả error code thì sinh ra các case riêng biệt, không gộp chung vào Happy Path.

=== QUY TẮC CẤU TRÚC PHÂN CẤP VÀ ĐỊNH DẠNG CÂY XMIND LỒNG DỌC (CRITICAL STRUCTURE) ===
Mỗi Test Case API BẮT BUỘC phải xuất theo cấu trúc LỒNG DỌC CẤP DƯỚI (Parent-Child Tree) bằng việc thụt lề liên tiếp:

- (Cấp 1: Nhóm API / Module Service)
  - (Cấp 2: Nhóm kiểm thử (Kiểm tra xác thực và token / Kiểm tra phân quyền / Kiểm tra Validate / Kiểm tra luồng chính / Kiểm tra ngoại lệ))
    - (Cấp 3: Endpoint / Field Name / Chức năng)
      - TC_API_(STT) - (Tên API/Endpoint) - (Mục đích kiểm thử)
        - Pre-condition: (Điều kiện login, token, URL endpoint, phân quyền API)
          - Steps & Data test: 1. Gọi API (Method) -> 2. Truyền Header Auth -> 3. Truyền Body JSON: {{ "sample": "data" }} -> 4. Send Request
            - Response (Kết quả mong đợi): 1. Trạng thái (Thành công/Thất bại) -> 2. Status Code (200/400/401/403/500) -> 3. Message trả về -> 4. Response Body Schema

TUYỆT ĐỐI BẮT BUỘC VỀ QUY TẮC SUB-TOPIC:
1. Pre-condition LÀ SUB-TOPIC TRỰC TIẾP (CẤP CON) CỦA TÊN TEST CASE.
2. Steps & Data test LÀ SUB-TOPIC TRỰC TIẾP (CẤP CON) CỦA PRE-CONDITION.
3. Response (Kết quả mong đợi) LÀ SUB-TOPIC TRỰC TIẾP (CẤP CON) CỦA STEPS & DATA TEST.
4. BÊN TRONG NỘI DUNG CÁC BƯỚC VÀ KẾT QUẢ: KHÔNG Enter tạo node mới, trình bày dạng chuỗi nối tiếp.

Mô tả chi tiết mã lỗi mặc định:
- 401: Token signature is invalid / Full authentication is required to access this resource / Token expired
- 403: Access is denied
- 404: Resource not found
- 405: Method not allowed
- 500: Internal server error

=== BỘ CHECKLIST VÀ QUY TẮC THIẾT KẾ TEST CASE CHI TIẾT ===
1. KIỂM TRA XÁC THỰC VÀ TOKEN:
1.1 Kiểm tra Xác thực và token:
+ Kiểm tra không truyền Authorization
+ Kiểm tra truyền Authorization hết hạn
+ Kiểm tra truyền Authorization hợp lệ
1.2 Kiểm tra URL & Method:
+ Kiểm tra đúng URL sai Method
+ Kiểm tra sai URL sai Method

2. VALIDATION REQUEST:
CHECKLIST VALIDATION THEO KIỂU DỮ LIỆU:
- String/Text: Required, null, empty, whitespace, Min Length, Max Length, vượt Max Length, ký tự không hợp lệ, format/pattern, giá trị hợp lệ, giá trị không hợp lệ, ký tự đặc biệt nếu rule có quy định.
- Number: Required, null, empty, whitespace, sai kiểu dữ liệu, số âm, số 0, số dương, Min, Max, Boundary Min-1/Min/Min+1/Max-1/Max/Max+1, vượt giới hạn, số thập phân, số chữ số nguyên/thập phân, format số nếu Spec quy định.
- List/Array: Required, null, empty, sai kiểu dữ liệu, phần tử sai kiểu dữ liệu, danh sách rỗng, số lượng phần tử Min/Max nếu có, phần tử hợp lệ, phần tử không hợp lệ, phần tử trùng lặp nếu có rule, kết hợp nhiều phần tử hợp lệ/không hợp lệ.
- Date/DateTime: Required, null, empty, whitespace, sai kiểu dữ liệu, sai format, ngày không tồn tại, Min Date, Max Date, Boundary Min-1/Min/Min+1/Max-1/Max/Max+1 nếu có, ngày quá khứ/tương lai nếu có rule, quan hệ giữa các trường ngày nếu có.
- Boolean: Required, null, empty, sai kiểu dữ liệu, true, false.
- Object: Required, null, empty, sai kiểu dữ liệu, thiếu field bắt buộc bên trong Object, field bên trong sai kiểu dữ liệu, Object hợp lệ.
- Enum/giá trị cố định: Required, null, empty, giá trị hợp lệ thuộc danh sách, giá trị không thuộc danh sách, sai kiểu dữ liệu.
- Path Parameter: Required, thiếu parameter, null, empty, whitespace, sai kiểu dữ liệu, format không hợp lệ, giá trị không tồn tại, giá trị hợp lệ.
- Query Parameter: Required, thiếu parameter, null, empty, whitespace, sai kiểu dữ liệu, format không hợp lệ, Min/Max/Boundary nếu có, giá trị hợp lệ, giá trị không hợp lệ.
- Header: Header bắt buộc thiếu, null/empty, sai format, giá trị không hợp lệ, giá trị hợp lệ.
- File/Download: Kiểm tra file không tồn tại, sai định dạng, file rỗng, file quá dung lượng nếu có rule; với response download kiểm tra link, status, tên file, định dạng, dung lượng và nội dung file theo Spec.
Lưu ý: Chỉ sinh Validation Test Case khi field có validation rule được mô tả trong tài liệu hoặc là field bắt buộc phải truyền theo schema. Kiểu dữ liệu chỉ dùng để xác định bộ rule validation phù hợp; không dùng kiểu dữ liệu để mặc định sinh tất cả các validation.

3. KIỂM TRA HAPPY CASE:
Sinh Test Case cho từng luồng API thành công được mô tả trong Spec.
Sử dụng Request hợp lệ, đúng Schema và thỏa mãn điều kiện nghiệp vụ.
Kiểm tra: API xử lý trạng thái thành công, Status Code 2xx theo Spec, Response Code/Message/Body đúng Spec, Dữ liệu Response đúng Schema, Dữ liệu được xử lý đúng theo chức năng, Mapping API → DB nếu Spec yêu cầu.
Mỗi luồng thành công khác nhau = 1 Test Case độc lập.
Không đưa các Business Rule đặc thù vào Happy Case nếu Rule đó cần được kiểm tra riêng tại Business Case.

4. KIỂM TRA BUSINESS CASE:
Kiểm tra API xử lý các trạng thái khác thành công (ví dụ như trạng thái của bản ghi khác success, trạng thái của dữ liệu khác success). Đồng thời sinh ra các Test Case riêng cho từng Business Rule đặc thù được mô tả trong Spec. Mỗi Business Rule đặc thù = 1 Test Case độc lập.

5. KIỂM TRA EXCEPTION CASE:
Các trường hợp ngoại lệ được mô tả trong Spec, bao gồm: 4xx, 5xx, Timeout, Payload quá lớn, Database Error simulation.

=== NỘI DUNG TÀI LIỆU API SPEC NGUỒN ===
- Kế hoạch:
{plan_content}

- Thiết kế API:
{design_doc}

- Nghiệp vụ BA:
{ba_doc}
"""

# === 4.1 PROMPTS DÀNH CHO REST API ===
PROMPT_AGENT1_EXTRACT_API_JSON_SCOPE = """
Bạn là Senior API QA Lead. Hãy đọc kỹ tài liệu API Spec (Swagger, Postman, Markdown) và Tài liệu Nghiệp vụ BA dưới đây.
Nhiệm vụ của bạn là LẬP BẢN TÓM TẮT PHẠM VI KIỂM THỬ API (API Scope Plan) và đóng gói dưới dạng CẤU TRÚC JSON CHUẨN. Tuyệt đối KHÔNG sinh Test Case chi tiết ở bước này.

YÊU CẦU ĐỊNH DẠNG JSON TRẢ VỀ (BẮT BUỘC):
Trả về duy nhất 1 khối JSON hợp lệ nằm trong thẻ ```json ... ``` theo Schema bên dưới:

{{
  "api_modules": [
    {{
      "module_name": "Tên Module/Service",
      "endpoints": [
        {{
          "endpoint_path": "/api/v1/...",
          "method": "POST/GET/PUT/DELETE",
          "summary": "Mô tả chức năng API",
          "headers": [{"name": "Header-Name", "required": true/false}],
          "request_fields": [
            {{
              "field_name": "Tên param/body field",
              "data_type": "string/number/boolean/array",
              "required": true/false,
              "location": "body/query/path/header",
              "field_logic": "Quy tắc validation hoặc format"
            }}
          ],
          "default_cases": ["Các case mặc định: Thiếu Auth Header, Hết hạn Token, Sai Method..."],
          "validation_cases": ["Các case validate đầu vào từng field"],
          "happy_path_cases": ["Các luồng gọi API thành công 2xx"],
          "business_rule_cases": ["Các quy tắc nghiệp vụ sâu, phân quyền, trạng thái bản ghi"],
          "exception_cases": ["Timeout, Payload quá lớn, Lỗi hệ thống 5xx"]
        }}
      ]
    }}
  ]
}}

=== TÀI LIỆU API SPEC & BA ĐẦU VÀO ===
{content}
"""

PROMPT_AGENT2_GEN_API_XMIND_FROM_JSON = """
Bạn là Senior API QA Automation Lead. Nhiệm vụ của bạn là đọc CẤU TRÚC JSON API SCOPE PLAN ĐÃ ĐƯỢC QA LEAD REVIEW VÀ BỔ SUNG LOGIC, sau đó dệt thành bộ Test Case API chi tiết BẰNG TIẾNG VIỆT theo CẤU TRÚC PHÂN CẤP THỤT LỀ (dùng dấu gạch đầu dòng '-').

=== TUYỆT ĐỐI CẤM (CRITICAL RULE) ===
1. KHÔNG XUẤT DẠNG JSON, KHÔNG BỌC TRONG json ... HOẶC markdown ... .
2. CHỈ TRẢ VỀ VĂN BẢN THUẦN (PLAIN TEXT) DẠNG CÂY THỤT LỀ DÙNG DẤU GẠCH ĐẦU DÒNG -.

=== QUY TẮC CẤU TRÚC XMIND LỒNG DỌC (CRITICAL STRUCTURE) ===
- [Tên Module/Service]
  - [Endpoint Path & Method]
    - [Nhóm kiểm thử: Case Mặc Định / Validate Request / Happy Path / Business Rules / Ngoại Lệ]
      - TC_API_(STT) - [Method] [Endpoint] - [Mục đích kiểm thử]
        - Pre-condition: [Token, Endpoint URL, Phân quyền]
          - Steps & Data test: 1. Gọi API (Method) -> 2. Truyền Header Auth -> 3. Truyền Body/Query: {{...}} -> 4. Send Request
            - Response (Kết quả mong đợi): 1. Status Code -> 2. Response Message/Body Schema -> 3. Xử lý DB/Log

=== DỮ LIỆU JSON API ĐÃ DUYỆT BỞI QA LEAD ===
{content}
"""

# ==========================================
# 5. GIAO DIỆN CHÍNH (STREAMLIT TABS)
# ==========================================
tab_ui, tab_api = st.tabs(["📱 Web & Mobile App (UI/UX)", "🔌 RESTful / SOAP API Testing"])

# --- TAB 1: WEB & MOBILE APP (UI/UX) ---
with tab_ui:
    st.subheader("📌 Kiểm thử Giao diện & Luồng Người dùng (Web/App)")
    st.caption("Hỗ trợ bóc tách SRS PDF/MD hoặc Tóm tắt Scope Plan (JSON) có tương tác Review trực tiếp trước khi sinh XMind.")

    ui_workflow_mode = st.radio(
        "🎯 Chọn phương thức xử lý UI:",
        ["1-Click Auto Batching (Luồng Cũ - Tự động cắt .md)", "Interactive Review Plan (Luồng Mới - Review JSON Scope & Rule)"],
        index=1,
        key="ui_workflow_mode"
    )

    uploaded_ui_file = st.file_uploader("Tải lên tài liệu SRS / Figma Layout Text (.pdf, .md, .txt)", type=["pdf", "md", "txt"], key="main_ui_file_uploader")

    if uploaded_ui_file is not None:
        st.markdown("---")
        
        # LUỒNG 1: 1-CLICK AUTO BATCHING
        if ui_workflow_mode.startswith("1-Click"):
            st.write("### 🤖 Bước 1: Agent 1 - Đọc & Cắt tài liệu bằng Python Parser (Instant)")
            
            if st.button("🚀 Kích hoạt Agent 1 (Bóc tách tài liệu -> File .md)", key="btn_agent1_ui_main"):
                with st.status("Python Parser đang phân tích Heading và cắt nhỏ file...", expanded=True) as status:
                    try:
                        base_name = uploaded_ui_file.name.rsplit('.', 1)[0]
                        file_bytes = uploaded_ui_file.read()
                        
                        if uploaded_ui_file.name.lower().endswith(".pdf"):
                            reader = PdfReader(io.BytesIO(file_bytes))
                            raw_content = "\n\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
                        else:
                            try:
                                raw_content = file_bytes.decode("utf-8")
                            except UnicodeDecodeError:
                                raw_content = file_bytes.decode("latin-1")

                        extracted_files, zip_path = python_smart_split_md(raw_content, base_name, max_chunk_size=5000)

                        st.session_state.split_ui_files = extracted_files
                        st.session_state.zip_ui_path = zip_path
                        st.session_state.step_ui = 2
                        status.update(label=f"✅ Đã cắt thành công {len(extracted_files)} phần nhỏ!", state="complete")
                    except Exception as e:
                        status.update(label="❌ Lỗi khi bóc tách file!", state="error")
                        st.error(str(e))

            if "split_ui_files" in st.session_state and st.session_state.split_ui_files:
                st.success(f"🎉 Đã hoàn tất bóc tách thành {len(st.session_state.split_ui_files)} phần nhỏ.")
                
                with open(st.session_state.zip_ui_path, "rb") as zf:
                    st.download_button(
                        label="📥 Tải xuống trọn bộ file Markdown con (.zip)",
                        data=zf,
                        file_name=os.path.basename(st.session_state.zip_ui_path),
                        mime="application/zip",
                        key="dl_split_ui_zip_main"
                    )

                st.write("📌 **Danh sách các file .md đã cắt theo từng Màn hình / Part:**")
                for item in st.session_state.split_ui_files:
                    with st.expander(f"📄 **{item.get('screen_name', item.get('file_name'))}** ({item['char_count']} chars)"):
                        st.code(item['content'], language="markdown")

                st.markdown("---")
                st.write("### 🤖 Bước 2: Agent 2 - Sinh Test Case UI Chi Tiết Cho Từng Màn Hình")
                
                if st.button("🚀 Kích hoạt Agent 2 (Chạy Multi-Agent Batching UI)", type="primary", key="btn_agent2_ui_main"):
                    if not api_key:
                        st.error("⚠️ Vui lòng nhập API Key ở thanh bên trái!")
                    else:
                        ui_items = st.session_state.split_ui_files
                        progress_bar = st.progress(0)
                        status_text = st.empty()
                        full_parsed_tree = []

                        for index, item in enumerate(ui_items):
                            screen_name = item.get('screen_name', f"Screen_{index+1}")
                            screen_content = item['content']
                            
                            status_text.info(f"⏳ [{index + 1}/{len(ui_items)}] Agent 2 đang phân tích: **{screen_name}** ({len(screen_content)} chars)...")
                            
                            ok, screen_result = call_qwen_max_agent(
                                content=screen_content,
                                api_key=api_key,
                                base_url=base_url,
                                model=model_name,
                                prompt_template=PROMPT_WEB_UI,
                                max_tokens=16384
                            )

                            if ok and screen_result.strip():
                                full_parsed_tree.append(screen_result.strip())
                            else:
                                st.warning(f"⚠️ Phần '{screen_name}' bị lỗi/timeout, bỏ qua.")

                            progress_bar.progress((index + 1) / len(ui_items))

                        complete_tree_text = "\n\n".join(full_parsed_tree)
                        st.session_state.parsed_tree_ui = complete_tree_text
                        st.session_state.step_ui = 3
                        status_text.success(f"✅ Hoàn tất sinh Test Case UI cho {len(ui_items)} phần!")

        # LUỒNG 2: INTERACTIVE REVIEW PLAN
        else:
            st.write("### 🤖 Bước 1: Agent 1 - Bóc Tách Requirement & Lập Rule Matrix Dạng JSON")
            
            if st.button("🚀 Kích hoạt Agent 1 (Tạo JSON Rule Matrix)", key="btn_agent1_interactive"):
                if not api_key:
                    st.error("⚠️ Vui lòng nhập API Key!")
                else:
                    with st.spinner("Agent 1 đang đọc toàn bộ Spec và áp dụng QA Test Design Rules..."):
                        raw_ui_text = extract_text_from_file(uploaded_ui_file)
                        ok, json_res = call_qwen_max_agent(
                            content=raw_ui_text,
                            api_key=api_key,
                            base_url=base_url,
                            model=model_name,
                            prompt_template=PROMPT_AGENT1_EXTRACT_RULE_MATRIX,
                            max_tokens=16384
                        )
                        
                        if ok:
                            st.session_state.raw_agent1_response = json_res
                            parsed_ok, parsed_json, parse_diag = extract_json_from_model_response(json_res)

                            if parsed_ok:
                                schema_ok, schema_diag = validate_rule_matrix_schema(parsed_json)
                                if schema_ok:
                                    st.session_state.interactive_scope_json = parsed_json
                                    st.success(f"✅ Agent 1 trả JSON hợp lệ. {schema_diag}")
                                else:
                                    log_error(f"[AGENT1] ❌ JSON hợp lệ nhưng sai Rule Matrix schema | {schema_diag}")
                                    st.error(f"❌ JSON hợp lệ nhưng sai schema Rule Matrix: {schema_diag}")
                                    st.session_state.raw_json_fallback = json_res
                                    with st.expander("🔎 Raw output Agent 1", expanded=False):
                                        st.code(json_res, language="text")
                            else:
                                log_error(f"[AGENT1] ❌ {parse_diag}")
                                st.error(f"❌ Agent 1 trả output nhưng parse JSON thất bại: {parse_diag}")
                                st.session_state.raw_json_fallback = json_res
                                with st.expander("🔎 Raw output Agent 1 để debug", expanded=True):
                                    st.code(json_res, language="text")
                                    st.download_button(
                                        "📥 Tải raw Agent 1 để debug",
                                        data=json_res,
                                        file_name="agent1_raw_output.txt",
                                        mime="text/plain",
                                        key="dl_agent1_raw_debug"
                                    )
                        else:
                            st.error("❌ Agent 1/API không trả được output.")
                            st.code(json_res, language="text")

            if "interactive_scope_json" in st.session_state:
                st.markdown("---")
                st.write("### 🔍 Bước 2: QA Lead Review Rule Matrix & Chọn Màn Hình")
                st.info("💡 Bạn có thể tích chọn Màn hình và chỉnh sửa trực tiếp từng Test Rule dưới dạng JSON. Agent 2 chỉ render Rule Matrix, không tự suy luận thêm.")

                plan_data = st.session_state.interactive_scope_json
                approved_screens = []

                for idx, screen in enumerate(plan_data.get("screens", [])):
                    with st.expander(f"📌 Màn hình {idx+1}: {screen.get('screen_name', 'Unnamed')}", expanded=True):
                        is_selected = st.checkbox(f"Đưa màn hình '{screen.get('screen_name')}' vào Scope Gen Test Case", value=True, key=f"chk_inter_{idx}")
                        
                        edited_json_str = st.text_area(
                            "TEST DESIGN / RULE MATRIX của màn hình (Sửa rule nếu cần):",
                            value=json.dumps(screen, ensure_ascii=False, indent=2),
                            height=300,
                            key=f"txt_json_inter_{idx}"
                        )
                        
                        if is_selected:
                            try:
                                updated_screen_json = json.loads(edited_json_str)
                                approved_screens.append(updated_screen_json)
                            except json.JSONDecodeError as e:
                                st.error(
                                    f"❌ JSON màn hình {idx+1} lỗi tại line {e.lineno}, col {e.colno}: {e.msg}"
                                )

                st.markdown("---")
                st.write("### 🤖 Bước 3: Agent 2 - Render Test Case XMind Từ Rule Matrix Đã Phê Duyệt")
                
                if st.button("🚀 Chốt Rule Matrix & Kích Hoạt Agent 2 Gen XMind", type="primary", key="btn_agent2_interactive"):
                    if not api_key:
                        st.error("⚠️ Vui lòng nhập API Key!")
                    elif not approved_screens:
                        st.error("⚠️ Bạn chưa chọn Màn hình nào!")
                    else:
                        with st.spinner("Agent 2 đang render từng Test Rule thành Test Case XMind..."):
                            final_approved_payload = json.dumps({"screens": approved_screens}, ensure_ascii=False)
                            
                            ok, tc_res = call_qwen_max_agent(
                                content=final_approved_payload,
                                api_key=api_key,
                                base_url=base_url,
                                model=model_name,
                                prompt_template=PROMPT_AGENT2_GEN_XMIND_FROM_RULE_MATRIX,
                                max_tokens=16384
                            )
                            
                            if ok:
                                st.session_state.parsed_tree_ui = tc_res
                                st.success("✅ Đã render xong toàn bộ Test Case UI theo Rule Matrix chốt!")
                            else:
                                st.error("❌ Lỗi khi sinh Test Case!")

        # BƯỚC XUẤT FILE CHUNG CHO CẢ 2 LUỒNG UI
        if "parsed_tree_ui" in st.session_state and st.session_state.parsed_tree_ui:
            st.markdown("---")
            st.write("### 📦 Xuất Kịch Bản Kiểm Thử (XMind / Excel)")
            
            output_xmind_path = os.path.join(tempfile.gettempdir(), f"{uploaded_ui_file.name.rsplit('.', 1)[0]}_UI_TestCases.xmind")
            ok_xmind, err_xmind = create_xmind_from_text(st.session_state.parsed_tree_ui, output_xmind_path, root_title="Bộ Test Case UI/UX & Web App")
            
            if ok_xmind:
                st.success("✅ Đã tạo thành công bộ Test Case cho Web/App!")
                
                with st.expander("🔍 Xem trước dạng Text cây thụt lề"):
                    st.code(st.session_state.parsed_tree_ui, language="text")
                    
                col_xmind, col_excel = st.columns(2)
                
                with col_xmind:
                    with open(output_xmind_path, "rb") as f:
                        st.download_button(
                            label="📥 Tải xuống sơ đồ UI (.XMind)",
                            data=f,
                            file_name=f"{uploaded_ui_file.name.rsplit('.', 1)[0]}_UI_TestCases.xmind",
                            mime="application/octet-stream",
                            key="dl_ui_main"
                        )
                
                with col_excel:
                    excel_bytes = convert_tree_to_ui_excel(st.session_state.parsed_tree_ui)
                    st.download_button(
                        label="📊 Tải xuống bảng Excel (.XLSX)",
                        data=excel_bytes,
                        file_name=f"{uploaded_ui_file.name.rsplit('.', 1)[0]}_UI_TestCases.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_ui_excel_main"
                    )
            else:
                st.error("❌ Không thể tạo file XMind. Chi tiết lỗi:")
                st.code(err_xmind, language="bash")

# --- TAB 2: API TESTING ---
with tab_api:
    st.subheader("📌 Kiểm thử RESTful API & Services")
    
    col1_api, col2_api = st.columns(2)
    with col1_api:
        api_design_file = st.file_uploader("1. Tài liệu Thiết kế API (Swagger / Postman / Spec)", type=["pdf", "md", "json", "yaml", "txt"], key="api_design_file")
    with col2_api:
        api_ba_file = st.file_uploader("2. Tài liệu Nghiệp vụ BA (nếu có)", type=["pdf", "md", "txt"], key="api_ba_file")

    if api_design_file or api_ba_file:
        st.markdown("---")
        st.write("### 🤖 Bước 1: Agent 1 - Đọc API Spec & Lập Plan Kiểm thử")
        
        if st.button("🚀 Kích hoạt Agent 1 (Phân tích API & Lập Plan)", key="btn_agent1_api"):
            if not api_key:
                st.error("⚠️ Vui lòng nhập API Key!")
            else:
                with st.spinner("Agent 1 đang phân tích danh sách Endpoints & Quy tắc nghiệp vụ..."):
                    api_design_text = extract_text_from_file(api_design_file)
                    api_ba_text = extract_text_from_file(api_ba_file)
                    
                    api_input = (
                        "=== API DESIGN / SPEC ===\n"
                        + (api_design_text if api_design_text else "Không cung cấp")
                        + "\n\n=== BA BUSINESS DOCUMENT ===\n"
                        + (api_ba_text if api_ba_text else "Không cung cấp")
                    )
                    prompt_plan = PROMPT_AGENT1_EXTRACT_API_JSON_SCOPE.replace("{content}", api_input)
                    
                    ok, plan_res = call_qwen_agent(prompt_plan, api_key, base_url, model_name)
                    if ok:
                        st.session_state.api_plan_code = plan_res
                        st.session_state.api_design_text = api_design_text
                        st.session_state.api_ba_text = api_ba_text
                        st.success("✅ Agent 1 đã phân tích API xong! Hãy review Plan bên dưới.")
                    else:
                        st.error("❌ Lỗi khi phân tích tài liệu API!")

        if "api_plan_code" in st.session_state:
            st.markdown("---")
            st.write("### 🔍 Bước 2: Review & Phê duyệt API Test Plan")
            
            edited_api_plan = st.text_area("Bản Kế hoạch API (Chỉnh sửa nếu cần):", value=st.session_state.api_plan_code, height=350, key="api_edited_plan")

            if st.button("✅ Phê duyệt Plan & Cho phép Agent 2 Sinh API Test Case XMind", type="primary", key="btn_approve_api_plan"):
                with st.spinner("Agent 2 đang sinh bộ Test Case API chi tiết..."):
                    full_prompt = PROMPT_API_SPEC.format(
                        plan_content=edited_api_plan,
                        design_doc=st.session_state.api_design_text,
                        ba_doc=st.session_state.api_ba_text
                    )
                    
                    ok, tc_res = call_qwen_agent(full_prompt, api_key, base_url, model_name, max_tokens=16384)
                    if ok:
                        st.session_state.api_parsed_tree = tc_res
                        st.success("✅ Đã sinh xong toàn bộ Test Case API!")
                    else:
                        st.error("❌ Lỗi khi sinh Test Case API!")

        if "api_parsed_tree" in st.session_state:
            st.markdown("---")
            st.write("### 📦 Bước 3: Xuất File Sơ Đồ Tư Duy (.xmind & .xlsx)")
            
            out_filename = "API_TestCases.xmind"
            output_xmind_path = os.path.join(tempfile.gettempdir(), out_filename)
            ok_xmind, err_xmind = create_xmind_from_text(st.session_state.api_parsed_tree, output_xmind_path, root_title="Bộ Test Case API Testing")
            
            if ok_xmind:
                st.success("✅ Đã tạo thành công sơ đồ XMind API!")
                with st.expander("🔍 Xem trước kết quả cây Test Case API"):
                    st.code(st.session_state.api_parsed_tree, language="text")
                    
                col_api_xmind, col_api_excel = st.columns(2)
                
                with col_api_xmind:
                    with open(output_xmind_path, "rb") as f:
                        st.download_button(
                            label="📥 Tải xuống File API .XMind (Tiếng Việt)",
                            data=f,
                            file_name=out_filename,
                            mime="application/octet-stream",
                            key="dl_api_xmind"
                        )
                        
                with col_api_excel:
                    api_excel_bytes = convert_tree_to_ui_excel(st.session_state.api_parsed_tree)
                    st.download_button(
                        label="📊 Tải xuống File API Excel (.XLSX)",
                        data=api_excel_bytes,
                        file_name="API_TestCases.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_api_excel"
                    )
            else:
                st.error(f"❌ Lỗi khi đóng gói XMind: {err_xmind}")