# -*- coding: utf-8 -*-
import json
import io
import os
import re
import sys
import tempfile
import time
import traceback
import zipfile
import hashlib
import copy
import unicodedata
import yaml
from dataclasses import dataclass, asdict
from collections import OrderedDict, Counter
from pypdf import PdfReader
from datetime import datetime
import streamlit as st
from openai import OpenAI
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==========================================
# 0. CẤU HÌNH PIPELINE V3.1 — SCALABLE LONG DOCUMENT
# ==========================================
# Agent 1: chia tài liệu nguyên văn thành các semantic chunk nhỏ, có context ở biên.
AGENT1_CHUNK_TARGET_CHARS = 12000
AGENT1_CHUNK_MAX_CHARS = 15000
AGENT1_CONTEXT_CHARS = 1200
AGENT1_MIN_RECURSIVE_CHARS = 2500
AGENT1_MAX_RECURSION_DEPTH = 5
AGENT1_API_RETRIES = 1

# Agent 2: render Rule Matrix theo batch rule để không chạm output limit.
AGENT2_BATCH_MAX_RULES = 18
AGENT2_BATCH_MAX_INPUT_CHARS = 18000
AGENT2_MAX_RECURSION_DEPTH = 5
AGENT2_API_RETRIES = 1

# Qwen output budget. Nếu vẫn chạm length, pipeline sẽ tự chia nhỏ và retry.
QWEN_MAX_OUTPUT_TOKENS = 16384

# API Agent 1/2: dùng pipeline riêng nhưng cùng nguyên tắc scalable + recursive split.
API_AGENT1_CHUNK_TARGET_CHARS = 10000
API_AGENT1_CHUNK_MAX_CHARS = 13000
API_AGENT1_CONTEXT_CHARS = 1200
API_AGENT1_MIN_RECURSIVE_CHARS = 2200
API_AGENT1_MAX_RECURSION_DEPTH = 6
API_AGENT1_API_RETRIES = 1

API_AGENT2_BATCH_MAX_RULES = 16
API_AGENT2_BATCH_MAX_INPUT_CHARS = 17000
API_AGENT2_MAX_RECURSION_DEPTH = 6
API_AGENT2_API_RETRIES = 1
API_ENDPOINT_HINT_LIMIT = 8

# ==========================================
# 1. CẤU HÌNH TRANG STREAMLIT & SIDEBAR
# ==========================================
st.set_page_config(
    page_title="Chuyển đổi Tài liệu sang XMind Test Case (Multi-Agent)",
    page_icon="🧠",
    layout="wide"
)

st.title("🧠 Công cụ Tự động Tạo Test Case XMind từ Tài liệu (Kiến trúc Multi-Agent)")
st.caption("Giải pháp Mã Nguồn Mở dành riêng cho QA Lead: Phân tích tài liệu/Figma, Review JSON Scope Plan và viết bộ Test Case chi tiết ra sơ đồ tư duy (.xmind)")

with st.sidebar:
    st.header("⚙️ Cấu hình Qwen API (DashScope)")
    api_key = st.text_input("Nhập API Key:", type="password", key="sidebar_api_key")
    
    base_url = st.text_input(
        "Base URL Endpoint:", 
        value="https://token-plan.ap-southeast-1.maas.aliyuncs.com/compatible-mode/v1",
        key="sidebar_base_url"
    )
    
    model_name = st.selectbox(
        "Chọn Model cho Agent 1/2:", 
        ["qwen-max", "qwen3.8-max", "qwen-plus"], 
        index=0,
        key="sidebar_model_name"
    )

    st.divider()
    st.markdown("### 🔄 Trạng thái Tiến trình")
    if "step_ui" not in st.session_state:
        st.session_state.step_ui = 1
    if "step_api" not in st.session_state:
        st.session_state.step_api = 1

    if st.button("🔴 Reset/Chạy lại từ đầu", key="sidebar_reset_btn"):
        st.session_state.clear()
        st.rerun()

# ==========================================
# 2. KHAI BÁO CÁC HÀM XỬ LÝ (HELPER FUNCTIONS)
# ==========================================
def log_info(msg: str):
    timestamp = datetime.now().strftime("%H:%M:%S")
    log_msg = f"[LOG {timestamp}] ℹ️ {msg}"
    print(log_msg)
    sys.stdout.flush()

def log_error(msg: str, exc: Exception = None):
    timestamp = datetime.now().strftime("%H:%M:%S")
    err_msg = f"[ERROR {timestamp}] ❌ {msg}"
    if exc:
        err_msg += f"\nDetails: {str(exc)}\n{traceback.format_exc()}"
    print(err_msg)
    sys.stdout.flush()
    return err_msg

def extract_text_from_file(uploaded_file) -> str:
    if uploaded_file is None:
        return ""
    file_bytes = uploaded_file.read()
    filename = uploaded_file.name.lower()
    
    if filename.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(file_bytes))
        full_pdf_text = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                full_pdf_text.append(f"--- TRANG {idx+1} ---\n{page_text}")
        return "\n\n".join(full_pdf_text)
    else:
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1")

def _log_response_summary(agent_name: str, elapsed: float, result: str, chunk_count: int, finish_reason: str | None, reasoning_chunk_count: int = 0):
    """Log chẩn đoán đầy đủ cho request Qwen mà không dump toàn bộ output ra console."""
    log_info(
        f"[{agent_name}] ✅ Request kết thúc | "
        f"Time={elapsed:.2f}s | Chunks={chunk_count} | "
        f"ReasoningChunks={reasoning_chunk_count} | "
        f"OutputChars={len(result):,} | "
        f"FinishReason={finish_reason or 'UNKNOWN'}"
    )


def call_qwen_agent(prompt: str, api_key: str, base_url: str, model: str, max_tokens: int = 16384) -> tuple[bool, str]:
    """Generic Qwen caller cho API/flow cũ. Có log finish_reason + kích thước output."""
    client = OpenAI(
        api_key=api_key,
        base_url=base_url,
        timeout=3600.0
    )
    start_time = time.time()
    log_info(
        f"[QWEN_AGENT] 🚀 Request | Model={model} | MaxTokens={max_tokens} | "
        f"PromptChars={len(prompt):,}"
    )

    try:
        response = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": prompt}],
            temperature=0.1,
            max_tokens=max_tokens,
            stream=True
        )

        full_response = []
        chunk_count = 0
        reasoning_chunk_count = 0
        finish_reason = None
        last_log_time = time.time()

        for chunk in response:
            if not chunk.choices:
                continue

            choice = chunk.choices[0]
            delta = choice.delta

            if getattr(delta, "reasoning_content", None):
                reasoning_chunk_count += 1

            if delta and delta.content:
                full_response.append(delta.content)
                chunk_count += 1

            if choice.finish_reason:
                finish_reason = choice.finish_reason

            now = time.time()
            if now - last_log_time >= 10:
                elapsed = int(now - start_time)
                log_info(
                    f"[QWEN_AGENT] ⏳ Streaming | Elapsed={elapsed}s | "
                    f"ContentChunks={chunk_count} | ReasoningChunks={reasoning_chunk_count}"
                )
                last_log_time = now

        elapsed_total = time.time() - start_time
        result = "".join(full_response)
        _log_response_summary("QWEN_AGENT", elapsed_total, result, chunk_count, finish_reason, reasoning_chunk_count)
        return True, result

    except Exception as e:
        elapsed_total = time.time() - start_time
        err_details = log_error(
            f"[QWEN_AGENT] ❌ Request failed after {elapsed_total:.2f}s | "
            f"Model={model} | MaxTokens={max_tokens}",
            e
        )
        return False, err_details


@dataclass
class QwenCallResult:
    ok: bool
    text: str
    finish_reason: str | None
    elapsed: float
    content_chunks: int
    reasoning_chunks: int
    error: str | None = None

    @property
    def complete(self) -> bool:
        return self.ok and self.finish_reason != "length" and bool(self.text.strip())


def call_qwen_max_agent_detailed(
    content: str,
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    max_tokens: int = QWEN_MAX_OUTPUT_TOKENS,
    agent_name: str = "QWEN_MAX_AGENT",
) -> QwenCallResult:
    """Qwen caller có metadata đầy đủ để pipeline tự xử lý length/retry/split.

    Tầng transport KHÔNG parse JSON và KHÔNG tự cứu output bị truncate.
    """
    client = OpenAI(
        api_key=api_key,
        base_url=base_url,
        timeout=3600.0
    )

    # Chỉ thay placeholder {content}; prompt có nhiều JSON literal với { }.
    full_prompt = prompt_template.replace("{content}", content)
    start_time = time.time()
    log_info(
        f"[{agent_name}] 🚀 Request | Model={model} | MaxTokens={max_tokens} | "
        f"PromptChars={len(full_prompt):,} | ContentChars={len(content):,}"
    )

    try:
        response = client.chat.completions.create(
            model=model,
            messages=[{"role": "user", "content": full_prompt}],
            temperature=0.1,
            max_tokens=max_tokens,
            stream=True
        )

        full_response = []
        chunk_count = 0
        reasoning_chunk_count = 0
        finish_reason = None
        last_log_time = time.time()

        for chunk in response:
            if not chunk.choices:
                continue

            choice = chunk.choices[0]
            delta = choice.delta

            if getattr(delta, "reasoning_content", None):
                reasoning_chunk_count += 1

            if delta and delta.content:
                full_response.append(delta.content)
                chunk_count += 1

            if choice.finish_reason:
                finish_reason = choice.finish_reason

            now = time.time()
            if now - last_log_time >= 10:
                elapsed = int(now - start_time)
                log_info(
                    f"[{agent_name}] ⏳ Streaming | Elapsed={elapsed}s | "
                    f"ContentChunks={chunk_count} | ReasoningChunks={reasoning_chunk_count}"
                )
                last_log_time = now

        elapsed_total = time.time() - start_time
        result = "".join(full_response)
        _log_response_summary(
            agent_name,
            elapsed_total,
            result,
            chunk_count,
            finish_reason,
            reasoning_chunk_count
        )

        if finish_reason == "length":
            error = (
                f"OUTPUT BỊ CẮT DO MAX TOKENS | OutputChars={len(result):,} | "
                f"MaxTokens={max_tokens}"
            )
            log_error(f"[{agent_name}] ❌ {error}")
            return QwenCallResult(
                ok=False,
                text=result,
                finish_reason=finish_reason,
                elapsed=elapsed_total,
                content_chunks=chunk_count,
                reasoning_chunks=reasoning_chunk_count,
                error=error,
            )

        if not result.strip():
            error = (
                f"Request kết thúc nhưng Content rỗng | "
                f"FinishReason={finish_reason or 'UNKNOWN'} | "
                f"ReasoningChunks={reasoning_chunk_count}"
            )
            log_error(f"[{agent_name}] ⚠️ {error}")
            return QwenCallResult(
                ok=False,
                text=result,
                finish_reason=finish_reason,
                elapsed=elapsed_total,
                content_chunks=chunk_count,
                reasoning_chunks=reasoning_chunk_count,
                error=error,
            )

        return QwenCallResult(
            ok=True,
            text=result,
            finish_reason=finish_reason,
            elapsed=elapsed_total,
            content_chunks=chunk_count,
            reasoning_chunks=reasoning_chunk_count,
            error=None,
        )

    except Exception as e:
        elapsed_total = time.time() - start_time
        err_details = log_error(
            f"[{agent_name}] ❌ Request failed/timeout after {elapsed_total:.2f}s | "
            f"Model={model} | MaxTokens={max_tokens}",
            e
        )
        return QwenCallResult(
            ok=False,
            text="",
            finish_reason=None,
            elapsed=elapsed_total,
            content_chunks=0,
            reasoning_chunks=0,
            error=err_details,
        )


def call_qwen_max_agent(
    content: str,
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    max_tokens: int = QWEN_MAX_OUTPUT_TOKENS,
) -> tuple[bool, str]:
    """Compatibility wrapper cho các flow cũ.

    Với finish_reason=length trả False + partial raw output để UI có thể debug,
    tuyệt đối không coi JSON truncate là response thành công.
    """
    result = call_qwen_max_agent_detailed(
        content=content,
        api_key=api_key,
        base_url=base_url,
        model=model,
        prompt_template=prompt_template,
        max_tokens=max_tokens,
        agent_name="QWEN_MAX_AGENT",
    )
    if result.complete:
        return True, result.text
    return False, result.text or (result.error or "Qwen request failed")


def extract_json_from_model_response(raw_text: str) -> tuple[bool, dict | list | None, str]:
    """Parse JSON root an toàn từ model response.

    KHÔNG scan các object con. Nếu root JSON bị truncate, phải báo parse fail thật sự,
    tránh nhặt nhầm một test_rule con rồi báo sai schema.
    """
    if raw_text is None:
        return False, None, "Raw response = None"

    text = raw_text.strip()
    if not text:
        return False, None, "Raw response rỗng"

    # Ưu tiên code fence hoàn chỉnh nếu model vẫn trả fence.
    fenced = re.search(r"```(?:json)?\s*(.*?)\s*```", text, flags=re.IGNORECASE | re.DOTALL)
    candidate = fenced.group(1).strip() if fenced else text

    try:
        return True, json.loads(candidate), "JSON parse OK (direct)"
    except json.JSONDecodeError as direct_err:
        direct_error = direct_err

    # Chỉ thử ROOT JSON đầu tiên để hỗ trợ text thừa trước JSON.
    decoder = json.JSONDecoder()
    root_obj = candidate.find("{")
    root_arr = candidate.find("[")
    root_positions = [p for p in (root_obj, root_arr) if p >= 0]

    if root_positions:
        root_start = min(root_positions)
        try:
            parsed, end = decoder.raw_decode(candidate[root_start:])
            trailing = candidate[root_start + end:].strip()
            diagnostic = "JSON parse OK (root raw_decode)"
            if trailing:
                diagnostic += f" | Ignored trailing chars={len(trailing)}"
            return True, parsed, diagnostic
        except json.JSONDecodeError:
            pass

    pos = getattr(direct_error, "pos", None)
    if pos is not None:
        line = getattr(direct_error, "lineno", None)
        col = getattr(direct_error, "colno", None)
        start = max(0, pos - 180)
        end = min(len(candidate), pos + 180)
        preview = candidate[start:end].replace("\n", "\\n")
        diagnostic = (
            f"JSON parse FAILED | {direct_error.msg} | "
            f"pos={pos} | line={line} | col={col} | "
            f"candidate_chars={len(candidate):,} | "
            f"context={preview!r}"
        )
    else:
        diagnostic = f"JSON parse FAILED | candidate_chars={len(candidate):,}"

    return False, None, diagnostic


def validate_rule_matrix_schema(data, strict: bool = False) -> tuple[bool, str]:
    """Kiểm tra schema Rule Matrix.

    strict=False: đủ cho chunk output.
    strict=True: dùng cho Final Rule Matrix sau merge/renumber.
    """
    if not isinstance(data, dict):
        return False, f"Root phải là object, nhận {type(data).__name__}"

    screens = data.get("screens")
    if not isinstance(screens, list):
        return False, "Thiếu field 'screens' hoặc 'screens' không phải array"

    required_rule_fields = {
        "rule_id", "target", "category", "rule_type", "rule_name",
        "test_objective", "source_requirement", "applied_qa_rule", "generation_reason"
    }
    allowed_categories = {"UI", "VALIDATION", "ACTION", "GRID", "POPUP", "BUSINESS_FLOW", "EXCEPTION"}
    allowed_rule_types = {"EXPLICIT", "DERIVED"}

    total_rules = 0
    rule_ids = set()

    for idx, screen in enumerate(screens, start=1):
        if not isinstance(screen, dict):
            return False, f"screens[{idx-1}] phải là object"

        screen_name = screen.get("screen_name")
        if strict and (not isinstance(screen_name, str) or not screen_name.strip()):
            return False, f"screens[{idx-1}].screen_name rỗng/không hợp lệ"

        if not isinstance(screen.get("test_rules"), list):
            return False, f"screens[{idx-1}] thiếu 'test_rules' hoặc không phải array"

        for ridx, rule in enumerate(screen["test_rules"], start=1):
            total_rules += 1
            if not isinstance(rule, dict):
                return False, f"screens[{idx-1}].test_rules[{ridx-1}] phải là object"

            missing = sorted(required_rule_fields - set(rule.keys()))
            if missing:
                return False, f"Rule {idx}.{ridx} thiếu field: {', '.join(missing)}"

            if rule.get("category") not in allowed_categories:
                return False, f"Rule {idx}.{ridx} category không hợp lệ: {rule.get('category')!r}"

            if rule.get("rule_type") not in allowed_rule_types:
                return False, f"Rule {idx}.{ridx} rule_type không hợp lệ: {rule.get('rule_type')!r}"

            if not str(rule.get("source_requirement", "")).strip():
                return False, f"Rule {idx}.{ridx} source_requirement rỗng"

            if rule.get("rule_type") == "DERIVED" and not str(rule.get("applied_qa_rule", "")).strip():
                return False, f"Rule {idx}.{ridx} DERIVED nhưng applied_qa_rule rỗng"

            if strict:
                rule_id = str(rule.get("rule_id", "")).strip()
                if not rule_id:
                    return False, f"Rule {idx}.{ridx} rule_id rỗng"
                if rule_id in rule_ids:
                    return False, f"Duplicate rule_id: {rule_id}"
                rule_ids.add(rule_id)

    return True, f"Schema OK | Screens={len(screens)} | TestRules={total_rules}"


def python_smart_split_md(raw_text: str, base_filename: str, max_chunk_size: int = 5000) -> tuple[list[dict], str]:
    sections = []
    heading_pattern = re.compile(r'^(#{1,3}\s+.+)$', re.MULTILINE)
    splits = heading_pattern.split(raw_text)

    if len(splits) > 1:
        current_title = f"{base_filename}_Part_1"
        current_content = ""

        for element in splits:
            if heading_pattern.match(element):
                if current_content.strip():
                    sections.append({"title": current_title, "content": current_content.strip()})
                current_title = re.sub(r'^[#\s]+', '', element).strip()
                current_content = element + "\n"
            else:
                current_content += element

        if current_content.strip():
            sections.append({"title": current_title, "content": current_content.strip()})
    else:
        sections.append({"title": base_filename, "content": raw_text})

    final_chunks = []
    for sec in sections:
        content = sec["content"]
        title = sec["title"]
        
        if len(content) > max_chunk_size:
            for i in range(0, len(content), max_chunk_size):
                part_text = content[i:i+max_chunk_size]
                part_title = f"{title}_Part_{i//max_chunk_size + 1}"
                final_chunks.append({"title": part_title, "content": part_text})
        else:
            final_chunks.append({"title": title, "content": content})

    temp_dir = tempfile.mkdtemp()
    extracted_files = []
    zip_filename = f"{base_filename}_MD_Parsed.zip"
    zip_path = os.path.join(tempfile.gettempdir(), zip_filename)

    with zipfile.ZipFile(zip_path, 'w', zipfile.ZIP_DEFLATED) as zip_out:
        for idx, item in enumerate(final_chunks):
            safe_title = re.sub(r'[\\/*?:"<>|]', '_', item["title"])
            sub_file_name = f"{idx+1:02d}_{safe_title}.md"
            file_path = os.path.join(temp_dir, sub_file_name)

            with open(file_path, "w", encoding="utf-8") as f:
                f.write(item["content"])

            zip_out.write(file_path, arcname=sub_file_name)

            extracted_files.append({
                "screen_name": item["title"],
                "file_name": sub_file_name,
                "file_path": file_path,
                "content": item["content"],
                "char_count": len(item["content"])
            })

    return extracted_files, zip_path


# ==========================================
# 2.1 PIPELINE V3.1 — LONG DOCUMENT BATCHING
# ==========================================

def _normalize_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", str(value or ""))
    value = re.sub(r"\s+", " ", value).strip().casefold()
    return value


def _is_heading_line(line: str) -> bool:
    s = line.strip()
    if not s:
        return False
    if re.match(r"^#{1,6}\s+\S", s):
        return True
    if re.match(r"^\d+(?:\.\d+){1,6}\s+\S", s):
        return True
    if re.match(r"^\d+\.\s+[A-Za-zÀ-ỹ]", s):
        return True
    # Các heading thường gặp trong tài liệu BA/RSD.
    if re.match(
        r"^(màn hình|mh\d*\s*:|logic\s+|luồng\s+|quy tắc nghiệp vụ|mô tả màn hình|"
        r"đặc tả|tóm tắt usecase|ma trận phân quyền|các kết nối|sơ đồ luồng|mockup|"
        r"===\s*source document|api\s+design|api\s+spec|ba\s+business|"
        r"get\s+/|post\s+/|put\s+/|patch\s+/|delete\s+/|head\s+/|options\s+/)",
        s,
        flags=re.IGNORECASE,
    ):
        return True
    return False


def _safe_split_large_block(text: str, max_chars: int) -> list[str]:
    """Fallback split cho một block quá lớn; ưu tiên biên newline/<br>/câu hơn hard-cut."""
    text = text.strip()
    if len(text) <= max_chars:
        return [text] if text else []

    parts = []
    remaining = text
    min_cut = max(800, int(max_chars * 0.55))

    while len(remaining) > max_chars:
        window = remaining[:max_chars]
        candidates = [
            window.rfind("\n"),
            window.lower().rfind("<br>"),
            window.rfind(". "),
            window.rfind("; "),
            window.rfind(", "),
            window.rfind(" "),
        ]
        cut = max(candidates)
        if cut < min_cut:
            cut = max_chars
        else:
            # giữ delimiter ở phần trước nếu là <br>
            if window[max(0, cut-3):cut+1].lower().endswith("<br>"):
                cut += 1
        part = remaining[:cut].strip()
        if part:
            parts.append(part)
        remaining = remaining[cut:].strip()

    if remaining:
        parts.append(remaining)
    return parts


def _document_to_semantic_blocks(raw_text: str, base_title: str) -> list[dict]:
    """Tách raw document thành block theo heading / table row / paragraph, không dùng LLM."""
    lines = raw_text.splitlines()
    blocks = []
    paragraph = []
    current_title = base_title

    def flush_paragraph():
        nonlocal paragraph
        if paragraph:
            block_text = "\n".join(paragraph).strip()
            if block_text:
                blocks.append({"title": current_title, "text": block_text})
            paragraph = []

    for line in lines:
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            continue

        if _is_heading_line(line):
            flush_paragraph()
            current_title = re.sub(r"^[#\s]+", "", stripped).strip() or current_title
            blocks.append({"title": current_title, "text": line.rstrip()})
            continue

        # Markdown table: giữ từng row là block để tránh chém giữa một requirement row.
        if stripped.startswith("|"):
            flush_paragraph()
            blocks.append({"title": current_title, "text": line.rstrip()})
            continue

        paragraph.append(line.rstrip())

    flush_paragraph()

    if not blocks and raw_text.strip():
        blocks = [{"title": base_title, "text": raw_text.strip()}]

    # Nếu một block đơn vẫn quá lớn thì fallback split an toàn.
    expanded = []
    for block in blocks:
        if len(block["text"]) <= AGENT1_CHUNK_MAX_CHARS:
            expanded.append(block)
        else:
            for part in _safe_split_large_block(block["text"], AGENT1_CHUNK_MAX_CHARS):
                expanded.append({"title": block["title"], "text": part})
    return expanded


def split_document_semantic(
    raw_text: str,
    base_title: str,
    target_chars: int = AGENT1_CHUNK_TARGET_CHARS,
    max_chars: int = AGENT1_CHUNK_MAX_CHARS,
    context_chars: int = AGENT1_CONTEXT_CHARS,
) -> list[dict]:
    """Pack semantic blocks thành chunk động.

    Mỗi chunk có CURRENT SOURCE riêng + context trước/sau để giữ dependency ở biên.
    Context chỉ dùng để hiểu, prompt cấm sinh rule chỉ từ context.
    """
    blocks = _document_to_semantic_blocks(raw_text, base_title)
    core_chunks = []
    current_parts = []
    current_len = 0
    current_title = base_title

    def flush_current():
        nonlocal current_parts, current_len, current_title
        if not current_parts:
            return
        core = "\n\n".join(current_parts).strip()
        if core:
            core_chunks.append({
                "title": current_title,
                "core_text": core,
                "char_count": len(core),
            })
        current_parts = []
        current_len = 0

    for block in blocks:
        block_text = block["text"].strip()
        if not block_text:
            continue
        block_len = len(block_text) + (2 if current_parts else 0)

        if not current_parts:
            current_title = block.get("title") or base_title
            current_parts = [block_text]
            current_len = len(block_text)
            continue

        proposed = current_len + block_len
        # target là điểm flush ưu tiên; max là hard ceiling.
        if proposed > max_chars or (current_len >= target_chars and proposed > target_chars):
            flush_current()
            current_title = block.get("title") or base_title
            current_parts = [block_text]
            current_len = len(block_text)
        else:
            current_parts.append(block_text)
            current_len = proposed
            if _is_heading_line(block_text.splitlines()[0] if block_text.splitlines() else ""):
                current_title = block.get("title") or current_title

    flush_current()

    chunks = []
    for idx, item in enumerate(core_chunks):
        prev_context = core_chunks[idx - 1]["core_text"][-context_chars:] if idx > 0 else ""
        next_context = core_chunks[idx + 1]["core_text"][:context_chars] if idx + 1 < len(core_chunks) else ""
        chunks.append({
            "chunk_id": f"C{idx + 1:03d}",
            "title": item["title"],
            "core_text": item["core_text"],
            "prev_context": prev_context,
            "next_context": next_context,
            "char_count": item["char_count"],
            "depth": 0,
        })

    return chunks


def build_agent1_chunk_payload(chunk: dict) -> str:
    """Đóng gói raw chunk + boundary context. Metadata không phải requirement."""
    return f"""=== CHUNK METADATA — KHÔNG PHẢI REQUIREMENT ===
chunk_id: {chunk.get('chunk_id')}
source_section_hint: {chunk.get('title')}
chunk_depth: {chunk.get('depth', 0)}

=== PREVIOUS CONTEXT — CHỈ DÙNG ĐỂ HIỂU, KHÔNG SINH RULE CHỈ TỪ ĐOẠN NÀY ===
{chunk.get('prev_context', '')}

=== CURRENT SOURCE — PHẠM VI CHÍNH ĐƯỢC PHÉP SINH TEST RULE ===
{chunk.get('core_text', '')}

=== NEXT CONTEXT — CHỈ DÙNG ĐỂ HIỂU, KHÔNG SINH RULE CHỈ TỪ ĐOẠN NÀY ===
{chunk.get('next_context', '')}
"""


def _cache_key(prefix: str, model: str, prompt_template: str, payload: str) -> str:
    material = f"{prefix}\n{model}\n{prompt_template}\n{payload}".encode("utf-8", errors="ignore")
    return hashlib.sha256(material).hexdigest()


def _split_chunk_for_retry(chunk: dict) -> list[dict]:
    """Recursive split semantic cho một chunk bị length/parse/schema fail."""
    core = chunk.get("core_text", "")
    if len(core) <= AGENT1_MIN_RECURSIVE_CHARS:
        return []

    target = max(AGENT1_MIN_RECURSIVE_CHARS, min(7000, len(core) // 2))
    max_chars = max(target + 800, min(8500, len(core)))
    local = split_document_semantic(
        core,
        base_title=chunk.get("title", "RetryChunk"),
        target_chars=target,
        max_chars=max_chars,
        context_chars=min(AGENT1_CONTEXT_CHARS, 800),
    )

    if len(local) <= 1:
        # Last-resort 50/50 split tại whitespace gần giữa.
        mid = len(core) // 2
        left_cut = core.rfind("\n", 0, mid)
        if left_cut < AGENT1_MIN_RECURSIVE_CHARS:
            left_cut = core.rfind(" ", 0, mid)
        if left_cut < AGENT1_MIN_RECURSIVE_CHARS:
            left_cut = mid
        parts = [core[:left_cut].strip(), core[left_cut:].strip()]
        local = []
        for p in parts:
            if p:
                local.append({
                    "title": chunk.get("title", "RetryChunk"),
                    "core_text": p,
                    "prev_context": "",
                    "next_context": "",
                    "char_count": len(p),
                    "depth": chunk.get("depth", 0) + 1,
                })

    parent_prev = chunk.get("prev_context", "")
    parent_next = chunk.get("next_context", "")
    children = []
    parent_id = chunk.get("chunk_id", "C")
    for idx, child in enumerate(local):
        child = dict(child)
        child["chunk_id"] = f"{parent_id}.{idx + 1}"
        child["depth"] = chunk.get("depth", 0) + 1
        if idx == 0 and parent_prev:
            child["prev_context"] = (parent_prev + "\n" + child.get("prev_context", ""))[-AGENT1_CONTEXT_CHARS:]
        if idx == len(local) - 1 and parent_next:
            child["next_context"] = (child.get("next_context", "") + "\n" + parent_next)[:AGENT1_CONTEXT_CHARS]
        children.append(child)
    return children


def process_agent1_chunk_recursive(
    chunk: dict,
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    cache: dict | None = None,
    diagnostics: list | None = None,
) -> tuple[bool, list[dict]]:
    """Agent 1 leaf processor.

    length hoặc JSON/schema fail do complexity => tự chia nhỏ và retry.
    Không bỏ qua chunk lỗi; nếu leaf vẫn fail, toàn pipeline fail để bảo toàn coverage.
    """
    diagnostics = diagnostics if diagnostics is not None else []
    payload = build_agent1_chunk_payload(chunk)
    key = _cache_key("agent1-v3.1", model, prompt_template, payload)

    if cache is not None and key in cache:
        diagnostics.append({
            "chunk_id": chunk.get("chunk_id"),
            "chars": len(chunk.get("core_text", "")),
            "depth": chunk.get("depth", 0),
            "status": "CACHE_HIT",
        })
        return True, [copy.deepcopy(cache[key])]

    call_result = None
    for attempt in range(AGENT1_API_RETRIES + 1):
        agent_name = f"AGENT1/{chunk.get('chunk_id')}"
        call_result = call_qwen_max_agent_detailed(
            content=payload,
            api_key=api_key,
            base_url=base_url,
            model=model,
            prompt_template=prompt_template,
            max_tokens=QWEN_MAX_OUTPUT_TOKENS,
            agent_name=agent_name,
        )
        if call_result.ok or call_result.finish_reason == "length":
            break
        if attempt < AGENT1_API_RETRIES:
            log_info(f"[{agent_name}] 🔁 Retry API lần {attempt + 2}/{AGENT1_API_RETRIES + 1}")
            time.sleep(2)

    assert call_result is not None

    reason_to_split = None
    parse_diag = None
    schema_diag = None

    if call_result.finish_reason == "length":
        reason_to_split = "MAX_TOKENS"
    elif not call_result.ok:
        diagnostics.append({
            "chunk_id": chunk.get("chunk_id"),
            "chars": len(chunk.get("core_text", "")),
            "depth": chunk.get("depth", 0),
            "status": "API_FAILED",
            "error": call_result.error,
        })
        return False, []
    else:
        parsed_ok, parsed_json, parse_diag = extract_json_from_model_response(call_result.text)
        if not parsed_ok:
            reason_to_split = "JSON_PARSE_FAIL"
        else:
            schema_ok, schema_diag = validate_rule_matrix_schema(parsed_json, strict=False)
            if not schema_ok:
                reason_to_split = "SCHEMA_FAIL"
            else:
                diagnostics.append({
                    "chunk_id": chunk.get("chunk_id"),
                    "chars": len(chunk.get("core_text", "")),
                    "depth": chunk.get("depth", 0),
                    "status": "OK",
                    "elapsed": round(call_result.elapsed, 2),
                    "finish_reason": call_result.finish_reason,
                    "output_chars": len(call_result.text),
                    "screens": len(parsed_json.get("screens", [])),
                    "rules": sum(len(s.get("test_rules", [])) for s in parsed_json.get("screens", [])),
                })
                if cache is not None:
                    cache[key] = copy.deepcopy(parsed_json)
                return True, [parsed_json]

    depth = int(chunk.get("depth", 0))
    can_split = depth < AGENT1_MAX_RECURSION_DEPTH and len(chunk.get("core_text", "")) > AGENT1_MIN_RECURSIVE_CHARS

    diagnostics.append({
        "chunk_id": chunk.get("chunk_id"),
        "chars": len(chunk.get("core_text", "")),
        "depth": depth,
        "status": "SPLIT_RETRY" if can_split else "FAILED_LEAF",
        "reason": reason_to_split,
        "parse_diag": parse_diag,
        "schema_diag": schema_diag,
        "finish_reason": call_result.finish_reason,
        "output_chars": len(call_result.text),
        "raw_tail": call_result.text[-1200:] if call_result.text else "",
    })

    if not can_split:
        return False, []

    children = _split_chunk_for_retry(chunk)
    if len(children) < 2:
        return False, []

    log_info(
        f"[AGENT1/{chunk.get('chunk_id')}] ✂️ Recursive split do {reason_to_split} | "
        f"{len(chunk.get('core_text', '')):,} chars -> "
        + " + ".join(f"{len(c.get('core_text', '')):,}" for c in children)
    )

    all_matrices = []
    for child in children:
        child_ok, child_matrices = process_agent1_chunk_recursive(
            chunk=child,
            api_key=api_key,
            base_url=base_url,
            model=model,
            prompt_template=prompt_template,
            cache=cache,
            diagnostics=diagnostics,
        )
        if not child_ok:
            return False, []
        all_matrices.extend(child_matrices)
    return True, all_matrices


def _rule_signature(rule: dict) -> tuple:
    """Dedup bảo thủ: chỉ loại rule rất giống nhau do overlap/retry."""
    return (
        _normalize_text(rule.get("target")),
        _normalize_text(rule.get("category")),
        _normalize_text(rule.get("rule_type")),
        _normalize_text(rule.get("rule_name")),
        _normalize_text(rule.get("source_requirement")),
        _normalize_text(rule.get("applied_qa_rule")),
    )


def _infer_rule_prefix(rules: list[dict], fallback: str) -> str:
    prefixes = []
    for rule in rules:
        rid = str(rule.get("rule_id", "")).strip()
        m = re.match(r"^(.+?)[-_](\d+)$", rid)
        if m:
            p = m.group(1).strip()
            if p and p.upper() not in {"TMP", "RULE", "R"}:
                prefixes.append(p)
    if prefixes:
        prefix = Counter(prefixes).most_common(1)[0][0]
    else:
        prefix = fallback
    prefix = unicodedata.normalize("NFKD", prefix).encode("ascii", "ignore").decode("ascii")
    prefix = re.sub(r"[^A-Za-z0-9_]+", "_", prefix).strip("_").upper()
    return prefix or fallback


# ============================================================
# PATCH V2 — CANONICAL SCREEN MERGE
# Thay thế hàm merge_rule_matrices hiện tại bằng block này.
# Dùng helper/import sẵn có trong file:
# _normalize_text, _rule_signature, _infer_rule_prefix,
# OrderedDict, copy, re, unicodedata
# ============================================================

def _screen_core_tokens(name: str) -> tuple[str, ...]:
    s = unicodedata.normalize("NFKD", str(name or "")).encode("ascii", "ignore").decode("ascii").lower()
    s = re.sub(r"[^a-z0-9]+", " ", s)

    generic = {
        "man", "hinh", "screen", "page", "module",
        "bidv", "trai", "phieu"
    }
    return tuple(t for t in s.split() if t and t not in generic)


def _stable_screen_code(screen: dict) -> str | None:
    name = str(screen.get("screen_name", "") or "").strip()

    # 1) Screen code xuất hiện trực tiếp trong screen_name.
    for code in re.findall(r"\b[A-Z][A-Z0-9]*(?:_[A-Z0-9]+){1,}\b", name):
        c = re.sub(r"_\d+$", "", code.upper())
        if c and not c.startswith("SCREEN"):
            return c

    # 2) Fallback từ rule_id, ví dụ:
    # BOND_ORDER_LIST-001 / BOND_ORDER_LIST_2-001 / BOND_ORDER_LIST_3-001
    # -> BOND_ORDER_LIST
    for rule in screen.get("test_rules", []) or []:
        rid = str(rule.get("rule_id", "") or "").strip().upper()
        if not rid:
            continue

        prefix = re.sub(r"-\d+$", "", rid)
        prefix = re.sub(r"_\d+$", "", prefix)

        if (
            prefix
            and "_" in prefix
            and not prefix.startswith("SCREEN")
            and re.fullmatch(r"[A-Z][A-Z0-9_]+", prefix)
        ):
            return prefix

    return None


def _screen_alias_equivalent(name_a: str, name_b: str) -> bool:
    """
    Merge bảo thủ:
      Danh sách lệnh đặt mua
      Màn hình Danh sách lệnh đặt mua Trái phiếu
    => cùng screen.

    Không merge:
      Danh sách lệnh đặt mua
      Màn hình xem chi tiết lệnh đặt mua
    """
    a = _screen_core_tokens(name_a)
    b = _screen_core_tokens(name_b)

    if not a or not b:
        return False

    if a == b:
        return True

    sa, sb = set(a), set(b)
    inter = len(sa & sb)
    union = len(sa | sb)
    jaccard = inter / max(1, union)

    # Token phân biệt loại màn hình.
    discriminators = {
        "danh", "sach", "chi", "tiet", "them", "moi",
        "chinh", "sua", "tao", "duyet", "xem"
    }

    if (sa & discriminators) != (sb & discriminators):
        return False

    return inter >= 4 and jaccard >= 0.85



def _is_technical_screen_code(name: str) -> bool:
    s = str(name or "").strip()
    return bool(re.fullmatch(r"[A-Z][A-Z0-9_]+", s)) and "_" in s


def _prefer_human_screen_name(current: str, candidate: str) -> str:
    """
    Giữ tên nghiệp vụ dễ đọc.
    Nếu current chỉ là code kỹ thuật và candidate là tên người đọc được -> dùng candidate.
    Nếu cả hai đều là tên người đọc được -> ưu tiên tên mô tả rõ hơn nhưng không quá dài.
    """
    cur = str(current or "").strip()
    cand = str(candidate or "").strip()
    if not cand:
        return cur
    if not cur:
        return cand

    if _is_technical_screen_code(cur) and not _is_technical_screen_code(cand):
        return cand
    if not _is_technical_screen_code(cur) and _is_technical_screen_code(cand):
        return cur

    # Ưu tiên tên có "màn hình" hoặc "trái phiếu" khi cùng alias,
    # vì thường mô tả domain rõ hơn.
    cur_score = (2 if "màn hình" in cur.lower() else 0) + (1 if "trái phiếu" in cur.lower() else 0)
    cand_score = (2 if "màn hình" in cand.lower() else 0) + (1 if "trái phiếu" in cand.lower() else 0)
    if cand_score > cur_score:
        return cand
    if cand_score == cur_score and len(cand) > len(cur) and len(cand) <= 80:
        return cand
    return cur


def merge_rule_matrices(matrices: list[dict]) -> tuple[dict, dict]:
    """
    Merge Python-only:
    - canonicalize alias cùng screen;
    - exact-dedup rule;
    - renumber rule_id.
    Không suy luận QA/business.
    """
    screens_map: OrderedDict[str, dict] = OrderedDict()
    seen_by_screen: dict[str, set] = {}
    screen_meta: OrderedDict[str, dict] = OrderedDict()

    raw_rules = 0
    exact_duplicates = 0
    screen_aliases_merged = 0

    for matrix in matrices:
        for screen in matrix.get("screens", []):
            name = str(screen.get("screen_name", "")).strip() or "Unnamed Screen"
            stable_code = _stable_screen_code(screen)

            canonical_key = None

            # A. Screen code là identity mạnh nhất.
            if stable_code:
                code_key = f"code::{stable_code.lower()}"
                if code_key in screens_map:
                    canonical_key = code_key
                else:
                    for k, meta in screen_meta.items():
                        if meta.get("stable_code") == stable_code:
                            canonical_key = k
                            break
                    if canonical_key is None:
                        canonical_key = code_key

            # B. Exact normalized name.
            if canonical_key is None:
                exact_key = f"name::{_normalize_text(name)}"
                if exact_key in screens_map:
                    canonical_key = exact_key

            # C. Alias gần như cùng tên, merge bảo thủ.
            if canonical_key is None:
                for k, meta in screen_meta.items():
                    if _screen_alias_equivalent(name, meta["screen_name"]):
                        canonical_key = k
                        screen_aliases_merged += 1
                        break

            # D. Screen mới.
            if canonical_key is None:
                canonical_key = f"name::{_normalize_text(name)}"

            if canonical_key not in screens_map:
                # stable_code chỉ dùng làm identity/key. Tên hiển thị vẫn là tên nghiệp vụ.
                display_name = name
                screens_map[canonical_key] = {
                    "screen_name": display_name,
                    "test_rules": []
                }
                seen_by_screen[canonical_key] = set()
                screen_meta[canonical_key] = {
                    "screen_name": name,
                    "stable_code": stable_code,
                }
            elif stable_code and not screen_meta[canonical_key].get("stable_code"):
                # Khi chunk sau tìm ra code thật, chỉ cập nhật identity nội bộ.
                # KHÔNG thay screen_name hiển thị thành code.
                screen_meta[canonical_key]["stable_code"] = stable_code

            # Cập nhật tên hiển thị nếu alias mới human-readable hơn.
            screens_map[canonical_key]["screen_name"] = _prefer_human_screen_name(
                screens_map[canonical_key].get("screen_name", ""),
                name,
            )

            for rule in screen.get("test_rules", []):
                raw_rules += 1
                sig = _rule_signature(rule)

                if sig in seen_by_screen[canonical_key]:
                    exact_duplicates += 1
                    continue

                seen_by_screen[canonical_key].add(sig)
                screens_map[canonical_key]["test_rules"].append(copy.deepcopy(rule))

    final_screens = list(screens_map.values())

    used_prefixes = set()
    for screen_idx, screen in enumerate(final_screens, start=1):
        fallback = f"SCREEN_{screen_idx:02d}"
        prefix = _infer_rule_prefix(screen.get("test_rules", []), fallback)

        base_prefix = prefix
        suffix = 2
        while prefix in used_prefixes:
            prefix = f"{base_prefix}_{suffix}"
            suffix += 1
        used_prefixes.add(prefix)

        for idx, rule in enumerate(screen.get("test_rules", []), start=1):
            rule["rule_id"] = f"{prefix}-{idx:03d}"

    final = {
        "test_design_version": "3.1",
        "screens": final_screens,
    }

    stats = {
        "input_matrices": len(matrices),
        "screens": len(final_screens),
        "raw_rules": raw_rules,
        "screen_aliases_merged": screen_aliases_merged,
        "exact_duplicates_removed": exact_duplicates,
        "final_rules": sum(len(s.get("test_rules", [])) for s in final_screens),
    }
    return final, stats



def run_agent1_document_pipeline(
    raw_text: str,
    base_filename: str,
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    cache: dict | None = None,
    progress_callback=None,
) -> tuple[bool, dict | None, dict]:
    """Scalable Agent 1 pipeline cho 70K/150K/300K+ chars."""
    started = time.time()
    initial_chunks = split_document_semantic(raw_text, base_filename)
    diagnostics = []
    all_matrices = []

    log_info(
        f"[AGENT1_PIPELINE] 📚 DocumentChars={len(raw_text):,} | "
        f"InitialChunks={len(initial_chunks)} | Target≈{AGENT1_CHUNK_TARGET_CHARS:,}"
    )

    for idx, chunk in enumerate(initial_chunks, start=1):
        if progress_callback:
            progress_callback(
                idx - 1,
                len(initial_chunks),
                f"Agent 1 đang xử lý chunk {idx}/{len(initial_chunks)} — {chunk.get('title')} ({chunk.get('char_count', 0):,} chars)"
            )

        ok, matrices = process_agent1_chunk_recursive(
            chunk=chunk,
            api_key=api_key,
            base_url=base_url,
            model=model,
            prompt_template=prompt_template,
            cache=cache,
            diagnostics=diagnostics,
        )
        if not ok:
            summary = {
                "ok": False,
                "document_chars": len(raw_text),
                "initial_chunks": len(initial_chunks),
                "elapsed": round(time.time() - started, 2),
                "diagnostics": diagnostics,
            }
            if progress_callback:
                progress_callback(idx - 1, len(initial_chunks), f"Agent 1 thất bại tại chunk {chunk.get('chunk_id')}")
            return False, None, summary

        all_matrices.extend(matrices)
        if progress_callback:
            progress_callback(idx, len(initial_chunks), f"Hoàn tất chunk {idx}/{len(initial_chunks)}")

    final_matrix, merge_stats = merge_rule_matrices(all_matrices)
    schema_ok, schema_diag = validate_rule_matrix_schema(final_matrix, strict=True)
    elapsed = time.time() - started

    summary = {
        "ok": schema_ok,
        "document_chars": len(raw_text),
        "initial_chunks": len(initial_chunks),
        "leaf_matrices": len(all_matrices),
        "elapsed": round(elapsed, 2),
        "merge": merge_stats,
        "schema": schema_diag,
        "diagnostics": diagnostics,
    }

    if not schema_ok:
        log_error(f"[AGENT1_PIPELINE] ❌ Final schema fail | {schema_diag}")
        return False, None, summary

    log_info(
        f"[AGENT1_PIPELINE] ✅ Completed | Time={elapsed:.2f}s | "
        f"Screens={merge_stats['screens']} | FinalRules={merge_stats['final_rules']} | "
        f"ExactDupRemoved={merge_stats['exact_duplicates_removed']}"
    )
    return True, final_matrix, summary


def split_screen_rules_for_agent2(
    screen: dict,
    max_rules: int = AGENT2_BATCH_MAX_RULES,
    max_input_chars: int = AGENT2_BATCH_MAX_INPUT_CHARS,
) -> list[dict]:
    """Pack rule theo count + estimated JSON chars."""
    rules = screen.get("test_rules", [])
    batches = []
    current = []
    current_chars = 0

    for rule in rules:
        rule_chars = len(json.dumps(rule, ensure_ascii=False))
        if current and (len(current) >= max_rules or current_chars + rule_chars > max_input_chars):
            batches.append({"screen_name": screen.get("screen_name", "Unnamed"), "test_rules": current})
            current = []
            current_chars = 0
        current.append(rule)
        current_chars += rule_chars

    if current:
        batches.append({"screen_name": screen.get("screen_name", "Unnamed"), "test_rules": current})
    return batches


def render_agent2_batch_recursive(
    batch_screen: dict,
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    batch_id: str,
    depth: int = 0,
    diagnostics: list | None = None,
    cache: dict | None = None,
) -> tuple[bool, list[str]]:
    diagnostics = diagnostics if diagnostics is not None else []
    payload = json.dumps({"test_design_version": "3.1", "screens": [batch_screen]}, ensure_ascii=False)
    key = _cache_key("agent2-v3.1", model, prompt_template, payload)

    if cache is not None and key in cache:
        diagnostics.append({"batch_id": batch_id, "status": "CACHE_HIT", "rules": len(batch_screen.get("test_rules", []))})
        return True, [cache[key]]

    call_result = None
    for attempt in range(AGENT2_API_RETRIES + 1):
        call_result = call_qwen_max_agent_detailed(
            content=payload,
            api_key=api_key,
            base_url=base_url,
            model=model,
            prompt_template=prompt_template,
            max_tokens=QWEN_MAX_OUTPUT_TOKENS,
            agent_name=f"AGENT2/{batch_id}",
        )
        if call_result.ok or call_result.finish_reason == "length":
            break
        if attempt < AGENT2_API_RETRIES:
            log_info(f"[AGENT2/{batch_id}] 🔁 Retry API lần {attempt + 2}/{AGENT2_API_RETRIES + 1}")
            time.sleep(2)

    assert call_result is not None

    if call_result.complete:
        diagnostics.append({
            "batch_id": batch_id,
            "status": "OK",
            "rules": len(batch_screen.get("test_rules", [])),
            "elapsed": round(call_result.elapsed, 2),
            "output_chars": len(call_result.text),
        })
        if cache is not None:
            cache[key] = call_result.text
        return True, [call_result.text]

    rules = batch_screen.get("test_rules", [])
    can_split = call_result.finish_reason == "length" and len(rules) > 1 and depth < AGENT2_MAX_RECURSION_DEPTH
    diagnostics.append({
        "batch_id": batch_id,
        "status": "SPLIT_RETRY" if can_split else "FAILED",
        "reason": "MAX_TOKENS" if call_result.finish_reason == "length" else (call_result.error or "UNKNOWN"),
        "rules": len(rules),
        "output_chars": len(call_result.text),
    })

    if not can_split:
        return False, []

    mid = max(1, len(rules) // 2)
    children = [rules[:mid], rules[mid:]]
    outputs = []
    for idx, child_rules in enumerate(children, start=1):
        if not child_rules:
            continue
        child_screen = {"screen_name": batch_screen.get("screen_name", "Unnamed"), "test_rules": child_rules}
        ok, child_outputs = render_agent2_batch_recursive(
            batch_screen=child_screen,
            api_key=api_key,
            base_url=base_url,
            model=model,
            prompt_template=prompt_template,
            batch_id=f"{batch_id}.{idx}",
            depth=depth + 1,
            diagnostics=diagnostics,
            cache=cache,
        )
        if not ok:
            return False, []
        outputs.extend(child_outputs)
    return True, outputs


def _parse_bullet_forest(text: str) -> list[dict]:
    clean = re.sub(r"^```[a-zA-Z]*\s*", "", text.strip())
    clean = re.sub(r"\s*```$", "", clean).strip()
    root = {"title": "__ROOT__", "children": []}
    stack = [(-1, root)]

    for raw_line in clean.splitlines():
        if not raw_line.strip():
            continue
        line = raw_line.replace("\t", "    ")
        indent = len(line) - len(line.lstrip(" "))
        level = indent // 2
        content = line.strip()
        if content.startswith("- ") or content.startswith("* "):
            content = content[2:].strip()
        if not content:
            continue

        node = {"title": content, "children": []}
        while stack and stack[-1][0] >= level:
            stack.pop()
        parent = stack[-1][1] if stack else root
        parent["children"].append(node)
        stack.append((level, node))
    return root["children"]


def _merge_bullet_nodes(target_children: list[dict], incoming: list[dict], depth: int = 0):
    """Merge cùng Screen/Category/Target ở 3 level đầu; từ testcase trở xuống luôn append."""
    if depth >= 3:
        target_children.extend(copy.deepcopy(incoming))
        return

    index = {_normalize_text(node.get("title")): node for node in target_children}
    for node in incoming:
        key = _normalize_text(node.get("title"))
        if key in index:
            _merge_bullet_nodes(index[key]["children"], node.get("children", []), depth + 1)
        else:
            cloned = copy.deepcopy(node)
            target_children.append(cloned)
            index[key] = cloned


def _renumber_testcases_in_nodes(nodes: list[dict]) -> int:
    counter = 0

    def walk(children):
        nonlocal counter
        for node in children:
            title = node.get("title", "")
            if re.match(r"^TC[_\- ]?\d+\s*[-:]", title, flags=re.IGNORECASE):
                counter += 1
                node["title"] = re.sub(
                    r"^TC[_\- ]?\d+",
                    f"TC_{counter:03d}",
                    title,
                    count=1,
                    flags=re.IGNORECASE,
                )
            walk(node.get("children", []))

    walk(nodes)
    return counter


def _serialize_bullet_nodes(nodes: list[dict], depth: int = 0) -> list[str]:
    lines = []
    for node in nodes:
        lines.append("  " * depth + "- " + str(node.get("title", "")).strip())
        lines.extend(_serialize_bullet_nodes(node.get("children", []), depth + 1))
    return lines


def merge_agent2_tree_outputs(outputs: list[str]) -> tuple[str, dict]:
    merged_nodes = []
    for output in outputs:
        forest = _parse_bullet_forest(output)
        _merge_bullet_nodes(merged_nodes, forest, depth=0)
    tc_count = _renumber_testcases_in_nodes(merged_nodes)
    merged_text = "\n".join(_serialize_bullet_nodes(merged_nodes)).strip()
    return merged_text, {"agent2_leaf_outputs": len(outputs), "testcases_renumbered": tc_count}


def run_agent2_rule_matrix_pipeline(
    approved_screens: list[dict],
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    cache: dict | None = None,
    progress_callback=None,
) -> tuple[bool, str, dict]:
    """Render Rule Matrix lớn theo batch, recursive split nếu chạm max_tokens."""
    started = time.time()
    initial_batches = []
    for screen_idx, screen in enumerate(approved_screens, start=1):
        for batch_idx, batch in enumerate(split_screen_rules_for_agent2(screen), start=1):
            initial_batches.append((f"S{screen_idx:02d}B{batch_idx:02d}", batch))

    diagnostics = []
    outputs = []

    log_info(
        f"[AGENT2_PIPELINE] 🧩 Screens={len(approved_screens)} | "
        f"InitialBatches={len(initial_batches)} | "
        f"Rules={sum(len(s.get('test_rules', [])) for s in approved_screens)}"
    )

    for idx, (batch_id, batch) in enumerate(initial_batches, start=1):
        if progress_callback:
            progress_callback(
                idx - 1,
                len(initial_batches),
                f"Agent 2 đang render batch {idx}/{len(initial_batches)} — {batch.get('screen_name')} ({len(batch.get('test_rules', []))} rules)"
            )
        ok, batch_outputs = render_agent2_batch_recursive(
            batch_screen=batch,
            api_key=api_key,
            base_url=base_url,
            model=model,
            prompt_template=prompt_template,
            batch_id=batch_id,
            diagnostics=diagnostics,
            cache=cache,
        )
        if not ok:
            return False, "", {
                "ok": False,
                "initial_batches": len(initial_batches),
                "elapsed": round(time.time() - started, 2),
                "diagnostics": diagnostics,
            }
        outputs.extend(batch_outputs)
        if progress_callback:
            progress_callback(idx, len(initial_batches), f"Hoàn tất batch {idx}/{len(initial_batches)}")

    merged_text, merge_stats = merge_agent2_tree_outputs(outputs)
    summary = {
        "ok": True,
        "initial_batches": len(initial_batches),
        "elapsed": round(time.time() - started, 2),
        "merge": merge_stats,
        "diagnostics": diagnostics,
    }
    log_info(
        f"[AGENT2_PIPELINE] ✅ Completed | Time={summary['elapsed']:.2f}s | "
        f"LeafOutputs={merge_stats['agent2_leaf_outputs']} | TestCases={merge_stats['testcases_renumbered']}"
    )
    return True, merged_text, summary


def create_xmind_from_text(tree_data: str, output_path: str, root_title: str = "Kế hoạch & Kịch bản Kiểm thử"):
    try:
        log_info(f"Tiến hành parse dữ liệu text sang cấu trúc file .xmind...")
        root_topic = {
            "id": "root_node",
            "title": root_title,
            "children": {"attached": []}
        }
        
        clean_data_str = re.sub(r'^```[a-zA-Z]*\n', '', tree_data, flags=re.MULTILINE)
        clean_data_str = re.sub(r'\n```$', '', clean_data_str, flags=re.MULTILINE).strip()
        
        node_id_counter = [1]
        lines = clean_data_str.split("\n")
        stack = [(0, root_topic)]

        for line in lines:
            if not line.strip():
                continue
                
            expanded_line = line.replace("\t", "    ")
            indent_spaces = len(expanded_line) - len(expanded_line.lstrip(" "))
            
            content = expanded_line.strip()
            
            if content.startswith("- "):
                content = content[2:].strip()
            elif content.startswith("* "):
                content = content[2:].strip()
                
            content = content.replace("\\n", "\n")

            if not content:
                continue

            level = (indent_spaces // 2) + 1

            new_node = {
                "id": f"node_{node_id_counter[0]}",
                "title": content,
                "children": {"attached": []}
            }
            node_id_counter[0] += 1

            while stack and stack[-1][0] >= level:
                stack.pop()

            parent_node = stack[-1][1]
            parent_node["children"]["attached"].append(new_node)
            stack.append((level, new_node))

        content_json = [{
            "id": "sheet_1",
            "title": "Sơ đồ Kiểm thử QA",
            "rootTopic": root_topic
        }]
        
        manifest_json = {"file-entries": {"content.json": {}, "metadata.json": {}}}
        metadata_json = {}

        with zipfile.ZipFile(output_path, 'w', zipfile.ZIP_DEFLATED) as zip_file:
            zip_file.writestr('content.json', json.dumps(content_json, ensure_ascii=False, indent=2))
            zip_file.writestr('manifest.json', json.dumps(manifest_json, ensure_ascii=False, indent=2))
            zip_file.writestr('metadata.json', json.dumps(metadata_json, ensure_ascii=False, indent=2))
            
        log_info(f"✅ Đã đóng gói thành công file XMind tại: {output_path}")
        return True, None
    except Exception as e:
        err_msg = log_error("Lỗi khi đóng gói file XMind", e)
        return False, err_msg

# ==============================================================================
# HÀM BÓC TÁCH CÂY XMIND VÀ XUẤT EXCEL CHUẨN THEO NHÁNH NHÓM (DỌC THEO TÊN MÀN HÌNH)
# ==============================================================================
# ==============================================================================
# HÀM BÓC TÁCH CÂY XMIND VÀ XUẤT EXCEL CHUẨN (CẬP NHẬT DOCSTRING & QUÉT ĐỘ SÂU)
# ==============================================================================
def convert_tree_to_ui_excel(tree_text: str) -> bytes:
    """
    Chuyển toàn bộ cây Test Case Web/XMind thành workbook Excel.

    FIX V3.3:
    - KHÔNG còn chỉ lấy root/screen đầu tiên.
    - Mỗi Screen Root có testcase được xuất thành 1 worksheet riêng.
    - Thêm worksheet Summary để đối soát số testcase theo màn hình.
    - Chỉ nhận node bắt đầu bằng TC_... / TC_API_... là testcase;
      không nhận nhầm leaf text có dấu " - ".
    - Quét đệ quy mọi Target/Sub-group nên không mất testcase do độ sâu cây thay đổi.
    - Tự kiểm tra tổng TC trong tree == tổng TC đã ghi Excel; lệch là raise lỗi,
      tránh xuất file thiếu testcase một cách im lặng.
    """
    clean_text = re.sub(r'^```[a-zA-Z]*\s*', '', str(tree_text or '').strip())
    clean_text = re.sub(r'\s*```$', '', clean_text).strip()
    if not clean_text:
        raise ValueError("Không có dữ liệu Test Case để xuất Excel.")

    # ----------------------------------------------------------
    # 1. Parse plain-text bullet tree -> node forest
    # ----------------------------------------------------------
    root_node = {"title": "__ROOT__", "children": []}
    stack = [(-1, root_node)]

    for raw_line in clean_text.splitlines():
        if not raw_line.strip():
            continue

        expanded = raw_line.replace("\t", "    ")
        indent = len(expanded) - len(expanded.lstrip(" "))
        level = indent // 2
        content = expanded.strip()
        if content.startswith("- ") or content.startswith("* "):
            content = content[2:].strip()
        if not content:
            continue

        node = {"title": content.replace('\\n', '\n'), "children": []}
        while stack and stack[-1][0] >= level:
            stack.pop()
        parent = stack[-1][1] if stack else root_node
        parent["children"].append(node)
        stack.append((level, node))

    # ----------------------------------------------------------
    # 2. Helpers
    # ----------------------------------------------------------
    tc_regex = re.compile(r'^TC(?:_API)?[_\- ]?\d+\s*(?:[-:]|$)', re.IGNORECASE)

    def is_test_case_node(node: dict) -> bool:
        return bool(tc_regex.match(str(node.get("title", "")).strip()))

    def count_testcases(node: dict) -> int:
        total = 1 if is_test_case_node(node) else 0
        for child in node.get("children", []):
            total += count_testcases(child)
        return total

    def has_test_cases(node: dict) -> bool:
        return count_testcases(node) > 0

    def clean_prefix(value: str, prefixes: list[str]) -> str:
        t = str(value or "").strip()
        low = t.lower()
        for prefix in prefixes:
            if low.startswith(prefix.lower()):
                return t[len(prefix):].strip()
        return t

    def extract_details(tc_node: dict) -> tuple[str, str, str]:
        """Lấy Pre-condition / Steps / Expected ở bất kỳ độ sâu con nào."""
        pre = ""
        steps = ""
        expected = ""

        def walk(node: dict):
            nonlocal pre, steps, expected
            for child in node.get("children", []):
                title = str(child.get("title", "")).strip()
                low = title.lower()

                if low.startswith(("pre-condition:", "precondition:", "tiền điều kiện:")):
                    if not pre:
                        pre = clean_prefix(title, ["Pre-condition:", "Precondition:", "Tiền điều kiện:"])
                elif low.startswith(("các bước thực hiện:", "các bước:", "steps:", "step:")):
                    if not steps:
                        steps = clean_prefix(title, ["Các bước thực hiện:", "Các bước:", "Steps:", "Step:"])
                elif low.startswith(("kết quả mong đợi:", "expected result:", "expected:", "response (kết quả mong đợi):")):
                    if not expected:
                        expected = clean_prefix(
                            title,
                            ["Kết quả mong đợi:", "Expected Result:", "Expected:", "Response (Kết quả mong đợi):"]
                        )

                walk(child)

        walk(tc_node)
        return pre, steps, expected

    def split_tc_title(title: str) -> tuple[str, str]:
        t = str(title or "").strip()
        m = re.match(r'^(TC(?:_API)?[_\- ]?\d+)\s*(?:[-:]\s*)?(.*)$', t, flags=re.IGNORECASE)
        if not m:
            return "", t
        return m.group(1).strip(), m.group(2).strip()

    def sanitize_sheet_title(title: str, used: set[str]) -> str:
        name = re.sub(r'[\\/*?:\[\]]', '_', str(title or "Screen")).strip()
        name = name or "Screen"
        name = name[:31]
        base = name
        idx = 2
        while name.lower() in used:
            suffix = f"_{idx}"
            name = (base[:31-len(suffix)] + suffix)[:31]
            idx += 1
        used.add(name.lower())
        return name

    # Screen roots = mọi top-level node có ít nhất một testcase.
    # Nếu model bọc thêm một root chung duy nhất, unwrap một cấp khi hợp lý.
    screen_nodes = [node for node in root_node["children"] if has_test_cases(node)]
    if len(screen_nodes) == 1:
        only = screen_nodes[0]
        child_roots = [child for child in only.get("children", []) if has_test_cases(child)]
        # Không unwrap nếu children trông giống category chuẩn.
        category_like = any(
            re.match(r'^\d+\.\s*(kiểm tra|kiem tra)', str(c.get("title", "")).strip(), flags=re.IGNORECASE)
            for c in child_roots
        )
        if child_roots and not category_like and all(count_testcases(c) > 0 for c in child_roots):
            screen_nodes = child_roots

    if not screen_nodes:
        raise ValueError("Không tìm thấy Test Case hợp lệ (TC_...) trong cây để xuất Excel.")

    expected_total = sum(count_testcases(screen) for screen in screen_nodes)

    # ----------------------------------------------------------
    # 3. Workbook + styles
    # ----------------------------------------------------------
    wb = Workbook()
    default_ws = wb.active
    wb.remove(default_ws)

    fill_green = PatternFill(start_color="81C784", fill_type="solid")
    fill_light_green = PatternFill(start_color="C8E6C9", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", fill_type="solid")
    fill_header_bg = PatternFill(start_color="4DD0E1", fill_type="solid")
    fill_pink = PatternFill(start_color="F8BBD0", fill_type="solid")
    fill_blue_exec = PatternFill(start_color="90CAF9", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='B0BEC5'),
        right=Side(style='thin', color='B0BEC5'),
        top=Side(style='thin', color='B0BEC5'),
        bottom=Side(style='thin', color='B0BEC5')
    )

    font_title = Font(name="Segoe UI", size=11, bold=True)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_body = Font(name="Segoe UI", size=10)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    headers = [
        ("A11", "External ID", fill_header_bg),
        ("B11", "Name", fill_pink),
        ("C11", "PreConditions", fill_pink),
        ("D11", "Importance", fill_pink),
        ("E11", "Step", fill_pink),
        ("F11", "Data test", fill_pink),
        ("G11", "Expected Result", fill_pink),
        ("H11", "Actual Result", fill_pink),
        ("I11", "Lần 1", fill_blue_exec),
        ("J11", "Lần 2", fill_blue_exec),
    ]
    col_widths = {'A': 15, 'B': 52, 'C': 36, 'D': 12, 'E': 55, 'F': 20, 'G': 55, 'H': 25, 'I': 10, 'J': 10}

    used_sheet_names = set()
    exported_total = 0
    summary_rows = []

    def write_sheet(screen: dict):
        nonlocal exported_total

        screen_name = str(screen.get("title", "Tên Màn Hình")).strip() or "Tên Màn Hình"
        ws = wb.create_sheet(title=sanitize_sheet_title(screen_name, used_sheet_names))
        ws.views.sheetView[0].showGridLines = True

        # Metadata block
        ws['D1'] = "KỊCH BẢN KIỂM THỬ *"
        ws['D1'].font = font_title
        ws['D1'].alignment = align_center
        ws['C2'] = "Tên màn hình/Tên chức năng"
        ws['C2'].font = font_bold
        ws['D2'] = screen_name
        ws['D2'].font = font_body

        for r in range(1, 5):
            for col in ["C", "D"]:
                ws[f"{col}{r}"].border = thin_border

        for pos, label, fill in headers:
            ws[pos] = label
            ws[pos].font = font_bold
            ws[pos].alignment = align_center
            ws[pos].border = thin_border
            ws[pos].fill = fill

        for col, width in col_widths.items():
            ws.column_dimensions[col].width = width

        current_row = 14
        screen_exported = 0

        def write_group_header(title: str, indent_level: int, strong: bool):
            nonlocal current_row
            if not title:
                return
            prefix = "  " * max(0, indent_level)
            ws.cell(row=current_row, column=2, value=f"{prefix}{title}").font = font_bold
            fill = fill_green if strong else fill_light_green
            for col in range(1, 11):
                c = ws.cell(row=current_row, column=col)
                c.fill = fill
                c.border = thin_border
            current_row += 1

        def write_tc(node: dict):
            nonlocal current_row, screen_exported, exported_total
            tc_id, tc_name = split_tc_title(node.get("title", ""))
            pre, steps, expected = extract_details(node)

            ws.cell(row=current_row, column=1, value=tc_id).alignment = align_center
            ws.cell(row=current_row, column=2, value=tc_name)
            ws.cell(row=current_row, column=3, value=pre)
            ws.cell(row=current_row, column=4, value=None)
            ws.cell(row=current_row, column=5, value=steps)
            ws.cell(row=current_row, column=6, value=None)
            ws.cell(row=current_row, column=7, value=expected)

            for col in range(1, 11):
                c = ws.cell(row=current_row, column=col)
                c.font = font_body
                c.fill = fill_white
                c.border = thin_border
                c.alignment = align_center if col in (1, 4, 9, 10) else align_left

            current_row += 1
            screen_exported += 1
            exported_total += 1

        def traverse(node: dict, depth: int = 0):
            if is_test_case_node(node):
                write_tc(node)
                return

            if not has_test_cases(node):
                return

            # Node dưới screen: category ở depth 0, target/subgroup ở depth >=1
            write_group_header(str(node.get("title", "")).strip(), depth, strong=(depth == 0))
            for child in node.get("children", []):
                traverse(child, depth + 1)

        for child in screen.get("children", []):
            traverse(child, depth=0)

        ws.freeze_panes = "A12"
        ws.auto_filter.ref = f"A11:J{max(11, current_row - 1)}"
        summary_rows.append((screen_name, screen_exported, ws.title))

    for screen in screen_nodes:
        write_sheet(screen)

    # ----------------------------------------------------------
    # 4. Safety invariant: XMind tree TC count == Excel TC count
    # ----------------------------------------------------------
    if exported_total != expected_total:
        raise ValueError(
            f"Excel export không đầy đủ: Tree có {expected_total} testcase nhưng Excel mới ghi {exported_total}."
        )

    # ----------------------------------------------------------
    # 5. Summary sheet
    # ----------------------------------------------------------
    summary = wb.create_sheet(title="Summary", index=0)
    summary['A1'] = "ĐỐI SOÁT TEST CASE"
    summary['A1'].font = Font(name="Segoe UI", size=12, bold=True)
    summary['A3'] = "Màn hình / Chức năng"
    summary['B3'] = "Số Test Case"
    summary['C3'] = "Worksheet"
    for cell in summary[3]:
        cell.font = font_bold
        cell.fill = fill_header_bg
        cell.border = thin_border
        cell.alignment = align_center

    row = 4
    for screen_name, tc_count, sheet_name in summary_rows:
        summary.cell(row=row, column=1, value=screen_name)
        summary.cell(row=row, column=2, value=tc_count)
        summary.cell(row=row, column=3, value=sheet_name)
        for col in range(1, 4):
            summary.cell(row=row, column=col).border = thin_border
            summary.cell(row=row, column=col).alignment = align_left if col != 2 else align_center
        row += 1

    summary.cell(row=row, column=1, value="TOTAL").font = font_bold
    summary.cell(row=row, column=2, value=exported_total).font = font_bold
    for col in range(1, 4):
        summary.cell(row=row, column=col).border = thin_border

    summary.column_dimensions['A'].width = 55
    summary.column_dimensions['B'].width = 18
    summary.column_dimensions['C'].width = 35

    try:
        log_info(
            f"[EXCEL_EXPORT] Screens={len(screen_nodes)} | TreeTC={expected_total} | ExcelTC={exported_total}"
        )
    except Exception:
        pass

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


# ==========================================
# 2.2 PIPELINE API V3.2 — CONTRACT QA BRAIN / LONG DOCUMENT
# ==========================================

API_RULE_FIELDS = {
    "rule_id", "target", "category", "rule_type", "rule_name",
    "test_objective", "test_condition", "expected_result",
    "source_requirement", "source_document", "applied_qa_rule", "generation_reason"
}
API_ALLOWED_CATEGORIES = {
    "AUTHENTICATION", "AUTHORIZATION", "METHOD_URL", "REQUEST_VALIDATION",
    "HAPPY_PATH", "BUSINESS_RULE", "RESPONSE_VALIDATION", "INTEGRATION", "EXCEPTION"
}
API_ALLOWED_RULE_TYPES = {"EXPLICIT", "DERIVED"}
API_ALLOWED_SOURCE_DOCUMENTS = {"API_SPEC", "BA", "BOTH"}
API_HTTP_METHODS = {"GET", "POST", "PUT", "PATCH", "DELETE", "HEAD", "OPTIONS", "UNMAPPED"}


def _try_load_structured_api_doc(text: str):
    if not text or not text.strip():
        return None
    try:
        obj = json.loads(text)
        if isinstance(obj, (dict, list)):
            return obj
    except Exception:
        pass
    try:
        obj = yaml.safe_load(text)
        if isinstance(obj, (dict, list)):
            return obj
    except Exception:
        pass
    return None


def extract_api_endpoint_index(design_text: str) -> list[dict]:
    """Trích endpoint index bằng Python, không suy luận QA.

    Hỗ trợ OpenAPI/Swagger JSON/YAML, Postman collection JSON và regex fallback.
    Index chỉ dùng làm metadata định danh endpoint cho Agent 1 chunk.
    """
    results = []
    seen = set()

    def add(method, path, summary="", module=""):
        method = str(method or "").upper().strip()
        path = str(path or "").strip()
        if method not in API_HTTP_METHODS - {"UNMAPPED"} or not path:
            return
        key = (method, _normalize_text(path))
        if key in seen:
            return
        seen.add(key)
        results.append({
            "method": method,
            "path": path,
            "summary": str(summary or "").strip(),
            "module": str(module or "").strip(),
        })

    obj = _try_load_structured_api_doc(design_text)
    if isinstance(obj, dict):
        # OpenAPI / Swagger
        paths = obj.get("paths")
        if isinstance(paths, dict):
            for path, path_item in paths.items():
                if not isinstance(path_item, dict):
                    continue
                for method, operation in path_item.items():
                    if str(method).lower() not in {"get", "post", "put", "patch", "delete", "head", "options"}:
                        continue
                    operation = operation if isinstance(operation, dict) else {}
                    tags = operation.get("tags") or []
                    module = tags[0] if isinstance(tags, list) and tags else ""
                    add(
                        method,
                        path,
                        operation.get("summary") or operation.get("operationId") or operation.get("description") or "",
                        module,
                    )

        # Postman collection
        if isinstance(obj.get("item"), list):
            def walk_items(items, parents=None):
                parents = parents or []
                for item in items or []:
                    if not isinstance(item, dict):
                        continue
                    name = item.get("name", "")
                    if isinstance(item.get("item"), list):
                        walk_items(item.get("item"), parents + ([name] if name else []))
                        continue
                    req = item.get("request")
                    if not isinstance(req, dict):
                        continue
                    method = req.get("method")
                    url = req.get("url")
                    if isinstance(url, dict):
                        raw_url = url.get("raw") or ""
                        path_parts = url.get("path")
                        if isinstance(path_parts, list) and path_parts:
                            path = "/" + "/".join(str(x) for x in path_parts)
                        else:
                            path = raw_url
                    else:
                        path = str(url or "")
                    module = " / ".join([p for p in parents if p])
                    add(method, path, name, module)
            walk_items(obj.get("item"))

    # Fallback từ text Markdown / raw Swagger / tài liệu mô tả.
    for m in re.finditer(r"(?im)^\s*(GET|POST|PUT|PATCH|DELETE|HEAD|OPTIONS)\s+([/][^\s|`\"']+)", design_text or ""):
        add(m.group(1), m.group(2))
    for m in re.finditer(r'(?im)^\s*[\"\']?(/[^\"\']+)[\"\']?\s*:\s*\{?\s*$', design_text or ""):
        path = m.group(1).strip()
        tail = (design_text or "")[m.end():m.end()+800]
        mm = re.search(r'(?im)[\"\']?(get|post|put|patch|delete|head|options)[\"\']?\s*:', tail)
        if mm:
            add(mm.group(1), path)

    return results


def _tokenize_for_api_hint(text: str) -> set[str]:
    norm = _normalize_text(text)
    return {
        t for t in re.findall(r"[a-z0-9_./-]+", norm)
        if len(t) >= 3 and t not in {"api", "the", "and", "for", "with", "request", "response"}
    }


def select_api_endpoint_hints(chunk_text: str, endpoint_index: list[dict], limit: int = API_ENDPOINT_HINT_LIMIT) -> list[dict]:
    """Chọn endpoint hints bằng lexical matching bảo thủ, không dùng LLM."""
    if not endpoint_index:
        return []
    raw = chunk_text or ""
    chunk_norm = _normalize_text(raw)
    chunk_tokens = _tokenize_for_api_hint(raw)
    scored = []
    for ep in endpoint_index:
        path = str(ep.get("path", ""))
        method = str(ep.get("method", ""))
        summary = str(ep.get("summary", ""))
        module = str(ep.get("module", ""))
        score = 0.0
        if path and _normalize_text(path) in chunk_norm:
            score += 100.0
        if method and re.search(rf"\b{re.escape(method.casefold())}\b", chunk_norm):
            score += 6.0
        ep_tokens = _tokenize_for_api_hint(" ".join([path, summary, module]))
        overlap = len(chunk_tokens & ep_tokens)
        if overlap:
            score += overlap * 2.0
        if score > 0:
            scored.append((score, ep))
    scored.sort(key=lambda x: (-x[0], str(x[1].get("path", ""))))
    return [copy.deepcopy(ep) for _, ep in scored[:limit]]


def build_api_combined_source(design_text: str, ba_text: str) -> str:
    parts = []
    if design_text and design_text.strip():
        parts.append("# === SOURCE DOCUMENT: API_SPEC ===\n" + design_text.strip())
    if ba_text and ba_text.strip():
        parts.append("# === SOURCE DOCUMENT: BA ===\n" + ba_text.strip())
    return "\n\n".join(parts).strip()


def build_api_agent1_chunk_payload(chunk: dict, endpoint_index: list[dict]) -> str:
    hints = select_api_endpoint_hints(chunk.get("core_text", ""), endpoint_index)
    hint_text = json.dumps(hints, ensure_ascii=False, indent=2) if hints else "[]"
    return f"""=== CHUNK METADATA — KHÔNG PHẢI REQUIREMENT ===
chunk_id: {chunk.get('chunk_id')}
source_section_hint: {chunk.get('title')}
chunk_depth: {chunk.get('depth', 0)}

=== ENDPOINT INDEX HINT — CHỈ ĐỊNH DANH, KHÔNG TỰ TẠO TEST RULE TỪ INDEX ===
{hint_text}

=== PREVIOUS CONTEXT — CHỈ DÙNG ĐỂ HIỂU, KHÔNG SINH RULE CHỈ TỪ ĐOẠN NÀY ===
{chunk.get('prev_context', '')}

=== CURRENT SOURCE — PHẠM VI CHÍNH ĐƯỢC PHÉP SINH API TEST RULE ===
{chunk.get('core_text', '')}

=== NEXT CONTEXT — CHỈ DÙNG ĐỂ HIỂU, KHÔNG SINH RULE CHỈ TỪ ĐOẠN NÀY ===
{chunk.get('next_context', '')}
"""


def validate_api_rule_matrix_schema(data, strict: bool = False) -> tuple[bool, str]:
    if not isinstance(data, dict):
        return False, f"Root phải là object, nhận {type(data).__name__}"
    modules = data.get("api_modules")
    if not isinstance(modules, list):
        return False, "Thiếu field 'api_modules' hoặc 'api_modules' không phải array"

    total_endpoints = 0
    total_rules = 0
    seen_ids = set()
    for midx, module in enumerate(modules):
        if not isinstance(module, dict):
            return False, f"api_modules[{midx}] phải là object"
        module_name = module.get("module_name")
        if strict and (not isinstance(module_name, str) or not module_name.strip()):
            return False, f"api_modules[{midx}].module_name rỗng"
        endpoints = module.get("endpoints")
        if not isinstance(endpoints, list):
            return False, f"api_modules[{midx}] thiếu endpoints array"
        for eidx, ep in enumerate(endpoints):
            total_endpoints += 1
            if not isinstance(ep, dict):
                return False, f"api_modules[{midx}].endpoints[{eidx}] phải là object"
            method = str(ep.get("method", "")).upper().strip()
            path = str(ep.get("endpoint_path", "")).strip()
            if method not in API_HTTP_METHODS:
                return False, f"Endpoint {midx}.{eidx} method không hợp lệ: {method!r}"
            if strict and not path:
                return False, f"Endpoint {midx}.{eidx} endpoint_path rỗng"
            rules = ep.get("test_rules")
            if not isinstance(rules, list):
                return False, f"Endpoint {midx}.{eidx} thiếu test_rules array"
            for ridx, rule in enumerate(rules):
                total_rules += 1
                if not isinstance(rule, dict):
                    return False, f"Rule {midx}.{eidx}.{ridx} phải là object"
                missing = sorted(API_RULE_FIELDS - set(rule.keys()))
                if missing:
                    return False, f"Rule {midx}.{eidx}.{ridx} thiếu field: {', '.join(missing)}"
                category = str(rule.get("category", "")).upper().strip()
                rule_type = str(rule.get("rule_type", "")).upper().strip()
                source_document = str(rule.get("source_document", "")).upper().strip()
                if category not in API_ALLOWED_CATEGORIES:
                    return False, f"Rule {midx}.{eidx}.{ridx} category không hợp lệ: {category!r}"
                if rule_type not in API_ALLOWED_RULE_TYPES:
                    return False, f"Rule {midx}.{eidx}.{ridx} rule_type không hợp lệ: {rule_type!r}"
                if source_document not in API_ALLOWED_SOURCE_DOCUMENTS:
                    return False, f"Rule {midx}.{eidx}.{ridx} source_document không hợp lệ: {source_document!r}"
                if strict:
                    for field in API_RULE_FIELDS:
                        if not isinstance(rule.get(field), str):
                            return False, f"Rule {midx}.{eidx}.{ridx}.{field} phải là string"
                    rid = rule.get("rule_id", "").strip()
                    if not rid:
                        return False, f"Rule {midx}.{eidx}.{ridx}.rule_id rỗng"
                    if rid in seen_ids:
                        return False, f"Trùng rule_id: {rid}"
                    seen_ids.add(rid)
                    if not rule.get("source_requirement", "").strip():
                        return False, f"Rule {rid} thiếu source_requirement"
                    if rule_type == "DERIVED" and not rule.get("applied_qa_rule", "").strip():
                        return False, f"Rule {rid} DERIVED thiếu applied_qa_rule"

    return True, f"API Schema OK | Modules={len(modules)} | Endpoints={total_endpoints} | TestRules={total_rules}"


def _api_rule_signature(rule: dict) -> tuple:
    return (
        _normalize_text(rule.get("target")),
        _normalize_text(rule.get("category")),
        _normalize_text(rule.get("rule_type")),
        _normalize_text(rule.get("rule_name")),
        _normalize_text(rule.get("test_condition")),
        _normalize_text(rule.get("expected_result")),
        _normalize_text(rule.get("applied_qa_rule")),
    )


def _merge_api_source_document(a: str, b: str) -> str:
    aa, bb = str(a or "").upper(), str(b or "").upper()
    if aa == bb:
        return aa or bb or "API_SPEC"
    if "BOTH" in {aa, bb}:
        return "BOTH"
    if {aa, bb} == {"API_SPEC", "BA"}:
        return "BOTH"
    return aa or bb or "API_SPEC"


def merge_api_rule_matrices(matrices: list[dict]) -> tuple[dict, dict]:
    """Merge Python-only theo module + METHOD + endpoint_path, dedup bảo thủ và renumber."""
    modules_map: OrderedDict[str, dict] = OrderedDict()
    seen_rule_map = {}
    raw_rules = 0
    duplicates = 0

    for matrix in matrices:
        for module in matrix.get("api_modules", []):
            module_name = str(module.get("module_name", "")).strip() or "UNMAPPED"
            mkey = _normalize_text(module_name)
            if mkey not in modules_map:
                modules_map[mkey] = {"module_name": module_name, "endpoints": OrderedDict()}
            endpoint_map = modules_map[mkey]["endpoints"]

            for ep in module.get("endpoints", []):
                method = str(ep.get("method", "UNMAPPED")).upper().strip() or "UNMAPPED"
                path = str(ep.get("endpoint_path", "UNMAPPED")).strip() or "UNMAPPED"
                ekey = (method, _normalize_text(path))
                if ekey not in endpoint_map:
                    endpoint_map[ekey] = {
                        "method": method,
                        "endpoint_path": path,
                        "summary": str(ep.get("summary", "")).strip(),
                        "test_rules": [],
                    }
                    seen_rule_map[(mkey, ekey)] = {}
                target_ep = endpoint_map[ekey]
                if not target_ep.get("summary") and ep.get("summary"):
                    target_ep["summary"] = str(ep.get("summary", "")).strip()

                sig_map = seen_rule_map[(mkey, ekey)]
                for rule in ep.get("test_rules", []):
                    raw_rules += 1
                    sig = _api_rule_signature(rule)
                    if sig in sig_map:
                        duplicates += 1
                        existing = target_ep["test_rules"][sig_map[sig]]
                        existing["source_document"] = _merge_api_source_document(existing.get("source_document"), rule.get("source_document"))
                        old_src = str(existing.get("source_requirement", "")).strip()
                        new_src = str(rule.get("source_requirement", "")).strip()
                        if new_src and _normalize_text(new_src) != _normalize_text(old_src):
                            existing["source_requirement"] = (old_src + " | " + new_src).strip(" |")
                        continue
                    sig_map[sig] = len(target_ep["test_rules"])
                    target_ep["test_rules"].append(copy.deepcopy(rule))

    final_modules = []
    global_rule_counter = 0
    unmapped_endpoints = 0
    for m_idx, module in enumerate(modules_map.values(), start=1):
        endpoints = list(module["endpoints"].values())
        for e_idx, ep in enumerate(endpoints, start=1):
            if ep.get("method") == "UNMAPPED" or ep.get("endpoint_path") == "UNMAPPED":
                unmapped_endpoints += 1
            method = ep.get("method", "API")
            slug = unicodedata.normalize("NFKD", ep.get("endpoint_path", "API")).encode("ascii", "ignore").decode("ascii")
            slug = re.sub(r"[^A-Za-z0-9]+", "_", slug).strip("_").upper()[-28:] or f"EP{e_idx:02d}"
            prefix = f"API_{method}_{slug}" if method != "UNMAPPED" else f"API_UNMAPPED_{m_idx:02d}_{e_idx:02d}"
            for local_idx, rule in enumerate(ep.get("test_rules", []), start=1):
                global_rule_counter += 1
                rule["rule_id"] = f"{prefix}-{local_idx:03d}"
        final_modules.append({"module_name": module["module_name"], "endpoints": endpoints})

    final = {"api_test_design_version": "3.2", "api_modules": final_modules}
    stats = {
        "input_matrices": len(matrices),
        "modules": len(final_modules),
        "endpoints": sum(len(m.get("endpoints", [])) for m in final_modules),
        "unmapped_endpoints": unmapped_endpoints,
        "raw_rules": raw_rules,
        "duplicates_removed": duplicates,
        "final_rules": global_rule_counter,
    }
    return final, stats


def _api_split_chunk_for_retry(chunk: dict) -> list[dict]:
    core = chunk.get("core_text", "")
    if len(core) <= API_AGENT1_MIN_RECURSIVE_CHARS:
        return []
    target = max(API_AGENT1_MIN_RECURSIVE_CHARS, min(6000, len(core) // 2))
    max_chars = max(target + 700, min(7600, len(core)))
    local = split_document_semantic(
        core,
        base_title=chunk.get("title", "API_RetryChunk"),
        target_chars=target,
        max_chars=max_chars,
        context_chars=min(API_AGENT1_CONTEXT_CHARS, 700),
    )
    if len(local) <= 1:
        mid = len(core) // 2
        cut = core.rfind("\n", 0, mid)
        if cut < API_AGENT1_MIN_RECURSIVE_CHARS:
            cut = core.rfind(" ", 0, mid)
        if cut < API_AGENT1_MIN_RECURSIVE_CHARS:
            cut = mid
        local = []
        for part in [core[:cut].strip(), core[cut:].strip()]:
            if part:
                local.append({
                    "title": chunk.get("title", "API_RetryChunk"),
                    "core_text": part,
                    "prev_context": "",
                    "next_context": "",
                    "char_count": len(part),
                    "depth": chunk.get("depth", 0) + 1,
                })
    children = []
    parent_id = chunk.get("chunk_id", "A")
    for idx, child in enumerate(local):
        child = dict(child)
        child["chunk_id"] = f"{parent_id}.{idx+1}"
        child["depth"] = chunk.get("depth", 0) + 1
        children.append(child)
    return children


def process_api_agent1_chunk_recursive(
    chunk: dict,
    endpoint_index: list[dict],
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    cache: dict | None = None,
    diagnostics: list | None = None,
) -> tuple[bool, list[dict]]:
    diagnostics = diagnostics if diagnostics is not None else []
    payload = build_api_agent1_chunk_payload(chunk, endpoint_index)
    key = _cache_key("api-agent1-v3.2", model, prompt_template, payload)
    if cache is not None and key in cache:
        diagnostics.append({"chunk_id": chunk.get("chunk_id"), "status": "CACHE_HIT", "chars": len(chunk.get("core_text", ""))})
        return True, [copy.deepcopy(cache[key])]

    call_result = None
    for attempt in range(API_AGENT1_API_RETRIES + 1):
        agent_name = f"API_AGENT1/{chunk.get('chunk_id')}"
        call_result = call_qwen_max_agent_detailed(
            content=payload,
            api_key=api_key,
            base_url=base_url,
            model=model,
            prompt_template=prompt_template,
            max_tokens=QWEN_MAX_OUTPUT_TOKENS,
            agent_name=agent_name,
        )
        if call_result.ok or call_result.finish_reason == "length":
            break
        if attempt < API_AGENT1_API_RETRIES:
            log_info(f"[{agent_name}] 🔁 Retry API lần {attempt+2}/{API_AGENT1_API_RETRIES+1}")
            time.sleep(2)

    assert call_result is not None
    reason_to_split = None
    parse_diag = None
    schema_diag = None
    parsed_json = None

    if call_result.finish_reason == "length":
        reason_to_split = "MAX_TOKENS"
    elif not call_result.ok:
        diagnostics.append({"chunk_id": chunk.get("chunk_id"), "status": "API_FAILED", "error": call_result.error})
        return False, []
    else:
        parsed_ok, parsed_json, parse_diag = extract_json_from_model_response(call_result.text)
        if not parsed_ok:
            reason_to_split = "JSON_PARSE_FAIL"
        else:
            schema_ok, schema_diag = validate_api_rule_matrix_schema(parsed_json, strict=False)
            if not schema_ok:
                reason_to_split = "SCHEMA_FAIL"
            else:
                modules = parsed_json.get("api_modules", [])
                rules = sum(len(ep.get("test_rules", [])) for m in modules for ep in m.get("endpoints", []))
                endpoints = sum(len(m.get("endpoints", [])) for m in modules)
                diagnostics.append({
                    "chunk_id": chunk.get("chunk_id"), "status": "OK",
                    "chars": len(chunk.get("core_text", "")), "depth": chunk.get("depth", 0),
                    "elapsed": round(call_result.elapsed, 2), "output_chars": len(call_result.text),
                    "endpoints": endpoints, "rules": rules,
                })
                if cache is not None:
                    cache[key] = copy.deepcopy(parsed_json)
                return True, [parsed_json]

    depth = int(chunk.get("depth", 0))
    can_split = depth < API_AGENT1_MAX_RECURSION_DEPTH and len(chunk.get("core_text", "")) > API_AGENT1_MIN_RECURSIVE_CHARS
    diagnostics.append({
        "chunk_id": chunk.get("chunk_id"), "status": "SPLIT_RETRY" if can_split else "FAILED_LEAF",
        "reason": reason_to_split, "parse_diag": parse_diag, "schema_diag": schema_diag,
        "finish_reason": call_result.finish_reason, "output_chars": len(call_result.text),
        "raw_tail": call_result.text[-1200:] if call_result.text else "",
    })
    if not can_split:
        return False, []
    children = _api_split_chunk_for_retry(chunk)
    if len(children) < 2:
        return False, []
    log_info(
        f"[API_AGENT1/{chunk.get('chunk_id')}] ✂️ Recursive split do {reason_to_split} | "
        f"{len(chunk.get('core_text', '')):,} chars -> " + " + ".join(f"{len(c.get('core_text', '')):,}" for c in children)
    )
    all_matrices = []
    for child in children:
        ok, mats = process_api_agent1_chunk_recursive(
            child, endpoint_index, api_key, base_url, model, prompt_template, cache, diagnostics
        )
        if not ok:
            return False, []
        all_matrices.extend(mats)
    return True, all_matrices


def run_api_agent1_document_pipeline(
    design_text: str,
    ba_text: str,
    base_filename: str,
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    cache: dict | None = None,
    progress_callback=None,
) -> tuple[bool, dict | None, dict]:
    started = time.time()
    combined = build_api_combined_source(design_text, ba_text)
    endpoint_index = extract_api_endpoint_index(design_text)
    initial_chunks = split_document_semantic(
        combined,
        base_title=base_filename,
        target_chars=API_AGENT1_CHUNK_TARGET_CHARS,
        max_chars=API_AGENT1_CHUNK_MAX_CHARS,
        context_chars=API_AGENT1_CONTEXT_CHARS,
    )
    diagnostics = []
    matrices = []
    log_info(
        f"[API_AGENT1_PIPELINE] 📚 DocumentChars={len(combined):,} | InitialChunks={len(initial_chunks)} | "
        f"EndpointIndex={len(endpoint_index)} | Target≈{API_AGENT1_CHUNK_TARGET_CHARS:,}"
    )
    for idx, chunk in enumerate(initial_chunks, start=1):
        if progress_callback:
            progress_callback(idx-1, len(initial_chunks), f"API Agent 1 chunk {idx}/{len(initial_chunks)} — {chunk.get('title')} ({chunk.get('char_count',0):,} chars)")
        ok, mats = process_api_agent1_chunk_recursive(
            chunk, endpoint_index, api_key, base_url, model, prompt_template, cache, diagnostics
        )
        if not ok:
            return False, None, {
                "ok": False, "document_chars": len(combined), "initial_chunks": len(initial_chunks),
                "endpoint_index": len(endpoint_index), "elapsed": round(time.time()-started,2), "diagnostics": diagnostics,
            }
        matrices.extend(mats)
        if progress_callback:
            progress_callback(idx, len(initial_chunks), f"Hoàn tất API chunk {idx}/{len(initial_chunks)}")

    final, merge_stats = merge_api_rule_matrices(matrices)
    schema_ok, schema_diag = validate_api_rule_matrix_schema(final, strict=True)
    summary = {
        "ok": schema_ok, "document_chars": len(combined), "initial_chunks": len(initial_chunks),
        "endpoint_index": len(endpoint_index), "leaf_matrices": len(matrices),
        "elapsed": round(time.time()-started,2), "merge": merge_stats, "schema": schema_diag,
        "diagnostics": diagnostics,
    }
    if not schema_ok:
        log_error(f"[API_AGENT1_PIPELINE] ❌ Final schema fail | {schema_diag}")
        return False, None, summary
    log_info(
        f"[API_AGENT1_PIPELINE] ✅ Completed | Time={summary['elapsed']:.2f}s | Modules={merge_stats['modules']} | "
        f"Endpoints={merge_stats['endpoints']} | FinalRules={merge_stats['final_rules']} | Unmapped={merge_stats['unmapped_endpoints']}"
    )
    return True, final, summary


def split_api_endpoint_rules_for_agent2(module_name: str, endpoint: dict) -> list[dict]:
    rules = endpoint.get("test_rules", [])
    batches, current, current_chars = [], [], 0
    for rule in rules:
        rc = len(json.dumps(rule, ensure_ascii=False))
        if current and (len(current) >= API_AGENT2_BATCH_MAX_RULES or current_chars + rc > API_AGENT2_BATCH_MAX_INPUT_CHARS):
            batches.append({
                "module_name": module_name,
                "endpoint": {
                    "method": endpoint.get("method", "UNMAPPED"),
                    "endpoint_path": endpoint.get("endpoint_path", "UNMAPPED"),
                    "summary": endpoint.get("summary", ""),
                    "test_rules": current,
                }
            })
            current, current_chars = [], 0
        current.append(rule)
        current_chars += rc
    if current:
        batches.append({
            "module_name": module_name,
            "endpoint": {
                "method": endpoint.get("method", "UNMAPPED"),
                "endpoint_path": endpoint.get("endpoint_path", "UNMAPPED"),
                "summary": endpoint.get("summary", ""),
                "test_rules": current,
            }
        })
    return batches


def render_api_agent2_batch_recursive(
    batch: dict,
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    batch_id: str,
    depth: int = 0,
    diagnostics: list | None = None,
    cache: dict | None = None,
) -> tuple[bool, list[str]]:
    diagnostics = diagnostics if diagnostics is not None else []
    payload = json.dumps({
        "api_test_design_version": "3.2",
        "api_modules": [{
            "module_name": batch.get("module_name", "UNMAPPED"),
            "endpoints": [batch.get("endpoint", {})],
        }]
    }, ensure_ascii=False)
    key = _cache_key("api-agent2-v3.2", model, prompt_template, payload)
    if cache is not None and key in cache:
        diagnostics.append({"batch_id": batch_id, "status": "CACHE_HIT", "rules": len(batch.get("endpoint", {}).get("test_rules", []))})
        return True, [cache[key]]

    call_result = None
    for attempt in range(API_AGENT2_API_RETRIES + 1):
        call_result = call_qwen_max_agent_detailed(
            content=payload, api_key=api_key, base_url=base_url, model=model,
            prompt_template=prompt_template, max_tokens=QWEN_MAX_OUTPUT_TOKENS,
            agent_name=f"API_AGENT2/{batch_id}",
        )
        if call_result.ok or call_result.finish_reason == "length":
            break
        if attempt < API_AGENT2_API_RETRIES:
            time.sleep(2)
    assert call_result is not None
    if call_result.complete:
        diagnostics.append({"batch_id": batch_id, "status": "OK", "rules": len(batch.get("endpoint", {}).get("test_rules", [])), "output_chars": len(call_result.text)})
        if cache is not None:
            cache[key] = call_result.text
        return True, [call_result.text]

    rules = batch.get("endpoint", {}).get("test_rules", [])
    can_split = call_result.finish_reason == "length" and len(rules) > 1 and depth < API_AGENT2_MAX_RECURSION_DEPTH
    diagnostics.append({"batch_id": batch_id, "status": "SPLIT_RETRY" if can_split else "FAILED", "reason": "MAX_TOKENS" if call_result.finish_reason == "length" else (call_result.error or "UNKNOWN"), "rules": len(rules)})
    if not can_split:
        return False, []
    mid = max(1, len(rules)//2)
    outputs = []
    for idx, child_rules in enumerate([rules[:mid], rules[mid:]], start=1):
        if not child_rules:
            continue
        child = copy.deepcopy(batch)
        child["endpoint"]["test_rules"] = child_rules
        ok, child_outputs = render_api_agent2_batch_recursive(
            child, api_key, base_url, model, prompt_template, f"{batch_id}.{idx}", depth+1, diagnostics, cache
        )
        if not ok:
            return False, []
        outputs.extend(child_outputs)
    return True, outputs


def _renumber_api_testcases_in_nodes(nodes: list[dict]) -> int:
    counter = 0
    def walk(children):
        nonlocal counter
        for node in children:
            title = str(node.get("title", ""))
            if re.match(r"^TC_API[_\- ]?\d+\s*[-:]", title, flags=re.IGNORECASE):
                counter += 1
                node["title"] = re.sub(r"^TC_API[_\- ]?\d+", f"TC_API_{counter:03d}", title, count=1, flags=re.IGNORECASE)
            walk(node.get("children", []))
    walk(nodes)
    return counter


def merge_api_agent2_tree_outputs(outputs: list[str]) -> tuple[str, dict]:
    merged_nodes = []
    for output in outputs:
        forest = _parse_bullet_forest(output)
        _merge_bullet_nodes(merged_nodes, forest, depth=0)
    tc_count = _renumber_api_testcases_in_nodes(merged_nodes)
    return "\n".join(_serialize_bullet_nodes(merged_nodes)).strip(), {
        "agent2_leaf_outputs": len(outputs), "testcases_renumbered": tc_count
    }


def run_api_agent2_rule_matrix_pipeline(
    approved_modules: list[dict],
    api_key: str,
    base_url: str,
    model: str,
    prompt_template: str,
    cache: dict | None = None,
    progress_callback=None,
) -> tuple[bool, str, dict]:
    started = time.time()
    initial_batches = []
    for midx, module in enumerate(approved_modules, start=1):
        for eidx, endpoint in enumerate(module.get("endpoints", []), start=1):
            for bidx, batch in enumerate(split_api_endpoint_rules_for_agent2(module.get("module_name", "UNMAPPED"), endpoint), start=1):
                initial_batches.append((f"M{midx:02d}E{eidx:03d}B{bidx:02d}", batch))
    diagnostics, outputs = [], []
    total_rules = sum(len(ep.get("test_rules", [])) for m in approved_modules for ep in m.get("endpoints", []))
    log_info(f"[API_AGENT2_PIPELINE] 🧩 Modules={len(approved_modules)} | Batches={len(initial_batches)} | Rules={total_rules}")
    for idx, (batch_id, batch) in enumerate(initial_batches, start=1):
        if progress_callback:
            ep = batch.get("endpoint", {})
            progress_callback(idx-1, len(initial_batches), f"API Agent 2 batch {idx}/{len(initial_batches)} — {ep.get('method')} {ep.get('endpoint_path')} ({len(ep.get('test_rules', []))} rules)")
        ok, batch_outputs = render_api_agent2_batch_recursive(
            batch, api_key, base_url, model, prompt_template, batch_id, diagnostics=diagnostics, cache=cache
        )
        if not ok:
            return False, "", {"ok": False, "initial_batches": len(initial_batches), "elapsed": round(time.time()-started,2), "diagnostics": diagnostics}
        outputs.extend(batch_outputs)
        if progress_callback:
            progress_callback(idx, len(initial_batches), f"Hoàn tất API batch {idx}/{len(initial_batches)}")
    merged_text, merge_stats = merge_api_agent2_tree_outputs(outputs)
    summary = {"ok": True, "initial_batches": len(initial_batches), "elapsed": round(time.time()-started,2), "merge": merge_stats, "diagnostics": diagnostics}
    log_info(f"[API_AGENT2_PIPELINE] ✅ Completed | Time={summary['elapsed']:.2f}s | TestCases={merge_stats['testcases_renumbered']}")
    return True, merged_text, summary


def convert_tree_to_api_excel(tree_text: str) -> bytes:
    """Excel API hỗ trợ nhiều Module/Endpoint; không bỏ các top-level module phía sau."""
    forest = _parse_bullet_forest(tree_text)
    wb = Workbook()
    ws = wb.active
    ws.title = "API Test Execution"
    ws.views.sheetView[0].showGridLines = True

    fill_green = PatternFill(start_color="81C784", fill_type="solid")
    fill_light_green = PatternFill(start_color="C8E6C9", fill_type="solid")
    fill_white = PatternFill(start_color="FFFFFF", fill_type="solid")
    fill_header_bg = PatternFill(start_color="4DD0E1", fill_type="solid")
    fill_pink = PatternFill(start_color="F8BBD0", fill_type="solid")
    fill_blue_exec = PatternFill(start_color="90CAF9", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='B0BEC5'), right=Side(style='thin', color='B0BEC5'),
        top=Side(style='thin', color='B0BEC5'), bottom=Side(style='thin', color='B0BEC5')
    )
    font_title = Font(name="Segoe UI", size=11, bold=True)
    font_bold = Font(name="Segoe UI", size=10, bold=True)
    font_body = Font(name="Segoe UI", size=10)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)

    ws['D1'] = "KỊCH BẢN KIỂM THỬ API *"
    ws['D1'].font = font_title
    ws['D1'].alignment = align_center
    ws['C2'] = "Phạm vi"
    ws['C2'].font = font_bold
    ws['D2'] = " / ".join([n.get("title", "") for n in forest[:3]]) if forest else "API Testing"
    for r in range(1, 5):
        for col in ["C", "D"]:
            ws[f"{col}{r}"].border = thin_border

    headers = [
        ("A11", "External ID", fill_header_bg), ("B11", "Name", fill_pink),
        ("C11", "PreConditions", fill_pink), ("D11", "Importance", fill_pink),
        ("E11", "Step", fill_pink), ("F11", "Data test", fill_pink),
        ("G11", "Expected Result", fill_pink), ("H11", "Actual Result", fill_pink),
        ("I11", "Lần 1", fill_blue_exec), ("J11", "Lần 2", fill_blue_exec),
    ]
    for pos, val, fill in headers:
        ws[pos] = val; ws[pos].font = font_bold; ws[pos].alignment = align_center; ws[pos].border = thin_border; ws[pos].fill = fill

    current_row = 14

    def clean_prefix(value, prefixes):
        t = str(value).strip()
        for p in prefixes:
            if t.lower().startswith(p.lower()):
                return t[len(p):].strip()
        return t

    def extract_details(node):
        pre = step = exp = ""
        for child in node.get("children", []):
            title = child.get("title", "")
            low = title.lower()
            if "pre-condition" in low or "precondition" in low or "tiền điều kiện" in low:
                pre = clean_prefix(title, ["Pre-condition:", "Precondition:", "Tiền điều kiện:"])
            elif "steps & data test" in low or low.startswith("steps") or "các bước" in low:
                step = clean_prefix(title, ["Steps & Data test:", "Steps:", "Các bước thực hiện:"])
            elif "response" in low or "kết quả mong đợi" in low or "expected" in low:
                exp = clean_prefix(title, ["Response (Kết quả mong đợi):", "Kết quả mong đợi:", "Expected Result:"])
            sp, ss, se = extract_details(child)
            pre = pre or sp; step = step or ss; exp = exp or se
        return pre, step, exp

    def is_tc(node):
        return bool(re.match(r"^TC_API[_\- ]?\d+", str(node.get("title", "")), flags=re.IGNORECASE))

    def has_tc(node):
        return is_tc(node) or any(has_tc(c) for c in node.get("children", []))

    def write_group(title, level):
        nonlocal current_row
        ws.cell(row=current_row, column=2, value=("  " * max(0, level-1)) + title).font = font_bold
        for col in range(1, 11):
            c = ws.cell(row=current_row, column=col)
            c.fill = fill_green if level <= 2 else fill_light_green
            c.border = thin_border
        current_row += 1

    def walk(node, level=1):
        nonlocal current_row
        if is_tc(node):
            full = node.get("title", "")
            parts = full.split(" - ", 1)
            tc_id = parts[0].strip()
            tc_name = parts[1].strip() if len(parts) > 1 else full
            pre, step, exp = extract_details(node)
            ws.cell(row=current_row, column=1, value=tc_id).alignment = align_center
            ws.cell(row=current_row, column=2, value=tc_name)
            ws.cell(row=current_row, column=3, value=pre)
            ws.cell(row=current_row, column=4, value=None)
            ws.cell(row=current_row, column=5, value=step)
            ws.cell(row=current_row, column=6, value=None)
            ws.cell(row=current_row, column=7, value=exp)
            for col in range(1, 11):
                c = ws.cell(row=current_row, column=col); c.font = font_body; c.fill = fill_white; c.border = thin_border
                if col not in [1,4]: c.alignment = align_left
            current_row += 1
            return
        if has_tc(node):
            write_group(str(node.get("title", "")), level)
            for child in node.get("children", []):
                walk(child, level+1)

    for root in forest:
        walk(root, 1)

    for col, width in {'A':18,'B':48,'C':34,'D':12,'E':52,'F':18,'G':52,'H':20,'I':10,'J':10}.items():
        ws.column_dimensions[col].width = width
    output = io.BytesIO(); wb.save(output); output.seek(0)
    return output.getvalue()


# ==========================================
# 3. KHAI BÁO PROMPTS WEB/APP
# ==========================================
# ==============================================================================
# 1. PROMPT DÀNH CHO BÓC TÁCH MỘT BƯỚC (DÙNG TRỰC TIẾP PROMPT_WEB_UI CHUẨN)
# ==============================================================================
# QA TEST DESIGN PIPELINE — PROMPT PACK v1
# Mục tiêu:
# Agent 1 = QA Brain: đọc Spec + áp dụng QA Test Design Rules -> Test Design / Rule Matrix
# Agent 2 = Renderer: chỉ biến Rule Matrix thành Test Case/XMind, không suy luận thêm
# Web = Renderer cuối nếu cần, không tự nghĩ Business Rule


# ============================================================
# 1. PROMPT AGENT 1 — EXTRACT + APPLY QA RULES
# ============================================================

# NEW 2 — QA TEST DESIGN PIPELINE v2

# ============================================================
# 1. AGENT 1 — QA BRAIN
# ============================================================

# ============================================================
# ============================================================
# WEB PROMPT PACK — OPTIMIZED B
#
# Multi-Agent flow:
#   WEB_AGENT1_QA_BRAIN -> Rule Matrix JSON
#   WEB_AGENT2_XMIND_RENDERER -> XMind plain-text tree
#
# Legacy / 1-Click flow:
#   WEB_DIRECT_TESTCASE_AGENT -> XMind plain-text tree trực tiếp
# ============================================================


# ============================================================
# 1. AGENT 1 — WEB_AGENT1_QA_BRAIN
# SRS CHUNK -> COMPACT RULE MATRIX
# ============================================================

PROMPT_AGENT1_EXTRACT_RULE_MATRIX = """
Bạn là WEB_AGENT1_QA_BRAIN — Senior QA Test Design Lead cho hệ thống Banking / Enterprise.

NHIỆM VỤ DUY NHẤT:
Đọc CURRENT SOURCE trong chunk, xác định requirement WEB có ý nghĩa kiểm thử và tạo TEST DESIGN / RULE MATRIX.
Bạn là Agent DUY NHẤT được phép phân tích requirement và áp dụng QA Test Design Technique.
WEB_AGENT2_XMIND_RENDERER phía sau CHỈ render, không được bổ sung Test Rule.

==================================================
1. CHUNK OWNERSHIP — CRITICAL
==================================================

Input có 3 vùng:
1. PREVIOUS CONTEXT: chỉ dùng để hiểu requirement ở biên.
2. CURRENT SOURCE: nguồn CHÍNH sở hữu Test Rule.
3. NEXT CONTEXT: chỉ dùng để hiểu requirement ở biên.

CHỈ tạo Rule khi target/action/constraint/behavior được kiểm thử có căn cứ trong CURRENT SOURCE.
Có thể dùng PREVIOUS/NEXT CONTEXT để hoàn thiện requirement thuộc CURRENT SOURCE.
KHÔNG tạo Rule nếu toàn bộ căn cứ chỉ nằm trong CONTEXT.

Đây là một chunk của tài liệu lớn:
- Coverage đầy đủ CURRENT SOURCE.
- Không cố đánh giá coverage toàn tài liệu.
- Không copy lại requirement đã thuộc chunk khác.
- Nếu cùng nội dung lặp do bảng/OCR/image text, chỉ giữ một objective.

==================================================
2. SCOPE — CHỈ TEST WEB
==================================================

Requirement WEB gồm:
- Screen/UI/Figma.
- Field/Textbox/Numeric/Dropdown/Datepicker/Checkbox.
- Button/Icon/Link/Action.
- Grid/Column/Pagination.
- Popup/Dialog/Toast.
- Search/Filter/Reset.
- Permission hiển thị/truy cập Web.
- Business Flow thực hiện từ giao diện.
- FE gọi API và FE xử lý response như một phần chức năng Web.

CRITICAL:
- API được nhắc trong SRS Web KHÔNG tự động trở thành API Test Scope.
- FE gọi API để Search/Dropdown/Hold/Confirm/Cancel/Load data... thuộc TÍNH NĂNG MÀN HÌNH.
- Không tự sinh Authentication/HTTP Method/Header/Schema test chỉ vì tài liệu có chữ API.
- Không lấy nội dung từ URL/“Request access” mà CURRENT SOURCE không cung cấp.
- Không bịa API/Status/Message/DB/Permission/Timeout/Edge Case.

==================================================
3. SCREEN OWNERSHIP
==================================================

Chỉ tạo screen khi có căn cứ rõ như:
- section riêng “Màn hình ...”;
- mã màn hình;
- bảng mô tả màn hình;
- mockup/control specification riêng.

KHÔNG tạo screen mới khi:
- chỉ được nhắc như màn hình đích điều hướng;
- chỉ xuất hiện trong link/reference;
- flow chỉ nói “chuyển sang màn hình X” nhưng không đặc tả màn X;
- cùng một màn hình được gọi bằng nhiều biến thể tên.

QUY TẮC ỔN ĐỊNH TÊN — BẮT BUỘC:
- Mã screen/module như BOND_ORDER_LIST chỉ là ID kỹ thuật để nhận diện cùng một màn hình, KHÔNG phải tên hiển thị bắt buộc.
- screen_name phải là TÊN NGHIỆP VỤ DỄ ĐỌC của màn hình, ưu tiên heading/definition chính trong Spec.
- Ví dụ:
  + screen code: BOND_ORDER_LIST
  + screen_name: "Danh sách lệnh đặt mua Trái phiếu"
  => đây là CÙNG MỘT MÀN HÌNH, KHÔNG được tạo hai screen.
- Sub-heading như "Mô tả màn hình", "Logic tìm kiếm", "Hold lại tiền"... KHÔNG phải screen mới.
- Nếu CURRENT SOURCE là phần tiếp nối của screen đang có trong context, BẮT BUỘC kế thừa cùng màn hình đó.
- Các biến thể:
  + "BOND_ORDER_LIST"
  + "Danh sách lệnh đặt mua"
  + "Màn hình Danh sách lệnh đặt mua"
  + "Danh sách lệnh đặt mua Trái phiếu"
  + "Màn hình Danh sách lệnh đặt mua Trái phiếu"
  nếu cùng business function thì phải coi là MỘT SCREEN.
- Không tạo screen root cho màn hình chỉ được nhắc là destination/reference.
- Nếu cần dùng screen code, chỉ dùng làm prefix rule_id/identity nội bộ; không dùng nó để tách thêm screen.

Mục tiêu: một màn hình nghiệp vụ chỉ có MỘT screen root dễ đọc xuyên suốt tất cả chunk.

==================================================
4. 5 CATEGORY NỘI BỘ — OWNERSHIP
==================================================

category CHỈ được là:
UI | VALIDATION | ACTION | BUSINESS_FLOW | EXCEPTION

KHÔNG output GRID hoặc POPUP như category riêng.
Grid/Popup/Search/Pagination chỉ là target/sub-feature.

------------------------------
4.1 UI
------------------------------
UI CHỈ kiểm tra PHẦN HIỂN THỊ TĨNH / PRESENTATION của màn hình:
- tiêu đề màn hình;
- breadcrumb;
- static label/text;
- bố cục/section;
- element tồn tại/visibility thuần túy;
- icon hình thức;
- grid header/column label;
- format hiển thị của dữ liệu READ-ONLY;
- width/ellipsis/tooltip của dữ liệu READ-ONLY;
- popup title/content/icon;
- empty/loading presentation nếu Spec mô tả.

TUYỆT ĐỐI KHÔNG XẾP VÀO UI đối với INPUT CONTROL:
- Placeholder;
- Default Value;
- Default Selection;
- giá trị "Tất cả" mặc định;
- danh sách option của Dropdown/Combobox;
- thứ tự option;
- selected value hiển thị sau khi chọn;
- Enable/Disable;
- Readonly/Editability;
- Required;
- input format/character/length;
- search bên trong Dropdown.

TẤT CẢ các mục trên của Field/Input Control phải thuộc VALIDATION.

COMPACT:
- Các static UI element chỉ cần “hiển thị đúng” và không có rule riêng có thể gom thành 1 objective UI tổng thể.
- Không sinh từng testcase UI chỉ để kiểm tra Placeholder/Default của từng Field.

Không đưa Validation/Click/Search/Business result/API result/Exception vào UI.

------------------------------
4.2 VALIDATION
------------------------------
VALIDATION sở hữu TOÀN BỘ behavior của Field/Input Control, kể cả trạng thái ban đầu.

BẮT BUỘC đưa vào VALIDATION nếu Spec mô tả:
- Placeholder;
- Default Value;
- Default Selection;
- giá trị mặc định "Tất cả";
- Required;
- Enable/Disable của Field;
- Readonly/Editability của Field;
- length/min/max;
- character/type constraint;
- format nhập;
- danh sách option của Dropdown/Combobox;
- thứ tự option;
- Single/Multi Select;
- Select All / bỏ Select All;
- Clear Selection;
- selected value hiển thị sau khi chọn;
- search bên trong Dropdown;
- dropdown dependency;
- date relation;
- field dependency.

Ví dụ:
"CIF placeholder = Nhập số CIF" -> VALIDATION.
"Dropdown Trái phiếu mặc định = Tất cả" -> VALIDATION.
"Dropdown hiển thị danh sách <Mã TP> - <Kỳ hạn>" -> VALIDATION.
"Chọn nhiều giá trị và hiển thị +N" -> VALIDATION.

Không biến Business Rule sau Submit thành Validation.
Không coi open dropdown/date picker là ACTION.

------------------------------
4.3 ACTION
------------------------------
Áp dụng cho Button/Icon/Link/Row Action/Button trong Popup.

Chỉ kiểm tra hành vi TRỰC TIẾP:
- label/icon/tooltip nếu có rule;
- visibility theo condition;
- enable/disable;
- click;
- mở/đóng popup;
- navigation trực tiếp;
- trigger đúng action/API nếu Spec mô tả.

Không kiểm tra tại ACTION:
- API response;
- success/error message sau xử lý;
- DB;
- business outcome.

Kết quả sau action thuộc BUSINESS_FLOW hoặc EXCEPTION.

------------------------------
4.4 BUSINESS_FLOW
------------------------------
Đây là “TÍNH NĂNG MÀN HÌNH” ở output cuối.

Bao gồm:
- Search/Filter/Reset.
- Pagination.
- Grid data mapping/count/refresh.
- Business Rule.
- State transition.
- Permission làm điều kiện thực hiện chức năng.
- FE request/parameter tới API.
- Successful response -> UI.
- Success toast/message.
- Điều hướng sau xử lý.
- Hold/Confirm/Cancel/Copy/Edit/Submit.
- Kết quả nghiệp vụ sau Button/Popup Confirm.

Grid:
- display/column/format -> UI.
- pagination/mapping/count/data result/refresh -> BUSINESS_FLOW.
- row button -> ACTION.

Popup:
- title/content/icon -> UI.
- Close/Confirm click behavior -> ACTION.
- kết quả nghiệp vụ sau Confirm -> BUSINESS_FLOW.

------------------------------
4.5 EXCEPTION
------------------------------
Chỉ abnormal/technical/error path có căn cứ:
- no permission + error;
- timeout;
- no response;
- server/system error;
- explicit error response;
- network error;
- duplicate/technical failure nếu Spec mô tả.

Business rejection bình thường theo rule -> BUSINESS_FLOW.
Technical failure -> EXCEPTION.

Không tự bổ sung exception.

==================================================
5. EXPLICIT / DERIVED
==================================================

EXPLICIT:
Behavior/rule được Spec mô tả trực tiếp.
applied_qa_rule = "EXPLICIT FROM SPEC".
generation_reason = "".

Ví dụ:
- “CIF chỉ nhập số”.
- “Độ dài CIF = 10”.
- “Mặc định = Tất cả”.
- “Timeout hiển thị message X”.

DERIVED:
Chỉ được suy ra TEST DATA/BOUNDARY từ constraint có thật trong Spec.

Cho phép:
- Length = N -> Boundary N-1/N/N+1.
- Min/Max -> Boundary quanh Min/Max.
- Required -> missing/empty phù hợp loại field.
- Allowed character/type -> valid + value trái constraint.
- Format/pattern -> đúng + sai format.
- Enum -> hợp lệ + ngoài tập nếu input thực tế có thể nhận giá trị ngoài tập.

DERIVED bắt buộc:
- source_requirement có constraint gốc.
- applied_qa_rule ghi đúng technique, ví dụ "Boundary Value Analysis".
- generation_reason cực ngắn, ví dụ "Derived từ Max Length=10".

KHÔNG DERIVE:
- Business Rule mới;
- API mới;
- Message/Status mới;
- Permission mới;
- DB behavior mới;
- Exception mới.

==================================================
6. ATOMICITY — ĐỦ NHƯNG KHÔNG SPAM
==================================================

Một Rule = một test objective có thể fail độc lập.

Tách khi khác:
- condition;
- expected behavior;
- target;
- parameter;
- business branch.

Ví dụ dropdown:
API load list / default / Select All / bỏ Select All / search
là các objective độc lập nếu Spec mô tả độc lập.

KHÔNG over-split:
- Một boundary set cùng objective có thể là MỘT Rule chứa N-1/N/N+1.
- Nhiều static UI element chỉ cần display đúng có thể là MỘT Rule UI tổng thể.

Không dùng “tương tự”, “các field khác”, “các case khác” để thay thế rule cần thiết.

==================================================
7. REQUIREMENT PRESERVATION — COMPACT
==================================================

source_requirement phải NGẮN nhưng giữ đủ dữ liệu làm thay đổi Test Case:
- default;
- min/max/length;
- condition;
- role/permission;
- state;
- request parameter/value;
- response/message;
- dependency;
- time/debounce;
- select/unselect behavior;
- format;
- mapping.

KHÔNG copy nguyên paragraph dài nếu có thể rút ngắn mà không mất meaning.

Ví dụ tốt:
"Dropdown Chi nhánh gọi API với searchStr=null, type=0, status=Active, page=null."

Ví dụ không đạt:
"Kiểm tra API của dropdown."

==================================================
8. DUPLICATE CONTROL
==================================================

Không tạo hai Rule có cùng:
- target;
- objective;
- condition;
- expected behavior.

Nếu cùng requirement xuất hiện ở nhiều bảng/OCR/section:
- hợp nhất thông tin;
- chỉ tạo một Rule cho cùng objective.

Một Button có thể có nhiều Rule KHÔNG duplicate nếu objective khác:
- visibility condition -> ACTION;
- click mở popup -> ACTION;
- confirm gửi request -> BUSINESS_FLOW;
- success reload/toast -> BUSINESS_FLOW;
- timeout -> EXCEPTION.

==================================================
9. RULE ID
==================================================

rule_id chỉ cần unique trong response hiện tại.
Nếu có screen code, ưu tiên prefix code.
Python sẽ renumber sau merge.

==================================================
10. LOCAL QUALITY GATE
==================================================

Trước output, kiểm tra CURRENT SOURCE:
1. Requirement có ý nghĩa kiểm thử đã được map chưa?
2. Mỗi Rule thuộc đúng 1 category?
3. Có duplicate objective không?
4. Có mất parameter/value/condition/message/default/format không?
5. DERIVED có constraint gốc không?
6. Có Rule nào chỉ từ best practice không?
7. Có tạo screen chỉ vì reference/navigation không?

Không bịa Rule để ép coverage.

==================================================
11. OUTPUT JSON — BẮT BUỘC
==================================================

CHỈ trả JSON hợp lệ.
KHÔNG markdown.
KHÔNG giải thích.
KHÔNG text trước/sau JSON.

Root bắt buộc:

{
  "test_design_version": "3.1",
  "screens": [
    {
      "screen_name": "",
      "test_rules": [
        {
          "rule_id": "",
          "target": "",
          "category": "UI | VALIDATION | ACTION | BUSINESS_FLOW | EXCEPTION",
          "rule_type": "EXPLICIT | DERIVED",
          "rule_name": "",
          "test_objective": "",
          "source_requirement": "",
          "applied_qa_rule": "",
          "generation_reason": ""
        }
      ]
    }
  ]
}

Nếu CURRENT SOURCE không có requirement đủ căn cứ:
{"test_design_version":"3.1","screens":[]}

==================================================
CHUNK SOURCE
==================================================

{content}
"""


# ============================================================
# 2. AGENT 2 — WEB_AGENT2_XMIND_RENDERER
# RULE MATRIX -> XMIND TREE
# ============================================================

PROMPT_AGENT2_GEN_XMIND_FROM_RULE_MATRIX = """
Bạn là WEB_AGENT2_XMIND_RENDERER — Test Case Renderer.

Bạn KHÔNG phải QA Analyst.
Bạn KHÔNG phân tích lại SRS.
Bạn KHÔNG bổ sung Test Rule.
Bạn KHÔNG mở rộng scope.

NHIỆM VỤ DUY NHẤT:
Chuyển từng test_rule trong TEST DESIGN / RULE MATRIX thành đúng 1 Test Case Web dạng cây XMind.

==================================================
1. SOURCE OF TRUTH
==================================================

Rule Matrix là SOURCE OF TRUTH DUY NHẤT.

Nếu có N test_rule hợp lệ -> phải có đúng N Test Case.

KHÔNG:
- thêm Rule/Test Case;
- bỏ Rule/Test Case;
- gộp Rule;
- đổi meaning;
- tự thêm Validation;
- tự thêm Business Rule;
- tự thêm API/Status/Message/DB;
- tự thêm Permission/Exception/Test Data không có căn cứ.

EXPLICIT -> render.
DERIVED -> render.
INFERRED -> không render.

==================================================
2. 5 NHÓM OUTPUT CỐ ĐỊNH
==================================================

Chỉ được tạo 5 category top-level và đúng thứ tự:

1. KIỂM TRA UI
2. KIỂM TRA VALIDATION
3. KIỂM TRA CHỨC NĂNG BUTTON / ACTION
4. KIỂM TRA TÍNH NĂNG MÀN HÌNH
5. KIỂM TRA NGOẠI LỆ

Mapping:
- UI -> 1. KIỂM TRA UI
- VALIDATION -> 2. KIỂM TRA VALIDATION
- ACTION -> 3. KIỂM TRA CHỨC NĂNG BUTTON / ACTION
- BUSINESS_FLOW -> 4. KIỂM TRA TÍNH NĂNG MÀN HÌNH
- EXCEPTION -> 5. KIỂM TRA NGOẠI LỆ

Chỉ tạo category có Test Case.

TƯƠNG THÍCH LEGACY nếu matrix cũ còn GRID/POPUP:
- GRID display/column/format -> UI; functional grid/pagination/mapping -> TÍNH NĂNG.
- POPUP title/content -> UI; direct button behavior -> ACTION; business result -> TÍNH NĂNG.
Không tạo top-level “Data Grid” hoặc “Popup”.

==================================================
3. CÁCH VIẾT TEST CASE
==================================================

Dùng:
- target;
- rule_name;
- test_objective;
- source_requirement;
- applied_qa_rule.

Tên Test Case:
TC_(STT) - [Target] - [Objective ngắn]

KHÔNG lặp screen_name trong tên Test Case vì Screen đã là Root.

Pre-condition:
- Chỉ ghi khi cần condition/role/state/dependency.
- Nếu không cần, có thể bỏ node Pre-condition.

Steps:
- Ngắn, executable.
- Dùng đúng Field/Button/Value/Condition từ Rule Matrix.
- Không copy nguyên source_requirement thành Steps.

Expected:
- Cụ thể, kiểm chứng được.
- Giữ chính xác message/value/state/format nếu Matrix có.
- Nếu Matrix không có thông tin cụ thể thì không tự bịa.

==================================================
4. CẤU TRÚC XMIND
==================================================

- [screen_name]
  - 1. KIỂM TRA UI
    - [Target]
      - TC_001 - [Target] - [Objective]
        - Pre-condition: ...
          - Các bước thực hiện: 1. ... -> 2. ... -> 3. ...
            - Kết quả mong đợi: 1. ... -> 2. ... -> 3. ...
  - 2. KIỂM TRA VALIDATION
    - ...
  - 3. KIỂM TRA CHỨC NĂNG BUTTON / ACTION
    - ...
  - 4. KIỂM TRA TÍNH NĂNG MÀN HÌNH
    - ...
  - 5. KIỂM TRA NGOẠI LỆ
    - ...

Cấp:
1. Screen
2. Category
3. Target
4. Test Case
5. Pre-condition (nếu có)
6. Các bước thực hiện
7. Kết quả mong đợi

Nếu không có Pre-condition:
Các bước thực hiện là child trực tiếp của Test Case.

Steps và Expected Result:
- mỗi loại nằm trên MỘT node;
- không Enter tạo node con cho từng step;
- dùng "->" để nối.

==================================================
5. OUTPUT DISCIPLINE
==================================================

CHỈ trả plain text dạng cây bằng dấu "-".
KHÔNG JSON.
KHÔNG markdown code block.
KHÔNG giải thích.
KHÔNG summary/statistics.
KHÔNG thêm/bỏ testcase.
Tất cả bằng tiếng Việt.

=== TEST DESIGN / RULE MATRIX JSON ===
{content}
"""


# ============================================================
# 3. LEGACY / 1-CLICK — WEB_DIRECT_TESTCASE_AGENT
# RAW WEB SRS CHUNK -> XMIND TREE
#
# Giữ tên biến PROMPT_WEB_UI để code cũ không phải sửa.
# ============================================================

PROMPT_WEB_UI = """
Bạn là WEB_DIRECT_TESTCASE_AGENT — Senior QA Lead sinh trực tiếp Test Case Web từ SRS/Markdown.

Đây là DIRECT GENERATOR cho luồng 1-Click.
Bạn đọc SOURCE CHUNK và sinh trực tiếp Test Case XMind.
KHÔNG output Rule Matrix/JSON trung gian.

==================================================
1. SOURCE OF TRUTH
==================================================

Chỉ dùng nội dung SOURCE CHUNK và context thực tế có trong chunk.

KHÔNG tự tạo:
Business Rule, Validation, API behavior, Status, Message, DB, Permission, Timeout, Edge Case.

API được nhắc trong SRS Web chỉ là một phần behavior của chức năng Web.
Không tự mở rộng thành API Test Suite.

Nếu nội dung lặp do bảng/OCR/image text:
- tổng hợp;
- không duplicate.

Nếu nội dung bị đánh dấu đã bỏ/xóa/strikethrough:
- không sinh Test Case, trừ khi section active khác mô tả lại behavior hiện hành.

==================================================
2. SCREEN
==================================================

Chỉ tạo Screen Root khi chunk có căn cứ rõ về màn hình.

Nếu có screen code:
- coi screen code là ID kỹ thuật để nhận biết cùng màn hình;
- Screen Root vẫn dùng tên nghiệp vụ dễ đọc từ heading/definition chính;
- không tạo thêm Root chỉ vì code khác cách viết với tên màn hình.

Nếu không có screen code:
- dùng heading/definition chính của màn hình;
- không dùng sub-heading/feature name làm Screen Root.

Ví dụ:
BOND_ORDER_LIST = Danh sách lệnh đặt mua = Màn hình Danh sách lệnh đặt mua Trái phiếu
nếu cùng spec thì chỉ là MỘT Root.

Không tạo screen mới chỉ vì:
- navigation tới màn khác;
- link/reference;
- tên màn hình đích được nhắc trong flow;
- sub-section Search/Hold/Confirm/Cancel.

==================================================
3. 5 NHÓM DUY NHẤT
==================================================

1. KIỂM TRA UI
2. KIỂM TRA VALIDATION
3. KIỂM TRA CHỨC NĂNG BUTTON / ACTION
4. KIỂM TRA TÍNH NĂNG MÀN HÌNH
5. KIỂM TRA NGOẠI LỆ

UI:
CHỈ static presentation: screen title/breadcrumb/static label/layout/grid header/read-only display/
tooltip/ellipsis/popup static content.

KHÔNG đưa Placeholder/Default Value/Default Selection/Dropdown option/Field Enable-Disable/
Field Readonly vào UI.

VALIDATION:
sở hữu toàn bộ behavior của Field/Input Control:
Placeholder, Default Value, Default Selection, Required, Enable/Disable, Readonly,
length/type/format, Dropdown option/order/select-all/multi-select/selected display/search/dependency,
Date relation và behavior trực tiếp khi nhập/chọn field.

ACTION:
visibility/enable/click/open-close popup/navigation/trigger trực tiếp.

TÍNH NĂNG:
search/filter/reset/pagination/grid data/business rule/FE API request-success response/
state transition/success toast/navigation sau xử lý.

NGOẠI LỆ:
timeout/no response/server error/permission error/technical failure có trong Spec.

Không tạo Data Grid hoặc Popup thành category top-level riêng.

==================================================
4. QA DERIVATION
==================================================

Được phép derive CHỈ từ constraint có thật:
- Length N -> N-1/N/N+1.
- Min/Max -> boundary.
- Required -> missing/empty phù hợp.
- Character/type -> valid + trái constraint.
- Format -> đúng/sai format.

Không derive Business Rule/API/Message/Permission/Exception.

Một boundary set cùng objective có thể nằm trong một Test Case để tránh spam.

==================================================
5. COVERAGE / DUPLICATE
==================================================

Cover requirement có ý nghĩa kiểm thử trong SOURCE CHUNK.
Một objective độc lập -> một Test Case.
Không dùng “tương tự/các case khác” để bỏ case.
Không tạo hai Test Case cùng target + objective + condition + expected.

Static UI chỉ cần display đúng có thể gom thành một Test Case tổng thể.
Rule riêng như format/tooltip/width/condition phải tách.

==================================================
6. CẤU TRÚC XMIND
==================================================

- [Screen Root]
  - 1. KIỂM TRA UI
    - [Target]
      - TC_001 - [Target] - [Objective]
        - Pre-condition: ...
          - Các bước thực hiện: 1. ... -> 2. ... -> 3. ...
            - Kết quả mong đợi: 1. ... -> 2. ... -> 3. ...
  - 2. KIỂM TRA VALIDATION
  - 3. KIỂM TRA CHỨC NĂNG BUTTON / ACTION
  - 4. KIỂM TRA TÍNH NĂNG MÀN HÌNH
  - 5. KIỂM TRA NGOẠI LỆ

Chỉ tạo category có Test Case.
Giữ đúng thứ tự 1 -> 5.

Tên Test Case:
TC_(STT) - [Target] - [Objective]

Không lặp Screen Name trong tên TC.

Pre-condition chỉ tạo khi cần.
Nếu không có Pre-condition, Steps là child trực tiếp của Test Case.

Steps/Expected:
- mỗi loại một node;
- không tách node từng bước;
- nối bằng "->".

==================================================
7. OUTPUT
==================================================

CHỈ plain text cây thụt lề bằng "-".
KHÔNG JSON.
KHÔNG markdown block.
KHÔNG giải thích.
KHÔNG summary/statistics.
Tất cả bằng tiếng Việt.

==================================================
SOURCE CHUNK
==================================================

{content}
"""


# ==========================================
# 4. PROMPTS API V3.2 — AGENT 1 QA BRAIN / AGENT 2 RENDERER
# ==========================================

PROMPT_API_AGENT1_RULE_MATRIX = """
Bạn là SENIOR API QA TEST DESIGN ENGINE / API QA BRAIN cho hệ thống Banking / Enterprise.

NHIỆM VỤ DUY NHẤT:
Đọc CURRENT SOURCE trong chunk và tạo API TEST DESIGN / RULE MATRIX.
Bạn là Agent DUY NHẤT được phép áp dụng API QA Test Design Rules.
Agent 2 phía sau CHỈ render Rule Matrix, không phân tích lại API Spec/BA.

==================================================
=== CHUNK RULE — CRITICAL
==================================================
Input có:
- ENDPOINT INDEX HINT: metadata Python trích từ API design, chỉ giúp định danh endpoint.
- PREVIOUS CONTEXT / NEXT CONTEXT: chỉ giúp hiểu requirement ở biên.
- CURRENT SOURCE: nguồn CHÍNH được phép sinh Test Rule.

CHỈ tạo Rule khi requirement/contract/constraint/behavior có căn cứ trong CURRENT SOURCE.
Được dùng CONTEXT để hoàn thiện ý nghĩa requirement nằm trong CURRENT SOURCE.
KHÔNG tạo Rule chỉ từ ENDPOINT INDEX HINT hoặc CONTEXT.

Nếu CURRENT SOURCE mô tả Business Rule nhưng không xác định được endpoint một cách chắc chắn:
- method = "UNMAPPED"
- endpoint_path = "UNMAPPED"
- KHÔNG đoán endpoint.

Đây là một chunk của tài liệu lớn. Coverage đầy đủ CURRENT SOURCE rồi kết thúc; không cố review toàn tài liệu.

==================================================
=== SOURCE PRIORITY / TRACEABILITY
==================================================
API_SPEC là source of truth cho:
- endpoint path / HTTP method
- security scheme nếu có
- header/query/path/body schema
- required/nullable/type/format/pattern/enum/min/max/length/items
- request/response schema
- documented status code / error code / message

BA là source of truth cho:
- business flow
- business condition
- permission/role nếu BA mô tả
- state transition
- data dependency
- integration behavior / downstream mapping
- business error/message nếu BA mô tả

Nếu cùng một rule có căn cứ ở cả hai tài liệu: source_document = "BOTH".
Không tự sửa mâu thuẫn giữa API_SPEC và BA. Nếu mâu thuẫn tạo ra hai expectation khác nhau, giữ traceability rõ để QA Lead review; không tự chọn một bên.

==================================================
=== ATOMIC RULE — KHÔNG TÓM TẮT
==================================================
1 independent test objective = 1 test_rule.

KHÔNG gom nhiều objective độc lập thành các câu mơ hồ như:
- "Kiểm tra validate request"
- "Kiểm tra các field"
- "Kiểm tra các status code"
- "Kiểm tra business rule"

Ví dụ endpoint có customerId required + maxLength 10:
- missing customerId là một DERIVED rule
- boundary maxLength 10 là một DERIVED rule khác
Không gộp nếu hai objective cần test độc lập.

Nhưng không over-split cùng một objective chỉ để tăng số rule.

==================================================
=== EXPLICIT / DERIVED
==================================================
EXPLICIT:
Requirement/behavior/validation/status/message được tài liệu mô tả trực tiếp.
applied_qa_rule = "EXPLICIT FROM SPEC" hoặc "EXPLICIT FROM BA".

DERIVED:
Contract/requirement cụ thể + một QA Rule cho phép bên dưới => Test Rule.
DERIVED bắt buộc có source_requirement + applied_qa_rule + generation_reason.

INFERRED / BEST PRACTICE không có căn cứ => KHÔNG output.

==================================================
=== API QA RULE LIBRARY — CHỈ DÙNG CÁC RULE NÀY
==================================================
1. METHOD / URL
- Documented Method/Path
- Wrong Method Negative: chỉ DERIVE khi method/path được định nghĩa rõ. Không tự khẳng định 405 nếu Spec không nói.
- Unknown/Wrong Path: chỉ tạo khi routing/not-found behavior hoặc status được tài liệu mô tả.

2. AUTHENTICATION
Chỉ áp dụng khi endpoint/global security xác định API cần authentication.
Có thể DERIVE:
- Missing Credential
- Invalid Credential / malformed credential
- Expired Credential chỉ khi Bearer/JWT/OAuth/token lifecycle có căn cứ
KHÔNG tự bịa 401/message nếu tài liệu không cung cấp.

3. AUTHORIZATION
Chỉ tạo khi role/scope/permission/branch/unit ownership được tài liệu mô tả.
Không tự tạo role matrix từ kiến thức chung.

4. REQUEST FIELD / PARAMETER
Dùng đúng location: header | path | query | body.
- required=true / required schema => Missing Field/Parameter
- nullable=false hoặc rule cấm null => Null invalid
- data type => Wrong Type
- minLength/maxLength => Boundary Value
- minimum/maximum/exclusiveMinimum/exclusiveMaximum => Numeric Boundary
- pattern/format => Valid/Invalid Format
- enum/allowed values => In Enum / Outside Enum
- array minItems/maxItems/uniqueItems/item type => Array Boundary / Duplicate / Wrong Item Type
- nested object required property => Missing Nested Required Field

KHÔNG tự tạo empty/whitespace/special char/unicode/trim nếu schema hoặc BA không cung cấp căn cứ phù hợp.
Required KHÔNG tự động đồng nghĩa null/empty invalid nếu contract không nói.

5. HAPPY PATH
Mỗi success scenario khác nhau được Spec/BA mô tả = một rule riêng.
Giữ nguyên documented request condition, status code, message, response behavior.

6. RESPONSE VALIDATION
Chỉ tạo từ response contract được tài liệu mô tả:
- documented status code
- response schema / required response field
- response field format/mapping
- business code/message
Không tự tạo status code ngoài tài liệu.

7. BUSINESS RULE
Mỗi business condition / state / dependency / permission đặc thù = rule riêng.
Không nhét business rule đặc thù vào Happy Path nếu nó cần objective riêng.

8. INTEGRATION
Chỉ khi tài liệu mô tả downstream/upstream system:
- request/response mapping
- downstream code mapping
- retry/timeout/fallback nếu được mô tả
- DB update/query nếu được mô tả
Không tự tạo behavior hệ thống ngoài tài liệu.

9. EXCEPTION
Chỉ khi tài liệu có căn cứ:
- timeout
- network/system error
- payload/file limit
- database error
- documented 4xx/5xx
- downstream exception
Không mặc định sinh 500/timeout/database error.

==================================================
=== PRESERVE CONTRACT DETAILS — CRITICAL
==================================================
Không được làm mất thông tin có giá trị kiểm thử.
Nếu source có các giá trị sau phải giữ trong source_requirement/test_condition/expected_result tương ứng:
- endpoint + method
- header/parameter/body field name + location
- required/nullable/type
- min/max/length/boundary
- pattern/format/enum
- exact status code
- exact business/error code
- exact message
- request/response mapping
- state/role/permission
- downstream code
- DB behavior

Không biến:
"codeType=BOND_TYPE, status=null"
thành:
"Kiểm tra API danh mục".

==================================================
=== EXPECTED RESULT — KHÔNG BỊA
==================================================
Mỗi rule phải có expected_result đủ để Agent 2 render.
- Nếu tài liệu có exact status/message/body => dùng đúng.
- Nếu DERIVED từ schema nhưng tài liệu không nêu status/message cụ thể => chỉ mô tả expectation ở mức contract, KHÔNG tự thêm code/message.
Ví dụ: "Request không thỏa contract vì thiếu field bắt buộc customerId; không khẳng định status code do source không cung cấp."

==================================================
=== OUTPUT ROOT SCHEMA — BẮT BUỘC
==================================================
CHỈ trả JSON hợp lệ. KHÔNG markdown fence. KHÔNG text trước/sau JSON.
Root bắt buộc là object có "api_modules" array.
KHÔNG trả endpoint object hoặc rule object ở root.

{
  "api_test_design_version": "3.2",
  "api_modules": [
    {
      "module_name": "",
      "endpoints": [
        {
          "method": "GET | POST | PUT | PATCH | DELETE | HEAD | OPTIONS | UNMAPPED",
          "endpoint_path": "",
          "summary": "",
          "test_rules": [
            {
              "rule_id": "TMP-001",
              "target": "",
              "category": "AUTHENTICATION | AUTHORIZATION | METHOD_URL | REQUEST_VALIDATION | HAPPY_PATH | BUSINESS_RULE | RESPONSE_VALIDATION | INTEGRATION | EXCEPTION",
              "rule_type": "EXPLICIT | DERIVED",
              "rule_name": "",
              "test_objective": "",
              "test_condition": "",
              "expected_result": "",
              "source_requirement": "",
              "source_document": "API_SPEC | BA | BOTH",
              "applied_qa_rule": "",
              "generation_reason": ""
            }
          ]
        }
      ]
    }
  ]
}

Nếu CURRENT SOURCE không có requirement có ý nghĩa kiểm thử:
{"api_test_design_version":"3.2","api_modules":[]}

==================================================
=== FINAL CHECK TRƯỚC OUTPUT
==================================================
- Đã coverage mọi contract/requirement có ý nghĩa kiểm thử trong CURRENT SOURCE?
- Có bỏ parameter/header/body/status/business condition quan trọng không?
- Có nén nhiều objective độc lập thành 1 rule không?
- EXPLICIT/DERIVED đúng chưa?
- DERIVED có applied_qa_rule chưa?
- Có tự bịa 401/403/404/405/500/message/timeout/DB behavior không?
- Có rule chỉ từ Endpoint Index Hint/Context không?
- Root có đúng api_modules array không?

=== INPUT CHUNK ===
{content}
"""


PROMPT_API_AGENT2_RENDERER = """
Bạn là API TEST CASE RENDERER.
Bạn KHÔNG phải API QA Analyst.
Bạn KHÔNG đọc lại API Spec/BA và KHÔNG áp dụng thêm QA Rule.

NHIỆM VỤ DUY NHẤT:
Chuyển từng test_rule trong API TEST DESIGN / RULE MATRIX thành đúng 1 API Test Case dạng cây.

==================================================
=== SOURCE OF TRUTH — CRITICAL
==================================================
Rule Matrix là source of truth DUY NHẤT.
KHÔNG:
- thêm/bỏ/gộp test_rule
- tự tạo auth case
- tự tạo wrong method/url
- tự tạo validation
- tự tạo status code/message
- tự tạo business rule
- tự tạo DB/integration behavior
- tự tạo exception

Nếu input có N test_rule => output đúng N testcase.

==================================================
=== MAPPING 1 RULE = 1 TEST CASE
==================================================
Dùng đúng:
- module_name
- method
- endpoint_path
- summary
- target
- category
- rule_name
- test_objective
- test_condition
- expected_result
- source_requirement

Không bịa giá trị mà Rule Matrix không có.
Nếu endpoint là UNMAPPED: giữ nguyên UNMAPPED, không đoán URL/method.

==================================================
=== CATEGORY TREE ORDER
==================================================
AUTHENTICATION → 1. Kiểm tra Xác thực
AUTHORIZATION → 2. Kiểm tra Phân quyền
METHOD_URL → 3. Kiểm tra Method & URL
REQUEST_VALIDATION → 4. Kiểm tra Validate Request
HAPPY_PATH → 5. Kiểm tra Luồng thành công
BUSINESS_RULE → 6. Kiểm tra Business Rules
RESPONSE_VALIDATION → 7. Kiểm tra Response
INTEGRATION → 8. Kiểm tra Tích hợp
EXCEPTION → 9. Kiểm tra Ngoại lệ
Chỉ tạo category có rule.

==================================================
=== CẤU TRÚC OUTPUT
==================================================
- [Module]
  - [METHOD] [Endpoint Path]
    - [Nhóm kiểm thử]
      - TC_API_001 - [METHOD] [Endpoint] - [rule_name]
        - Pre-condition: chỉ điều kiện được Rule Matrix cung cấp hoặc điều kiện tối thiểu không suy diễn
          - Steps & Data test: 1. Chuẩn bị request theo test_condition -> 2. Gọi đúng method/endpoint trong Matrix -> 3. Truyền dữ liệu/header/param/body đúng target/rule -> 4. Send Request
            - Response (Kết quả mong đợi): dùng đúng expected_result; nếu Matrix không có exact status/message thì KHÔNG tự thêm

Pre-condition là child trực tiếp của Test Case.
Steps & Data test là child trực tiếp của Pre-condition.
Response là child trực tiếp của Steps & Data test.
Steps/Response nằm trên một dòng, không Enter tạo node mới.

==================================================
=== OUTPUT
==================================================
CHỈ plain text dạng cây bằng dấu "-".
KHÔNG JSON.
KHÔNG markdown code block.
KHÔNG giải thích/thống kê.
Tất cả bằng tiếng Việt.

=== API TEST DESIGN / RULE MATRIX ===
{content}
"""

# Alias để code cũ/reference cũ không bị NameError; flow V3.2 dùng trực tiếp prompt mới.
PROMPT_AGENT1_EXTRACT_API_JSON_SCOPE = PROMPT_API_AGENT1_RULE_MATRIX
PROMPT_AGENT2_GEN_API_XMIND_FROM_JSON = PROMPT_API_AGENT2_RENDERER
PROMPT_API_SPEC = PROMPT_API_AGENT2_RENDERER

# ==========================================
# 5. GIAO DIỆN CHÍNH (STREAMLIT TABS)
# ==========================================
tab_ui, tab_api = st.tabs(["📱 Web & Mobile App (UI/UX)", "🔌 RESTful / SOAP API Testing"])

# --- TAB 1: WEB & MOBILE APP (UI/UX) ---
with tab_ui:
    st.subheader("📌 Kiểm thử Giao diện & Luồng Người dùng (Web/App)")
    st.caption("Hỗ trợ bóc tách SRS PDF/MD hoặc Tóm tắt Scope Plan (JSON) có tương tác Review trực tiếp trước khi sinh XMind.")

    ui_workflow_mode = st.radio(
        "🎯 Chọn phương thức xử lý UI:",
        ["1-Click Auto Batching (Luồng Cũ - Tự động cắt .md)", "Interactive Review Plan (Luồng Mới - Review JSON Scope & Rule)"],
        index=1,
        key="ui_workflow_mode"
    )

    uploaded_ui_file = st.file_uploader("Tải lên tài liệu SRS / Figma Layout Text (.pdf, .md, .txt)", type=["pdf", "md", "txt"], key="main_ui_file_uploader")

    if uploaded_ui_file is not None:
        st.markdown("---")
        
        # LUỒNG 1: 1-CLICK AUTO BATCHING
        if ui_workflow_mode.startswith("1-Click"):
            st.write("### 🤖 Bước 1: Agent 1 - Đọc & Cắt tài liệu bằng Python Parser (Instant)")
            
            if st.button("🚀 Kích hoạt Agent 1 (Bóc tách tài liệu -> File .md)", key="btn_agent1_ui_main"):
                with st.status("Python Parser đang phân tích Heading và cắt nhỏ file...", expanded=True) as status:
                    try:
                        base_name = uploaded_ui_file.name.rsplit('.', 1)[0]
                        file_bytes = uploaded_ui_file.read()
                        
                        if uploaded_ui_file.name.lower().endswith(".pdf"):
                            reader = PdfReader(io.BytesIO(file_bytes))
                            raw_content = "\n\n".join([p.extract_text() for p in reader.pages if p.extract_text()])
                        else:
                            try:
                                raw_content = file_bytes.decode("utf-8")
                            except UnicodeDecodeError:
                                raw_content = file_bytes.decode("latin-1")

                        extracted_files, zip_path = python_smart_split_md(raw_content, base_name, max_chunk_size=5000)

                        st.session_state.split_ui_files = extracted_files
                        st.session_state.zip_ui_path = zip_path
                        st.session_state.step_ui = 2
                        status.update(label=f"✅ Đã cắt thành công {len(extracted_files)} phần nhỏ!", state="complete")
                    except Exception as e:
                        status.update(label="❌ Lỗi khi bóc tách file!", state="error")
                        st.error(str(e))

            if "split_ui_files" in st.session_state and st.session_state.split_ui_files:
                st.success(f"🎉 Đã hoàn tất bóc tách thành {len(st.session_state.split_ui_files)} phần nhỏ.")
                
                with open(st.session_state.zip_ui_path, "rb") as zf:
                    st.download_button(
                        label="📥 Tải xuống trọn bộ file Markdown con (.zip)",
                        data=zf,
                        file_name=os.path.basename(st.session_state.zip_ui_path),
                        mime="application/zip",
                        key="dl_split_ui_zip_main"
                    )

                st.write("📌 **Danh sách các file .md đã cắt theo từng Màn hình / Part:**")
                for item in st.session_state.split_ui_files:
                    with st.expander(f"📄 **{item.get('screen_name', item.get('file_name'))}** ({item['char_count']} chars)"):
                        st.code(item['content'], language="markdown")

                st.markdown("---")
                st.write("### 🤖 Bước 2: Agent 2 - Sinh Test Case UI Chi Tiết Cho Từng Màn Hình")
                
                if st.button("🚀 Kích hoạt Agent 2 (Chạy Multi-Agent Batching UI)", type="primary", key="btn_agent2_ui_main"):
                    if not api_key:
                        st.error("⚠️ Vui lòng nhập API Key ở thanh bên trái!")
                    else:
                        ui_items = st.session_state.split_ui_files
                        progress_bar = st.progress(0)
                        status_text = st.empty()
                        full_parsed_tree = []

                        for index, item in enumerate(ui_items):
                            screen_name = item.get('screen_name', f"Screen_{index+1}")
                            screen_content = item['content']
                            
                            status_text.info(f"⏳ [{index + 1}/{len(ui_items)}] Agent 2 đang phân tích: **{screen_name}** ({len(screen_content)} chars)...")
                            
                            ok, screen_result = call_qwen_max_agent(
                                content=screen_content,
                                api_key=api_key,
                                base_url=base_url,
                                model=model_name,
                                prompt_template=PROMPT_WEB_UI,
                                max_tokens=16384
                            )

                            if ok and screen_result.strip():
                                full_parsed_tree.append(screen_result.strip())
                            else:
                                st.warning(f"⚠️ Phần '{screen_name}' bị lỗi/timeout, bỏ qua.")

                            progress_bar.progress((index + 1) / len(ui_items))

                        complete_tree_text = "\n\n".join(full_parsed_tree)
                        st.session_state.parsed_tree_ui = complete_tree_text
                        st.session_state.step_ui = 3
                        status_text.success(f"✅ Hoàn tất sinh Test Case UI cho {len(ui_items)} phần!")

        # LUỒNG 2: INTERACTIVE REVIEW PLAN
        else:
            st.write("### 🤖 Bước 1: Agent 1 V3.1 - QA Brain Batching & Rule Matrix JSON")
            
            if st.button("🚀 Kích hoạt Agent 1 (Tạo JSON Rule Matrix V3.1)", key="btn_agent1_interactive"):
                if not api_key:
                    st.error("⚠️ Vui lòng nhập API Key!")
                else:
                    # Không để state cũ làm người dùng hiểu nhầm khi run mới fail.
                    st.session_state.pop("interactive_scope_json", None)
                    st.session_state.pop("raw_json_fallback", None)
                    st.session_state.pop("raw_agent1_response", None)
                    st.session_state.pop("agent1_pipeline_summary", None)

                    raw_ui_text = extract_text_from_file(uploaded_ui_file)
                    base_filename = uploaded_ui_file.name.rsplit('.', 1)[0]
                    st.info(
                        f"📚 Tài liệu: {len(raw_ui_text):,} chars. "
                        f"V3.1 sẽ semantic-chunk + recursive split nếu chạm Max Tokens; không bỏ qua chunk lỗi."
                    )

                    progress_bar = st.progress(0)
                    status_text = st.empty()

                    def _agent1_progress(done, total, message):
                        progress_bar.progress(min(1.0, done / max(1, total)))
                        status_text.info(message)

                    if "agent1_chunk_cache" not in st.session_state:
                        st.session_state.agent1_chunk_cache = {}

                    with st.spinner("Agent 1 V3.1 đang phân tích tài liệu theo batch semantic..."):
                        ok, final_matrix, pipeline_summary = run_agent1_document_pipeline(
                            raw_text=raw_ui_text,
                            base_filename=base_filename,
                            api_key=api_key,
                            base_url=base_url,
                            model=model_name,
                            prompt_template=PROMPT_AGENT1_EXTRACT_RULE_MATRIX,
                            cache=st.session_state.agent1_chunk_cache,
                            progress_callback=_agent1_progress,
                        )

                    st.session_state.agent1_pipeline_summary = pipeline_summary

                    if ok and final_matrix is not None:
                        progress_bar.progress(1.0)
                        status_text.success("✅ Agent 1 V3.1 hoàn tất toàn bộ document.")
                        st.session_state.interactive_scope_json = final_matrix
                        st.session_state.raw_agent1_response = json.dumps(final_matrix, ensure_ascii=False, indent=2)

                        merge_stats = pipeline_summary.get("merge", {})
                        st.success(
                            f"✅ Rule Matrix hợp lệ | Screens={merge_stats.get('screens', 0)} | "
                            f"Rules={merge_stats.get('final_rules', 0)} | "
                            f"InitialChunks={pipeline_summary.get('initial_chunks', 0)} | "
                            f"LeafMatrices={pipeline_summary.get('leaf_matrices', 0)} | "
                            f"Time={pipeline_summary.get('elapsed', 0)}s"
                        )
                    else:
                        status_text.error("❌ Agent 1 V3.1 chưa hoàn thành toàn bộ tài liệu. Không đưa partial Rule Matrix sang Agent 2.")
                        st.error("❌ Có chunk không thể xử lý hoàn chỉnh sau retry/split. Xem diagnostics bên dưới.")

                    with st.expander("🧪 Diagnostics Agent 1 V3.1", expanded=not ok):
                        st.json(pipeline_summary)
                        st.download_button(
                            "📥 Tải diagnostics Agent 1",
                            data=json.dumps(pipeline_summary, ensure_ascii=False, indent=2),
                            file_name="agent1_v31_diagnostics.json",
                            mime="application/json",
                            key="dl_agent1_v31_diag",
                        )

            if "interactive_scope_json" in st.session_state:
                st.markdown("---")
                st.write("### 🔍 Bước 2: QA Lead Review Rule Matrix & Chọn Màn Hình")
                st.info("💡 Bạn có thể tích chọn Màn hình và chỉnh sửa trực tiếp từng Test Rule dưới dạng JSON. Agent 2 chỉ render Rule Matrix, không tự suy luận thêm.")

                plan_data = st.session_state.interactive_scope_json
                approved_screens = []

                for idx, screen in enumerate(plan_data.get("screens", [])):
                    with st.expander(f"📌 Màn hình {idx+1}: {screen.get('screen_name', 'Unnamed')}", expanded=True):
                        is_selected = st.checkbox(f"Đưa màn hình '{screen.get('screen_name')}' vào Scope Gen Test Case", value=True, key=f"chk_inter_{idx}")
                        
                        edited_json_str = st.text_area(
                            "TEST DESIGN / RULE MATRIX của màn hình (Sửa rule nếu cần):",
                            value=json.dumps(screen, ensure_ascii=False, indent=2),
                            height=300,
                            key=f"txt_json_inter_{idx}"
                        )
                        
                        if is_selected:
                            try:
                                updated_screen_json = json.loads(edited_json_str)
                                approved_screens.append(updated_screen_json)
                            except json.JSONDecodeError as e:
                                st.error(
                                    f"❌ JSON màn hình {idx+1} lỗi tại line {e.lineno}, col {e.colno}: {e.msg}"
                                )

                st.markdown("---")
                st.write("### 🤖 Bước 3: Agent 2 V3.1 - Batch Renderer Từ Rule Matrix Đã Phê Duyệt")
                
                if st.button("🚀 Chốt Rule Matrix & Kích Hoạt Agent 2 Gen XMind", type="primary", key="btn_agent2_interactive"):
                    if not api_key:
                        st.error("⚠️ Vui lòng nhập API Key!")
                    elif not approved_screens:
                        st.error("⚠️ Bạn chưa chọn Màn hình nào!")
                    else:
                        st.session_state.pop("parsed_tree_ui", None)
                        st.session_state.pop("agent2_pipeline_summary", None)

                        progress_bar_a2 = st.progress(0)
                        status_text_a2 = st.empty()

                        def _agent2_progress(done, total, message):
                            progress_bar_a2.progress(min(1.0, done / max(1, total)))
                            status_text_a2.info(message)

                        if "agent2_batch_cache" not in st.session_state:
                            st.session_state.agent2_batch_cache = {}

                        with st.spinner("Agent 2 đang render Rule Matrix theo batch; tự chia nhỏ nếu chạm Max Tokens..."):
                            ok, tc_res, agent2_summary = run_agent2_rule_matrix_pipeline(
                                approved_screens=approved_screens,
                                api_key=api_key,
                                base_url=base_url,
                                model=model_name,
                                prompt_template=PROMPT_AGENT2_GEN_XMIND_FROM_RULE_MATRIX,
                                cache=st.session_state.agent2_batch_cache,
                                progress_callback=_agent2_progress,
                            )

                        st.session_state.agent2_pipeline_summary = agent2_summary

                        if ok and tc_res.strip():
                            progress_bar_a2.progress(1.0)
                            status_text_a2.success("✅ Agent 2 hoàn tất toàn bộ Rule Matrix.")
                            st.session_state.parsed_tree_ui = tc_res
                            merge2 = agent2_summary.get("merge", {})
                            st.success(
                                f"✅ Đã render xong | Batches={agent2_summary.get('initial_batches', 0)} | "
                                f"LeafOutputs={merge2.get('agent2_leaf_outputs', 0)} | "
                                f"TestCases={merge2.get('testcases_renumbered', 0)} | "
                                f"Time={agent2_summary.get('elapsed', 0)}s"
                            )
                        else:
                            status_text_a2.error("❌ Agent 2 chưa render đủ toàn bộ Rule Matrix.")
                            st.error("❌ Có batch Agent 2 bị lỗi và pipeline đã dừng để tránh xuất bộ testcase thiếu.")

                        with st.expander("🧪 Diagnostics Agent 2 V3.1", expanded=not ok):
                            st.json(agent2_summary)
                            st.download_button(
                                "📥 Tải diagnostics Agent 2",
                                data=json.dumps(agent2_summary, ensure_ascii=False, indent=2),
                                file_name="agent2_v31_diagnostics.json",
                                mime="application/json",
                                key="dl_agent2_v31_diag",
                            )

        # BƯỚC XUẤT FILE CHUNG CHO CẢ 2 LUỒNG UI
        if "parsed_tree_ui" in st.session_state and st.session_state.parsed_tree_ui:
            st.markdown("---")
            st.write("### 📦 Xuất Kịch Bản Kiểm Thử (XMind / Excel)")
            
            output_xmind_path = os.path.join(tempfile.gettempdir(), f"{uploaded_ui_file.name.rsplit('.', 1)[0]}_UI_TestCases.xmind")
            ok_xmind, err_xmind = create_xmind_from_text(st.session_state.parsed_tree_ui, output_xmind_path, root_title="Bộ Test Case UI/UX & Web App")
            
            if ok_xmind:
                st.success("✅ Đã tạo thành công bộ Test Case cho Web/App!")
                
                with st.expander("🔍 Xem trước dạng Text cây thụt lề"):
                    st.code(st.session_state.parsed_tree_ui, language="text")
                    
                col_xmind, col_excel = st.columns(2)
                
                with col_xmind:
                    with open(output_xmind_path, "rb") as f:
                        st.download_button(
                            label="📥 Tải xuống sơ đồ UI (.XMind)",
                            data=f,
                            file_name=f"{uploaded_ui_file.name.rsplit('.', 1)[0]}_UI_TestCases.xmind",
                            mime="application/octet-stream",
                            key="dl_ui_main"
                        )
                
                with col_excel:
                    excel_bytes = convert_tree_to_ui_excel(st.session_state.parsed_tree_ui)
                    st.download_button(
                        label="📊 Tải xuống bảng Excel (.XLSX)",
                        data=excel_bytes,
                        file_name=f"{uploaded_ui_file.name.rsplit('.', 1)[0]}_UI_TestCases.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_ui_excel_main"
                    )
            else:
                st.error("❌ Không thể tạo file XMind. Chi tiết lỗi:")
                st.code(err_xmind, language="bash")

# --- TAB 2: API TESTING — V3.2 SCALABLE RULE MATRIX ---
with tab_api:
    st.subheader("📌 Kiểm thử RESTful API & Services — API QA Brain V3.2")
    st.caption("Agent 1 đọc API Spec + BA theo batch, áp dụng API QA Rules và tạo API Rule Matrix. Agent 2 chỉ render 1 Rule = 1 Test Case. Hỗ trợ tài liệu lớn bằng recursive batching.")

    col1_api, col2_api = st.columns(2)
    with col1_api:
        api_design_file = st.file_uploader(
            "1. Tài liệu Thiết kế API (OpenAPI/Swagger/Postman/Markdown)",
            type=["pdf", "md", "json", "yaml", "yml", "txt"], key="api_design_file"
        )
    with col2_api:
        api_ba_file = st.file_uploader(
            "2. Tài liệu Nghiệp vụ BA (nếu có)",
            type=["pdf", "md", "json", "yaml", "yml", "txt"], key="api_ba_file"
        )

    if api_design_file or api_ba_file:
        st.markdown("---")
        st.write("### 🤖 Bước 1: API Agent 1 — QA Brain → API Rule Matrix")

        if st.button("🚀 Kích hoạt API Agent 1 (Tạo API Rule Matrix)", key="btn_agent1_api_v32"):
            if not api_key:
                st.error("⚠️ Vui lòng nhập API Key!")
            else:
                for key in ["api_rule_matrix", "api_agent1_summary", "api_parsed_tree", "api_agent2_summary"]:
                    st.session_state.pop(key, None)

                api_design_text = extract_text_from_file(api_design_file)
                api_ba_text = extract_text_from_file(api_ba_file)
                base_name = (
                    (api_design_file.name.rsplit('.', 1)[0] if api_design_file else None)
                    or (api_ba_file.name.rsplit('.', 1)[0] if api_ba_file else "API_Document")
                )

                if "api_agent1_cache_v32" not in st.session_state:
                    st.session_state.api_agent1_cache_v32 = {}

                progress = st.progress(0)
                status = st.empty()

                def _api_a1_progress(done, total, message):
                    progress.progress(done / total if total else 0)
                    status.info(message)

                with st.spinner("API Agent 1 đang đọc Spec/BA theo batch và áp dụng QA Test Design Rules..."):
                    ok, matrix, summary = run_api_agent1_document_pipeline(
                        design_text=api_design_text,
                        ba_text=api_ba_text,
                        base_filename=base_name,
                        api_key=api_key,
                        base_url=base_url,
                        model=model_name,
                        prompt_template=PROMPT_API_AGENT1_RULE_MATRIX,
                        cache=st.session_state.api_agent1_cache_v32,
                        progress_callback=_api_a1_progress,
                    )

                st.session_state.api_agent1_summary = summary
                if ok and matrix:
                    st.session_state.api_rule_matrix = matrix
                    merge = summary.get("merge", {})
                    progress.progress(1.0)
                    status.success(
                        f"✅ API Agent 1 hoàn tất: {merge.get('modules',0)} modules | "
                        f"{merge.get('endpoints',0)} endpoints | {merge.get('final_rules',0)} rules | "
                        f"Unmapped endpoints={merge.get('unmapped_endpoints',0)}"
                    )
                else:
                    status.error("❌ API Agent 1 thất bại ở một chunk. Không publish Rule Matrix thiếu coverage.")
                    st.error("Xem diagnostics bên dưới để biết chunk/parse/schema/max-token lỗi.")

        if "api_agent1_summary" in st.session_state:
            with st.expander("🧪 Diagnostics API Agent 1", expanded=False):
                st.json(st.session_state.api_agent1_summary)

        if "api_rule_matrix" in st.session_state:
            st.markdown("---")
            st.write("### 🔍 Bước 2: QA Lead Review API Rule Matrix & Chọn Endpoint")
            st.info("Bạn có thể bỏ chọn endpoint hoặc sửa JSON từng endpoint. Agent 2 sẽ chỉ render các rule đã chốt, không tự đọc lại Spec/BA.")

            matrix = st.session_state.api_rule_matrix
            approved_modules_map = OrderedDict()
            endpoint_counter = 0

            for midx, module in enumerate(matrix.get("api_modules", []), start=1):
                module_name = module.get("module_name", "UNMAPPED")
                st.markdown(f"#### 📦 Module {midx}: {module_name}")
                for eidx, endpoint in enumerate(module.get("endpoints", []), start=1):
                    endpoint_counter += 1
                    method = endpoint.get("method", "UNMAPPED")
                    path = endpoint.get("endpoint_path", "UNMAPPED")
                    rules_count = len(endpoint.get("test_rules", []))
                    with st.expander(f"🔌 {method} {path} — {rules_count} rules", expanded=False):
                        selected = st.checkbox(
                            f"Đưa endpoint {method} {path} vào bộ Test Case",
                            value=True, key=f"api_ep_sel_{midx}_{eidx}"
                        )
                        edited = st.text_area(
                            "API Rule Matrix endpoint (có thể sửa):",
                            value=json.dumps(endpoint, ensure_ascii=False, indent=2),
                            height=360,
                            key=f"api_ep_json_{midx}_{eidx}"
                        )
                        if selected:
                            try:
                                ep_json = json.loads(edited)
                                temp = {"api_test_design_version": "3.2", "api_modules": [{"module_name": module_name, "endpoints": [ep_json]}]}
                                valid, diag = validate_api_rule_matrix_schema(temp, strict=True)
                                if not valid:
                                    st.error(f"❌ Endpoint JSON/schema lỗi: {diag}")
                                else:
                                    if module_name not in approved_modules_map:
                                        approved_modules_map[module_name] = {"module_name": module_name, "endpoints": []}
                                    approved_modules_map[module_name]["endpoints"].append(ep_json)
                            except Exception as e:
                                st.error(f"❌ JSON endpoint không hợp lệ: {e}")

            approved_modules = list(approved_modules_map.values())
            approved_rules = sum(len(ep.get("test_rules", [])) for m in approved_modules for ep in m.get("endpoints", []))
            st.caption(f"Đã chọn {sum(len(m.get('endpoints', [])) for m in approved_modules)} / {endpoint_counter} endpoints — {approved_rules} rules")

            if st.button("✅ Chốt API Rule Matrix & Kích hoạt Agent 2", type="primary", key="btn_api_agent2_v32"):
                if not approved_modules:
                    st.error("⚠️ Chưa chọn endpoint hợp lệ nào.")
                else:
                    if "api_agent2_cache_v32" not in st.session_state:
                        st.session_state.api_agent2_cache_v32 = {}
                    progress2 = st.progress(0)
                    status2 = st.empty()

                    def _api_a2_progress(done, total, message):
                        progress2.progress(done / total if total else 0)
                        status2.info(message)

                    with st.spinner("API Agent 2 đang render API Rule Matrix theo batch..."):
                        ok, tree, summary2 = run_api_agent2_rule_matrix_pipeline(
                            approved_modules=approved_modules,
                            api_key=api_key,
                            base_url=base_url,
                            model=model_name,
                            prompt_template=PROMPT_API_AGENT2_RENDERER,
                            cache=st.session_state.api_agent2_cache_v32,
                            progress_callback=_api_a2_progress,
                        )
                    st.session_state.api_agent2_summary = summary2
                    if ok and tree.strip():
                        st.session_state.api_parsed_tree = tree
                        progress2.progress(1.0)
                        status2.success(f"✅ API Agent 2 hoàn tất — {summary2.get('merge',{}).get('testcases_renumbered',0)} testcase")
                    else:
                        status2.error("❌ API Agent 2 thất bại ở một batch; không publish output thiếu.")

        if "api_agent2_summary" in st.session_state:
            with st.expander("🧪 Diagnostics API Agent 2", expanded=False):
                st.json(st.session_state.api_agent2_summary)

        if "api_parsed_tree" in st.session_state and st.session_state.api_parsed_tree:
            st.markdown("---")
            st.write("### 📦 Bước 3: Xuất API Test Case (.xmind & .xlsx)")
            out_filename = "API_TestCases_V3_2.xmind"
            output_xmind_path = os.path.join(tempfile.gettempdir(), out_filename)
            ok_xmind, err_xmind = create_xmind_from_text(
                st.session_state.api_parsed_tree,
                output_xmind_path,
                root_title="Bộ Test Case API Testing V3.2"
            )
            if ok_xmind:
                st.success("✅ Đã tạo thành công bộ API Test Case!")
                with st.expander("🔍 Xem trước cây API Test Case"):
                    st.code(st.session_state.api_parsed_tree, language="text")
                col_api_xmind, col_api_excel = st.columns(2)
                with col_api_xmind:
                    with open(output_xmind_path, "rb") as f:
                        st.download_button(
                            "📥 Tải API XMind", data=f, file_name=out_filename,
                            mime="application/octet-stream", key="dl_api_xmind_v32"
                        )
                with col_api_excel:
                    api_excel_bytes = convert_tree_to_api_excel(st.session_state.api_parsed_tree)
                    st.download_button(
                        "📊 Tải API Excel", data=api_excel_bytes,
                        file_name="API_TestCases_V3_2.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        key="dl_api_excel_v32"
                    )
            else:
                st.error(f"❌ Lỗi khi đóng gói XMind: {err_xmind}")

